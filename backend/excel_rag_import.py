"""Import the synthetic IT-flow Excel batch with row-level RAG provenance."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

import server
from backend import public_dataset_import


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEMO_ROOT = Path("/Users/irene/Documents/Codex/2026-07-22/dfd/outputs/it_data_flow_demo")
DEFAULT_WORKBOOKS = tuple(
    path
    for path in sorted((DEMO_ROOT / "01_upload" / "source_workbooks").glob("*.xlsx"))
    if not path.name.startswith("~$")
)
PACKETS_PATH = DEMO_ROOT / "06_ready_for_mia" / "mia_knowledge_packets.jsonl"
BATCH_KEY = "synthetic-it-data-flow-20260728"
ROUTE_KEYS = {"rte_dover_calais": "dover-calais", "rte_newhaven_dieppe": "newhaven-dieppe", "rte_amsterdam_newcastle": "newcastle-ijmuiden"}


def index_workbooks(workbooks):
    indexed = {}
    for workbook_path in workbooks:
        workbook = load_workbook(workbook_path, read_only=False, data_only=True)
        for sheet in workbook.worksheets:
            headers = [cell.value for cell in sheet[1]]
            for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not any(value is not None for value in values):
                    continue
                payload = {str(header): value for header, value in zip(headers, values) if header is not None}
                raw_id = next((str(value) for key, value in payload.items() if key.endswith("_row_id") and value), None)
                if raw_id:
                    indexed[raw_id] = {"workbook": workbook_path.name, "sheet": sheet.title, "excel_row": excel_row, "payload": payload}
        workbook.close()
    return indexed


def load_mia_packets(path=PACKETS_PATH):
    return [json.loads(line) for line in Path(path).read_text(encoding="utf-8").splitlines() if line.strip()]


def packet_to_rag_record(packet, source_row):
    raw_id = packet["lineage"]["raw_record_id"]
    evidence = str(packet.get("evidence") or "Structured source event")
    content = (
        f"Synthetic IT data-flow event {packet['knowledge_event_id']}: {packet['event_intent']}. "
        f"Route: {packet.get('route_name', 'All signals')}. Product: {packet.get('product_signal', 'not stated')}. "
        f"Context: {packet.get('travel_context', 'not stated')}. Evidence: {evidence}"
    )
    source_key = f"synthetic-it-flow-{packet['source_family']}"
    return {
        "source_key": source_key,
        "source_name": f"Synthetic IT Flow - {packet['source_family'].replace('_', ' ').title()}",
        "source_type": "synthetic_internal_demo",
        "company_name": "DFDS",
        "route_key": ROUTE_KEYS.get(packet.get("route_id"), "all"),
        "evidence_kind": f"it_flow_{packet['entity_grain']}",
        "original_content": content,
        "translated_content": "Synthetic, de-identified and consent-permitted IT data-flow demonstration event.",
        "published_at": "2026-07-28",
        "metadata": {
            "import_batch": BATCH_KEY,
            "knowledge_event_id": packet["knowledge_event_id"],
            "canonical_fact_id": packet["canonical_fact_id"],
            "raw_record_id": raw_id,
            "source_packet": packet["lineage"]["source_packet"],
            "workbook": source_row["workbook"],
            "sheet": source_row["sheet"],
            "excel_row": source_row["excel_row"],
            "consent_state": packet["consent_state"],
            "de_identified": packet["de_identified"],
            "classification_confidence": packet["classification_confidence"],
            "synthetic_source": True,
            "source_version": "2026-07-28",
        },
        "keywords": [
            str(value)
            for value in ("synthetic", "it-data-flow", packet["source_family"], packet["event_intent"], packet.get("route_name") or "all")
            if value
        ],
        "source_tier": "synthetic_demo",
        "is_synthetic": True,
        "source_version": "2026-07-28",
    }


def run_import(*, dry_run=False):
    source_rows = index_workbooks(DEFAULT_WORKBOOKS)
    packets = load_mia_packets()
    records = [packet_to_rag_record(packet, source_rows[packet["lineage"]["raw_record_id"]]) for packet in packets]
    summary = {"raw_excel_rows": len(source_rows), "mia_ready_records": len(records), "batch_key": BATCH_KEY}
    if dry_run:
        return summary
    server.load_local_env()
    conn = server.connect_db()
    if conn is None:
        raise RuntimeError("PostgreSQL is not connected")
    try:
        server.execute_schema(conn)
        with conn.cursor() as cur:
            cur.execute("INSERT INTO ingestion_batches (batch_key, source_label, is_synthetic) VALUES (%s, %s, TRUE) ON CONFLICT (batch_key) DO UPDATE SET source_label = EXCLUDED.source_label RETURNING id", (BATCH_KEY, "Synthetic IT Data Flow Excel batch"))
            batch_id = cur.fetchone()["id"]
            for raw_id, row in source_rows.items():
                cur.execute("INSERT INTO raw_source_rows (batch_id, raw_record_key, workbook_name, sheet_name, excel_row, raw_payload) VALUES (%s, %s, %s, %s, %s, %s::jsonb) ON CONFLICT (batch_id, raw_record_key) DO UPDATE SET raw_payload = EXCLUDED.raw_payload", (batch_id, raw_id, row["workbook"], row["sheet"], row["excel_row"], json.dumps(row["payload"], default=str)))
        conn.commit()
        summary.update(public_dataset_import.upsert_records(conn, records))
        return summary
    finally:
        conn.close()


def main(argv=None):
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)
    print(json.dumps(run_import(dry_run=args.dry_run), ensure_ascii=False))


if __name__ == "__main__":
    main()
