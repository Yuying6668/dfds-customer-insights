"""Safe file-batch intake and bounded spreadsheet cleaning for IT Data Flow."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import shutil
import uuid
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook


ALLOWED_SUFFIXES = {".xlsx", ".xls", ".csv", ".txt", ".pdf", ".doc", ".docx"}
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_BATCH_BYTES = 30 * 1024 * 1024
MAX_BATCH_FILES = 10
PREVIEW_ROW_LIMIT = 100
RESTRICTED_FIELD_PATTERN = re.compile(
    r"email|customer_?ref|crm_?contact|loyalty_?no|account_?hash|device_?hash|"
    r"alias_?value|phone|address|first_?name|last_?name",
    re.IGNORECASE,
)
MISSING_VALUE_TOKENS = {"", "-", "--", "?", "n/a", "na", "null", "none", "not available", "unknown"}
MISSING_VALUE_NUMBERS = {-999, -9999}
DATE_FIELDS = {
    "travel_date", "published", "review_date", "departure_date", "scheduled_departure",
    "payment_date", "event_time", "source_received_at", "consent_source_timestamp",
    "created_at", "updated_at",
}
DATE_FIELD_PATTERN = re.compile(r"(?:^|_)(?:date|datetime|timestamp)(?:$|_)", re.IGNORECASE)
KEY_FIELD_PATTERN = re.compile(r"(?:booking|payment|leg|response|case|customer|crm|loyalty|account|device).*(?:id|ref|no|hash)$", re.IGNORECASE)
AMBIGUOUS_REFERENCE_FIELDS = {"case_or_partner_ref", "customer_ref", "email_or_loyalty"}
ROUTE_ALIASES = {
    "dvrcalais": "Dover-Calais",
    "dovercalais": "Dover-Calais",
    "dkdover": "Dover-Calais",
    "newhavendieppe": "Newhaven-Dieppe",
    "newcastleijmuiden": "Newcastle-IJmuiden",
    "newcastleamsterdam": "Newcastle-IJmuiden",
    "copenhagenoslo": "Copenhagen-Oslo",
    "jersey": "Jersey",
}


class UploadValidationError(ValueError):
    """A client-supplied upload batch is structurally unsafe or unsupported."""


def _utc_timestamp():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _batch_version():
    return f"v{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')}"


def _normalise_header(value, index, used_headers):
    base = re.sub(r"[^\w]+", "_", str(value or "").strip().casefold(), flags=re.UNICODE).strip("_")
    base = base or f"column_{index + 1}"
    candidate = base
    suffix = 2
    while candidate in used_headers:
        candidate = f"{base}_{suffix}"
        suffix += 1
    used_headers.add(candidate)
    return candidate


def _is_missing_value(value):
    if value is None:
        return True
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value in MISSING_VALUE_NUMBERS
    return isinstance(value, str) and value.strip().casefold() in MISSING_VALUE_TOKENS | {"-999", "-999.0", "-9999", "-9999.0"}


def _normalise_cell(value):
    if _is_missing_value(value):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (bool, int, float)):
        return value
    return str(value).strip()


def _standardize_route(value):
    compact = re.sub(r"[^a-z]", "", str(value).casefold())
    if compact in ROUTE_ALIASES:
        return ROUTE_ALIASES[compact], True, False
    return "Unmapped", False, True


def _standardize_date(value):
    source_text = str(value).strip()
    timezone_label = None
    if re.search(r"(?:^|\s)(UTC|GMT)$", source_text, re.I):
        timezone_label = "UTC"
        text = re.sub(r"\s+(UTC|GMT)$", "+00:00", source_text, flags=re.I)
    elif source_text.endswith("Z"):
        timezone_label = "UTC"
        text = source_text[:-1] + "+00:00"
    else:
        offset_match = re.search(r"([+-]\d{2}:?\d{2})$", source_text)
        timezone_label = f"UTC{offset_match.group(1)}" if offset_match else None
        text = source_text
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        parsed = None
    if parsed is None:
        for pattern in (
            "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m-%d-%Y", "%d.%m.%Y",
            "%d-%m-%Y %H:%M", "%d/%m/%Y %H:%M", "%B %d %Y", "%b %d %Y",
        ):
            try:
                parsed = datetime.strptime(text, pattern)
                break
            except ValueError:
                continue
    if parsed is None:
        return source_text, False, True, timezone_label
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    else:
        parsed = parsed.astimezone(timezone.utc)
    return parsed.strftime("%Y-%m-%dT%H:%M:%SZ"), True, False, timezone_label


def _is_date_field(column):
    return column in DATE_FIELDS or bool(DATE_FIELD_PATTERN.search(column))


def _standardize_amount(value):
    match = re.fullmatch(r"(EUR|GBP|DKK)\s+([\d.,]+)", str(value).strip(), re.I)
    if not match:
        return str(value).strip(), False, True
    currency, raw_number = match.groups()
    if currency.upper() == "DKK":
        number = raw_number.replace(".", "").replace(",", ".")
    else:
        number = raw_number.replace(".", "").replace(",", ".") if raw_number.count(",") == 1 else raw_number
    try:
        return f"{float(number):.2f} {currency.upper()}", True, False
    except ValueError:
        return str(value).strip(), False, True


def _standardize_rating(value):
    text = str(value).strip().replace(",", ".")
    if re.fullmatch(r"[1-5](?:\.0)?/5", text):
        text = text.split("/", 1)[0]
    try:
        score = float(text)
    except ValueError:
        return str(value).strip(), False, True
    return (f"{score:.1f}", True, False) if 0 <= score <= 5 else (str(value).strip(), False, True)


def _standardize_enum(column, value):
    normalized = str(value).strip().casefold()
    if column == "marketing_consent":
        values = {"y": "Yes", "yes": "Yes", "true": "Yes", "n": "No", "no": "No", "false": "No", "unknown": "Unknown"}
        if normalized in values:
            return values[normalized], True, values[normalized] == "Unknown"
    elif column == "platform":
        values = {"ios": "iOS", "android": "Android"}
        if normalized in values:
            return values[normalized], True, False
    elif column == "channel" and normalized in {"dfds app", "mia's cruises app", "mias cruises app"}:
        return "Mia's Cruises App", True, False
    return str(value).strip(), False, True


def _standardize_row(row):
    mapped_values = 0
    mapping_exceptions = 0
    checks = {"invalidDateValues": 0, "invalidAmountValues": 0, "invalidRatingValues": 0, "unmappedRouteValues": 0, "enumValuesStandardized": 0, "enumExceptions": 0}
    exception_details = []
    detected_timezones = []
    for column, value in list(row.items()):
        if value is None:
            continue
        if column in {"route", "route_text", "route_query", "route_mentioned", "crossing"}:
            row[column], mapped, exception = _standardize_route(value)
            checks["unmappedRouteValues"] += int(exception)
        elif _is_date_field(column):
            row[column], mapped, exception, detected_timezone = _standardize_date(value)
            if detected_timezone and mapped:
                detected_timezones.append(detected_timezone)
            checks["invalidDateValues"] += int(exception)
        elif column in {"gross_amount", "amount_text"}:
            row[column], mapped, exception = _standardize_amount(value)
            checks["invalidAmountValues"] += int(exception)
        elif column in {"rating", "score"}:
            row[column], mapped, exception = _standardize_rating(value)
            checks["invalidRatingValues"] += int(exception)
        elif column in {"marketing_consent", "platform", "channel"}:
            row[column], mapped, exception = _standardize_enum(column, value)
            checks["enumValuesStandardized"] += int(mapped)
            checks["enumExceptions"] += int(exception)
        else:
            continue
        mapped_values += int(mapped)
        mapping_exceptions += int(exception)
        if exception:
            reason = "Unmapped route" if column in {"route", "route_text", "route_query", "route_mentioned", "crossing"} else "Invalid date" if _is_date_field(column) else "Invalid amount" if column in {"gross_amount", "amount_text"} else "Invalid rating" if column in {"rating", "score"} else "Unrecognized value"
            exception_details.append({"field": column, "value": str(value).strip(), "reason": reason})
    if detected_timezones:
        row["timezone"] = " / ".join(dict.fromkeys(detected_timezones))
    return mapped_values, mapping_exceptions, checks, exception_details


def _masked_preview(rows, columns):
    restricted = {column for column in columns if RESTRICTED_FIELD_PATTERN.search(column)}
    if not restricted:
        return rows
    return [
        {
            column: "Restricted" if column in restricted and value is not None else value
            for column, value in row.items()
        }
        for row in rows
    ]


def clean_sheet_rows(rows: Iterable[Iterable[object]], preview_row_limit=PREVIEW_ROW_LIMIT, include_exception_rows=False):
    """Return stable cleaned columns and a masked, limited preview without altering source files."""
    row_iterator = iter(rows)
    header_values = next(row_iterator, ())
    if header_values is None:
        header_values = ()

    used_headers = set()
    columns = [_normalise_header(value, index, used_headers) for index, value in enumerate(header_values)]
    if not columns:
        return {
            "columns": [],
            "dataRows": 0,
            "previewRows": [],
            "cleaning": {
                "receivedRows": 0,
                "cleanedRows": 0,
                "blankRowsRemoved": 0,
                "duplicateRowsRemoved": 0,
                "missingValuesStandardized": 0,
                "mappedValues": 0,
                "mappingExceptions": 0,
                "previewRows": 0,
            },
        }

    cleaned_rows = []
    exception_rows = []
    row_signatures = set()
    blank_rows_removed = 0
    duplicate_rows_removed = 0
    missing_values_standardized = 0
    mapped_values = 0
    mapping_exceptions = 0
    quality_checks = {
        "invalidDateValues": 0,
        "invalidAmountValues": 0,
        "invalidRatingValues": 0,
        "unmappedRouteValues": 0,
        "enumValuesStandardized": 0,
        "enumExceptions": 0,
        "missingKeyValues": 0,
        "duplicateKeyValues": 0,
        "mixedAgeFormats": 0,
        "ambiguousReferenceColumns": sum(column in AMBIGUOUS_REFERENCE_FIELDS for column in columns),
    }
    key_values = {}
    age_value_formats = set()
    received_rows = 0

    for raw_values in row_iterator:
        received_rows += 1
        values = list(raw_values or ())[: len(columns)]
        values.extend([None] * (len(columns) - len(values)))
        row = {column: _normalise_cell(value) for column, value in zip(columns, values)}

        if not any(value is not None for value in row.values()):
            blank_rows_removed += 1
            continue

        missing_values_standardized += sum(_is_missing_value(value) for value in values)
        row_mapped_values, row_mapping_exceptions, row_checks, row_exception_details = _standardize_row(row)
        if row.get("timezone") and "timezone" not in columns:
            columns.append("timezone")
            for cleaned_row in cleaned_rows:
                cleaned_row["timezone"] = None
        mapped_values += row_mapped_values
        mapping_exceptions += row_mapping_exceptions
        for key, count in row_checks.items():
            quality_checks[key] += count

        for column, value in row.items():
            if KEY_FIELD_PATTERN.search(column):
                if value is None:
                    quality_checks["missingKeyValues"] += 1
                elif value in key_values.setdefault(column, set()):
                    quality_checks["duplicateKeyValues"] += 1
                else:
                    key_values[column].add(value)
            if column == "age_value" and value:
                age_value_formats.add("range" if re.fullmatch(r"\d{1,3}\s*-\s*\d{1,3}", str(value)) else "single")

        signature = json.dumps(row, ensure_ascii=False, sort_keys=True, default=str)
        if signature in row_signatures:
            duplicate_rows_removed += 1
            continue
        row_signatures.add(signature)
        cleaned_rows.append(row)
        if include_exception_rows and row_exception_details:
            exception_rows.append({**row, "exceptionDetails": row_exception_details})

    quality_checks["mixedAgeFormats"] = max(0, len(age_value_formats) - 1)

    preview_rows = _masked_preview(
        cleaned_rows if preview_row_limit is None else cleaned_rows[:preview_row_limit],
        columns,
    )
    result = {
        "columns": columns,
        "dataRows": received_rows,
        "previewRows": preview_rows,
        "cleaning": {
            "receivedRows": received_rows,
            "cleanedRows": len(cleaned_rows),
            "blankRowsRemoved": blank_rows_removed,
            "duplicateRowsRemoved": duplicate_rows_removed,
            "missingValuesStandardized": missing_values_standardized,
            "mappedValues": mapped_values,
            "mappingExceptions": mapping_exceptions,
            **quality_checks,
            "previewRows": len(preview_rows),
        },
    }
    if include_exception_rows:
        restricted = {column for column in columns if RESTRICTED_FIELD_PATTERN.search(column)}
        result["exceptionRows"] = [
            {
                **{column: "Restricted" if column in restricted and row.get(column) is not None else row.get(column) for column in columns},
                "exceptionDetails": [
                    {**detail, "value": "Restricted" if detail["field"] in restricted else detail["value"]}
                    for detail in row["exceptionDetails"]
                ],
            }
            for row in exception_rows
        ]
    return result


def _build_lifecycle(profile_files):
    sheets = [sheet for file_entry in profile_files for sheet in file_entry["sheets"]]
    received = sum(sheet["cleaning"]["receivedRows"] for sheet in sheets)
    cleaned = sum(sheet["cleaning"]["cleanedRows"] for sheet in sheets)
    exceptions = sum(sheet["cleaning"].get("mappingExceptions", 0) for sheet in sheets)
    fields = {column for sheet in sheets for column in sheet["columns"]}
    return [
        {"key": "upload", "label": "Upload", "state": "complete", "detail": f"{len(profile_files)} files received"},
        {"key": "validate", "label": "Validate", "state": "complete", "detail": f"{received:,} rows checked"},
        {"key": "classify", "label": "Classify", "state": "complete", "detail": f"{len(fields)} fields classified"},
        {"key": "map", "label": "Map", "state": "needs_review" if exceptions else "complete", "detail": f"{exceptions:,} mapping exceptions"},
        {"key": "review", "label": "Review", "state": "active", "detail": "Save cleaned data to approve"},
        {"key": "ready_for_mia", "label": "Ready for Mia", "state": "pending", "detail": f"{cleaned:,} cleaned rows"},
    ]


def _profile_excel(path):
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        return [
            {"name": worksheet.title, **clean_sheet_rows(worksheet.iter_rows(values_only=True))}
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _profile_text_table(path, original_name):
    with Path(path).open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
        return [{"name": Path(original_name).stem, **clean_sheet_rows(csv.reader(source))}]


def profile_uploaded_file(path, original_name):
    """Return profile metadata for an already persisted source file."""
    suffix = Path(original_name).suffix.casefold()
    profile = {
        "name": Path(original_name).name,
        "format": suffix.lstrip(".").upper() or "FILE",
        "parserState": "received_unparsed",
        "sheets": [],
    }
    if suffix == ".xlsx":
        profile["parserState"] = "profiled"
        profile["sheets"] = _profile_excel(path)
    elif suffix in {".csv", ".txt"}:
        profile["parserState"] = "profiled"
        profile["sheets"] = _profile_text_table(path, original_name)
    return profile


def _validate_file_parts(files):
    if not files or len(files) > MAX_BATCH_FILES:
        raise UploadValidationError("Upload one to ten files per batch")

    total_bytes = 0
    checksums = set()
    staged = []
    for original_name, payload in files:
        name = Path(str(original_name or "")).name
        suffix = Path(name).suffix.casefold()
        if not name or name in {".", ".."}:
            raise UploadValidationError("Each file needs a valid filename")
        if suffix not in ALLOWED_SUFFIXES:
            raise UploadValidationError(f"Unsupported file type: {suffix or 'unknown'}")
        if not isinstance(payload, bytes) or not payload:
            raise UploadValidationError(f"File is empty: {name}")
        if len(payload) > MAX_FILE_BYTES:
            raise UploadValidationError(f"File exceeds 10 MB: {name}")
        total_bytes += len(payload)
        if total_bytes > MAX_BATCH_BYTES:
            raise UploadValidationError("Batch exceeds 30 MB")
        checksum = hashlib.sha256(payload).hexdigest()
        if checksum in checksums:
            raise UploadValidationError(f"Duplicate file in batch: {name}")
        checksums.add(checksum)
        staged.append((name, suffix, payload, checksum))
    return staged


def prepare_batch(files, storage_root, batch_id=None, owner_id=None):
    """Persist immutable originals and return a JSON-safe profile manifest."""
    staged = _validate_file_parts(files)
    batch_id = str(uuid.UUID(str(batch_id))) if batch_id else str(uuid.uuid4())
    batch_dir = Path(storage_root) / batch_id
    original_dir = batch_dir / "original"
    if batch_dir.exists():
        raise UploadValidationError("Upload batch already exists")

    try:
        original_dir.mkdir(parents=True, exist_ok=False)
        profile_files = []
        for original_name, suffix, payload, checksum in staged:
            file_id = str(uuid.uuid4())
            stored_name = f"{file_id}{suffix}"
            stored_path = original_dir / stored_name
            stored_path.write_bytes(payload)
            profile = profile_uploaded_file(stored_path, original_name)
            profile_files.append(
                {
                    "id": file_id,
                    "checksum": checksum,
                    "sizeBytes": len(payload),
                    "storedName": stored_name,
                    **profile,
                }
            )

        batch = {
            "id": batch_id,
            "ownerId": str(owner_id) if owner_id else None,
            "version": _batch_version(),
            "state": "ready_for_validation",
            "fileCount": len(profile_files),
            "profiledRowCount": sum(sheet["dataRows"] for item in profile_files for sheet in item["sheets"]),
            "totalBytes": sum(item["sizeBytes"] for item in profile_files),
            "receivedAt": _utc_timestamp(),
        }
        batch["lifecycle"] = _build_lifecycle(profile_files)
        manifest = {"batch": batch, "files": profile_files}
        (batch_dir / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return manifest
    except Exception:
        shutil.rmtree(batch_dir, ignore_errors=True)
        raise


def load_batch_manifest(storage_root, batch_id):
    """Load one persisted batch from a UUID-only path without exposing file bytes."""
    normalized_id = str(uuid.UUID(str(batch_id)))
    manifest_path = Path(storage_root) / normalized_id / "manifest.json"
    if not manifest_path.is_file():
        raise FileNotFoundError(normalized_id)
    return json.loads(manifest_path.read_text(encoding="utf-8"))


def uploaded_sheet_page(storage_root, batch_id, file_id, sheet_name, page, page_size, exceptions_only=False):
    """Return one masked, cleaned page from an uploaded spreadsheet or text table."""
    manifest = load_batch_manifest(storage_root, batch_id)
    file_entry = next((item for item in manifest["files"] if item["id"] == str(file_id)), None)
    if file_entry is None:
        raise FileNotFoundError(str(file_id))
    if file_entry["parserState"] != "profiled":
        raise UploadValidationError("This file format cannot be previewed as rows")

    source_path = Path(storage_root) / manifest["batch"]["id"] / "original" / file_entry["storedName"]
    suffix = Path(file_entry["name"]).suffix.casefold()
    if suffix == ".xlsx":
        workbook = load_workbook(source_path, read_only=True, data_only=False, keep_links=False)
        try:
            if sheet_name not in workbook.sheetnames:
                raise FileNotFoundError(sheet_name)
            cleaned = clean_sheet_rows(workbook[sheet_name].iter_rows(values_only=True), preview_row_limit=None, include_exception_rows=exceptions_only)
        finally:
            workbook.close()
    elif suffix in {".csv", ".txt"}:
        expected_sheet_name = Path(file_entry["name"]).stem
        if sheet_name != expected_sheet_name:
            raise FileNotFoundError(sheet_name)
        with source_path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
            cleaned = clean_sheet_rows(csv.reader(source), preview_row_limit=None, include_exception_rows=exceptions_only)
    else:
        raise UploadValidationError("This file format cannot be previewed as rows")

    all_rows = cleaned["exceptionRows"] if exceptions_only else cleaned["previewRows"]
    total_rows = len(all_rows) if exceptions_only else cleaned["cleaning"]["cleanedRows"]
    offset = (page - 1) * page_size
    return {
        "sheet": sheet_name,
        "columns": cleaned["columns"],
        "rows": all_rows[offset : offset + page_size],
        "page": page,
        "pageSize": page_size,
        "totalRows": total_rows,
        "totalPages": max(1, (total_rows + page_size - 1) // page_size),
        "exceptionsOnly": exceptions_only,
    }


def _cleaned_batch_sheets(storage_root, batch_id, file_id=None):
    manifest = load_batch_manifest(storage_root, batch_id)
    batch_dir = Path(storage_root) / manifest["batch"]["id"] / "original"
    sheets = []
    file_entries = [entry for entry in manifest["files"] if not file_id or entry.get("id") == file_id]
    if file_id and not file_entries:
        raise FileNotFoundError(file_id)
    for file_entry in file_entries:
        if file_entry["parserState"] != "profiled":
            continue
        source_path = batch_dir / file_entry["storedName"]
        suffix = Path(file_entry["name"]).suffix.casefold()
        if suffix == ".xlsx":
            workbook = load_workbook(source_path, read_only=True, data_only=False, keep_links=False)
            try:
                for worksheet in workbook.worksheets:
                    sheets.append({"fileName": file_entry["name"], "sheetName": worksheet.title, **clean_sheet_rows(worksheet.iter_rows(values_only=True), preview_row_limit=None)})
            finally:
                workbook.close()
        elif suffix in {".csv", ".txt"}:
            with source_path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
                sheets.append({"fileName": file_entry["name"], "sheetName": Path(file_entry["name"]).stem, **clean_sheet_rows(csv.reader(source), preview_row_limit=None)})
    if not sheets:
        raise UploadValidationError("This batch has no spreadsheet data available for export")
    return manifest, sheets


def _unique_sheet_name(workbook, preferred_name):
    base = re.sub(r"[\\/*?:\[\]]", "_", preferred_name)[:31] or "Cleaned data"
    candidate = base
    suffix = 2
    while candidate in workbook.sheetnames:
        candidate = f"{base[:28]}_{suffix}"
        suffix += 1
    return candidate


def save_cleaned_batch(storage_root, batch_id):
    """Persist a clean, masked Excel export for the uploaded batch."""
    manifest, sheets = _cleaned_batch_sheets(storage_root, batch_id)
    cleaned_dir = Path(storage_root) / manifest["batch"]["id"] / "cleaned"
    cleaned_dir.mkdir(exist_ok=True)
    filename = f"cleaned_{manifest['batch']['version']}.xlsx"
    output_path = cleaned_dir / filename
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets:
        worksheet = workbook.create_sheet(_unique_sheet_name(workbook, sheet["sheetName"]))
        worksheet.append(sheet["columns"])
        for row in sheet["previewRows"]:
            worksheet.append([row.get(column, "") for column in sheet["columns"]])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
    workbook.save(output_path)
    workbook.close()
    manifest["batch"]["cleanedWorkbook"] = filename
    manifest["batch"]["cleanedSavedAt"] = _utc_timestamp()
    unresolved_exceptions = any(
        stage.get("key") == "map" and stage.get("state") == "needs_review"
        for stage in manifest["batch"].get("lifecycle", [])
    )
    for stage in manifest["batch"].get("lifecycle", []):
        if stage.get("key") == "review":
            stage["state"] = "needs_review" if unresolved_exceptions else "complete"
            stage["detail"] = "Resolve mapping exceptions before approval" if unresolved_exceptions else "Cleaned data approved and saved"
        elif stage.get("key") == "ready_for_mia":
            stage["state"] = "pending" if unresolved_exceptions else "complete"
            stage["detail"] = "Awaiting exception resolution and approval" if unresolved_exceptions else "Approved dataset available to Mia"
    (Path(storage_root) / manifest["batch"]["id"] / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest, output_path


def export_cleaned_batch(storage_root, batch_id, export_format, file_id=None):
    """Create an exportable clean-data artifact and return its bytes and download metadata."""
    export_format = export_format.casefold()
    manifest, sheets = _cleaned_batch_sheets(storage_root, batch_id, file_id=file_id)
    version = manifest["batch"]["version"]
    if export_format == "xlsx":
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet in sheets:
            worksheet = workbook.create_sheet(_unique_sheet_name(workbook, sheet["sheetName"]))
            worksheet.append(sheet["columns"])
            for row in sheet["previewRows"]:
                worksheet.append([row.get(column, "") for column in sheet["columns"]])
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        source_name = Path(sheets[0]["fileName"]).stem if sheets else "upload"
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"standardised_{source_name}_{version}.xlsx"
    if export_format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        for sheet in sheets:
            writer.writerow([sheet["fileName"], sheet["sheetName"]])
            writer.writerow(sheet["columns"])
            writer.writerows([[row.get(column, "") for column in sheet["columns"]] for row in sheet["previewRows"]])
            writer.writerow([])
        source_name = Path(sheets[0]["fileName"]).stem if sheets else "upload"
        return output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", f"standardised_{source_name}_{version}.csv"
    if export_format == "pdf":
        lines = [f"Cleaned data export {version}"]
        for sheet in sheets:
            lines.append(f"{sheet['fileName']} / {sheet['sheetName']} ({sheet['cleaning']['cleanedRows']} rows)")
            lines.append(" | ".join(sheet["columns"]))
            lines.extend(" | ".join(str(row.get(column, "")) for column in sheet["columns"]) for row in sheet["previewRows"])
        content = "\n".join(lines).encode("latin-1", "replace")
        return _simple_pdf(content), "application/pdf", f"cleaned_{version}.pdf"
    raise UploadValidationError("Choose Excel, CSV, or PDF for export")


def _simple_pdf(content):
    lines = content.decode("latin-1").splitlines()
    pages = [lines[index : index + 48] for index in range(0, len(lines), 48)] or [["Cleaned data export"]]
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", None, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    page_ids = []
    for page_lines in pages:
        text = [b"BT /F1 8 Tf 42 800 Td 10 TL"]
        for line in page_lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("latin-1", "replace")
            text.append(b"(" + escaped + b") Tj T*")
        text.append(b"ET")
        content_id = len(objects) + 1
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(b"\n".join(text)), b"\n".join(text)))
        page_id = len(objects) + 1
        page_ids.append(page_id)
        objects.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_id)
    objects[1] = b"<< /Type /Pages /Count %d /Kids [%s] >>" % (len(page_ids), b" ".join(f"{page_id} 0 R".encode() for page_id in page_ids))
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + value + b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    output.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode())
    return bytes(output)
