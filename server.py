#!/usr/bin/env python3
import json
import hashlib
import math
import os
import re
import secrets
import ssl
import time
import urllib.error
import urllib.request
import uuid
import zlib
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from itertools import islice
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook

from backend import hybrid_retrieval, rag_embeddings
from backend.agent_monitoring_demo import DEMO_BATCH, DEMO_EVENT_COUNT, build_demo_events, summarize_route_monitoring
from backend.identity import hash_password, normalize_username, validate_registration_username, verify_password
from backend.mia_graph.conversation import build_conversation_graph
from backend.mia_graph.checkpoint import checkpoint_session
from backend.mia_graph.dataset import run_dataset_graph
from backend.mia_graph.triggers import MiaTriggerScheduler
from backend.mia_graph.evaluation import (
    aggregate_human_calibration,
    build_evaluation_graph,
    build_evaluation_run,
    parse_judge_verdict,
    validate_human_calibration_record,
)
from backend.mia_graph.observability import RedactedTracer, configured_langfuse_client, graph_audit_record, langfuse_batch_deletion_sink, record_graph_result_trace
from backend.mia_graph.policy import assess_untrusted_text
from backend.review_workflow import review_item_payload_for_graph, transition_for_action
from backend.deletion_propagation import BatchDeletedError, assert_batch_active, propagate_batch_deletion, register_lineage
from backend.production_governance import evaluate_readiness, load_governance, pseudonymous_actor_id, readiness_report
from backend.upload_http import parse_multipart_files
from backend.upload_ingest import MAX_BATCH_BYTES, UploadValidationError, _cleaned_batch_sheets, export_cleaned_batch, load_batch_manifest, prepare_batch, save_cleaned_batch, uploaded_sheet_page
from work.dataset_run import DatasetRunError, build_analysis_snapshot, event_envelope, publish_run
from work.dataset_run_outbox import claim as claim_outbox, complete as complete_outbox, enqueue as enqueue_outbox


try:
    import certifi
    import psycopg
    from pgvector import Vector
    from pgvector.psycopg import register_vector
    from psycopg.rows import dict_row
except ImportError:
    certifi = None
    psycopg = None
    Vector = None
    register_vector = None
    dict_row = None


ROOT = Path(__file__).resolve().parent
TOKEN_PATTERN = re.compile(r"[a-z0-9]+")


def json_response_body(payload):
    """Encode API responses containing PostgreSQL datetime values safely."""
    return json.dumps(
        payload,
        ensure_ascii=False,
        default=lambda value: value.isoformat() if hasattr(value, "isoformat") else str(value),
    ).encode("utf-8")


def tokenize(value: str) -> list[str]:
    """Return lowercase alphanumeric tokens for transparent local search."""
    return TOKEN_PATTERN.findall(value.lower())


def score_document(query: str, document: dict) -> int:
    """Score each distinct query token against searchable audit fields."""
    query_tokens = set(tokenize(query))
    searchable = " ".join(
        [document.get("id", ""), document.get("text", "")]
        + list(document.get("metadata", {}).values())
        + list(document.get("raw_answers", {}).values())
    )
    document_tokens = set(tokenize(searchable))
    return sum(token in document_tokens for token in query_tokens)
APP_DIR = ROOT / "app"
SCHEMA_PATH = ROOT / "db" / "schema.sql"
SEED_PATH = ROOT / "data" / "seed.json"
SMALLTALK_PATH = ROOT / "data" / "smalltalk.json"
DEEPSEEK_URL = "https://api.deepseek.com/chat/completions"
LOGIN_APP_ORIGINS = {
    "http://127.0.0.1:11337",
    "http://localhost:11337",
    "http://127.0.0.1:11338",
    "http://localhost:11338",
}
VECTOR_DIMS = 16
HTTPS_CONTEXT = ssl.create_default_context(cafile=certifi.where()) if certifi is not None else None
SOURCE_WORKBOOK_DIR = Path("/Users/irene/Documents/Codex/2026-07-22/dfd/outputs/it_data_flow_demo/01_upload/source_workbooks")
UPLOAD_STORAGE_DIR = ROOT / "data" / "upload_batches"
DATASET_RUN_WORKER = ThreadPoolExecutor(max_workers=1, thread_name_prefix="dataset-run")
MIA_GRAPH_CHECKPOINTER = None
MIA_TRIGGER_SCHEDULER = None
PRODUCTION_MODE_ENABLED = os.environ.get("MIA_PRODUCTION_MODE", "0").lower() in {"1", "true", "yes"}
GOVERNANCE_AUDIT_EVENTS = []
SOURCE_WORKBOOKS = {
    "booking_payment_export_v4.xlsx": ["Bookings", "Payments", "Service cases"],
    "crm_loyalty_snapshot.xlsx": ["Customers", "Consent", "Identity aliases"],
    "voice_and_operations_batch.xlsx": ["Trip legs", "App events", "Survey responses", "Reviews"],
}
RESTRICTED_SOURCE_FIELD = re.compile(r"email|customer_ref|crm_contact|loyalty_no|account_hash|device_hash|alias_value", re.I)

PROJECT_MEMORY_SEEDS = [
    {
        "memory_key": "dfds-rag-answering-policy-2026-07-20",
        "memory_layer": "long_term",
        "memory_type": "answering_policy",
        "title": "DFDS report assistant answering policy",
        "memory_text": (
            "This DFDS Customer Intelligence Platform focuses on passenger ferry customers, "
            "not freight or logistics. The floating report assistant must answer with the "
            "current frontend page context, the PostgreSQL Evidence Knowledge Base, historical "
            "business insights, and project memory together. Answers must be business-friendly, "
            "and must be written only in the user's input language. Do not add an English-first "
            "section or translation-style summary unless the manager explicitly asks for a "
            "translation. The assistant must show evidence IDs, source names, "
            "ratings or review counts when available, and avoid inventing facts not supported "
            "by the dashboard or database."
        ),
        "source": "user_requirement",
        "priority": 5,
        "evidence_refs": [],
        "metadata": {
            "memory_model": {
                "short_term": "recent chat turns in chat_messages",
                "medium_term": "active project preferences, UI rules, evaluation feedback, and batch state",
                "long_term": "durable product scope, answering policy, reusable business insight rules, and platform architecture",
            },
            "inspired_by": "Claude Code style project memory plus reusable knowledge notes",
        },
    },
    {
        "memory_key": "dfds-passenger-profile-synthetic-pg-2026-07-23",
        "memory_layer": "long_term",
        "memory_type": "data_source_policy",
        "title": "Synthetic Passenger Profile PG usage policy",
        "memory_text": (
            "The Passenger Profile sub-agent uses a synthetic PostgreSQL star-schema dataset "
            "with 1,000 passenger profiles and 1,800 trip records. It supports marketing "
            "segmentation, route affinity, travel context, product preference, and ACTAR "
            "recommendations. Mia may use it as a prototyping and hypothesis-generation source, "
            "but must label it as synthetic and must not present it as real DFDS CRM, booking, "
            "or survey evidence."
        ),
        "source": "user_requirement",
        "priority": 5,
        "evidence_refs": [],
        "metadata": {
            "source_key": "synthetic-passenger-profile-pg",
            "sub_agent": "Passenger Profile",
            "actar": "Action area, Customer target, Trigger signal, Analysis, Recommendation",
            "synthetic": True,
        },
    },
    {
        "memory_key": "dfds-platform-memory-architecture-2026-07-28",
        "memory_layer": "long_term",
        "memory_type": "platform_architecture",
        "title": "DFDS platform memory architecture",
        "memory_text": (
            "The shared platform should use short-term memory for the current chat turn and page state, "
            "medium-term memory for active upload batches, passenger-profile slices, and feedback queues, "
            "and long-term memory for durable product rules, validated insights, and reusable answer policies. "
            "Memory is shared across Overview, Passenger Profile, Survey CSV, Data Basis, IT Data Flow, Review Console, "
            "and Mia so the product behaves like one connected intelligence layer."
        ),
        "source": "user_requirement",
        "priority": 5,
        "evidence_refs": [],
        "metadata": {
            "memory_model": {
                "short_term": "current chat turns, current page state, and temporary UI state",
                "medium_term": "active upload batches, passenger-profile slices, and feedback queues",
                "long_term": "durable product rules, validated insights, and reusable platform conventions",
            },
            "surfaces": [
                "Overview",
                "Passenger Profile",
                "Survey CSV",
                "Data Basis",
                "IT Data Flow",
                "Review Console",
                "Mia",
            ],
        },
    },
    {
        "memory_key": "dfds-it-data-flow-architecture-2026-07-28",
        "memory_layer": "long_term",
        "memory_type": "process_documentation",
        "title": "DFDS IT data flow architecture",
        "memory_text": (
            "Uploaded txt, pdf, doc, and docx files should move through a clear flow: upload, validate, clean, model, "
            "serve, and collect feedback. Source versions, reliability checks, and visible Excel fields should stay "
            "explicit, while inferred dimensions remain model-side only. The IT flow page exists so the team can inspect "
            "the technical path behind the business dashboard."
        ),
        "source": "user_requirement",
        "priority": 4,
        "evidence_refs": [],
        "metadata": {
            "source_systems": [
                "Upload inbox",
                "Version monitor",
                "Cleaning / normalization",
                "Star-schema model",
                "Dashboard views",
                "Mia chat context",
            ],
            "governance": [
                "keep original filenames",
                "track source versions",
                "surface visible Excel fields",
                "use human feedback before reuse",
            ],
        },
    },
]

VALID_REVIEW_STATUSES = {"pending", "pending_human_review", "approved", "needs_changes", "corrected", "rejected", "withdrawn"}

REVIEW_CONSOLE_SEED_ITEMS = [
    {
        "id": "seed-20260716-route-filter-scope",
        "review_key": "seed-20260716-route-filter-scope",
        "title": "Route filter only appears where route focus changes content",
        "layer": "scope",
        "status": "approved",
        "severity": "medium",
        "source": "dashboard generation",
        "source_key": "dashboard-generation",
        "route_key": "all",
        "reason": "July 16 logs required route focus to be hidden on App Reviews, Competitors, Survey CSV, Update Log, and Data Basis.",
        "recommendation": "Keep route controls limited to Overview, Customer Voice, and Recommendations.",
        "publish_state": "verified",
        "evidence_chain": ["input context", "UI route behavior", "supervisor judgment", "visibility decision"],
        "artifacts": ["mias-cruises-customer-insights-logs/2026-07-16.md"],
        "metadata": {"seeded_from_log": True, "log_date": "2026-07-16"},
    },
    {
        "id": "seed-20260717-source-limitations",
        "review_key": "seed-20260717-source-limitations",
        "title": "Public source limitations are labeled honestly",
        "layer": "collection",
        "status": "needs_changes",
        "severity": "high",
        "source": "evidence import",
        "source_key": "evidence-import",
        "route_key": "all",
        "reason": "July 17 logs say Google Reviews are location-level signals, Reddit is discussion signal, and Firecrawl search-result snapshots are directional evidence.",
        "recommendation": "Show collection method and limitation labels before using source evidence in recommendations.",
        "publish_state": "internal_only",
        "evidence_chain": ["collection attempt", "source limitation", "normalized source row", "supervisor judgment"],
        "artifacts": ["mias-cruises-customer-insights-logs/2026-07-17.md"],
        "metadata": {"seeded_from_log": True, "log_date": "2026-07-17"},
    },
    {
        "id": "seed-20260720-rag-context",
        "review_key": "seed-20260720-rag-context",
        "title": "RAG answers use database, page context, memory, and fallback evidence",
        "layer": "retrieval",
        "status": "pending",
        "severity": "high",
        "source": "RAG retrieval",
        "source_key": "rag-retrieval",
        "route_key": "all",
        "reason": "July 20 logs introduced PostgreSQL + pgvector, project memories, recent chat history, and frontend fallback evidence.",
        "recommendation": "Flag answers that cite irrelevant evidence or miss obvious visible dashboard evidence.",
        "publish_state": "internal_only",
        "evidence_chain": ["page context", "evidence_items", "project_memories", "chat history", "supervisor judgment"],
        "artifacts": ["outputs/dfds-rag-validation-set.xlsx", "mias-cruises-customer-insights-logs/2026-07-20.md"],
        "metadata": {"seeded_from_log": True, "log_date": "2026-07-20"},
    },
    {
        "id": "seed-20260721-mia-weather-language",
        "review_key": "seed-20260721-mia-weather-language",
        "title": "Mia weather and multilingual guardrails do not leak evidence cards",
        "layer": "language",
        "status": "pending",
        "severity": "high",
        "source": "chat answer",
        "source_key": "chat-answer",
        "route_key": "all",
        "reason": "July 21 logs require weather, small talk, and off-topic prompts to skip retrieval; non-English evidence summaries must avoid long English evidence bodies.",
        "recommendation": "Test weather, greeting, Chinese Dover-Calais, and Danish prompts before publishing chat changes.",
        "publish_state": "internal_only",
        "evidence_chain": ["input message", "intent guard", "route detection", "language formatting", "supervisor judgment"],
        "artifacts": ["mias-cruises-customer-insights-logs/2026-07-21.md"],
        "metadata": {"seeded_from_log": True, "log_date": "2026-07-21"},
    },
    {
        "id": "seed-20260718-21-cache-release",
        "review_key": "seed-20260718-21-cache-release",
        "title": "Frontend cache-busting and Update Log sequencing are checked",
        "layer": "release",
        "status": "needs_changes",
        "severity": "medium",
        "source": "release hygiene",
        "source_key": "release-hygiene",
        "route_key": "all",
        "reason": "July 18 through July 21 logs repeatedly show browser cache caused old dashboard or Mia logic to stay visible.",
        "recommendation": "Require app/index.html, app/app.js, app/data/update-log.mjs, and the daily log to move together for frontend releases.",
        "publish_state": "internal_only",
        "evidence_chain": ["module change", "cache version", "Update Log entry", "daily log", "supervisor judgment"],
        "artifacts": ["mias-cruises-customer-insights-logs/2026-07-18.md", "mias-cruises-customer-insights-logs/2026-07-21.md"],
        "metadata": {"seeded_from_log": True, "log_date": "2026-07-21"},
    },
]


def load_local_env():
    env_path = ROOT / ".env"
    if not env_path.exists():
        return

    for line in env_path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        key, value = stripped.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


def load_smalltalk_playbook(path=SMALLTALK_PATH):
    path = Path(path)
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def build_smalltalk_records(playbook=None):
    playbook = playbook or load_smalltalk_playbook()
    if not playbook:
        return []

    collected_at = playbook.get("collectedAt")
    route_key = playbook.get("routeKey", "smalltalk")
    source_key = playbook.get("sourceKey", "mia-smalltalk-playbook")
    source_name = playbook.get("sourceName", "Mia Small Talk Playbook")
    source_type = playbook.get("sourceType", "conversation_playbook")
    source_status = playbook.get("sourceStatus", "Curated")
    source_notes = playbook.get("sourceNotes", "Curated small-talk theme cards for bridge replies.")
    themes = playbook.get("themes", [])
    chunk_count = max(len(themes), 1)

    records = []
    for index, theme in enumerate(themes, start=1):
        title = str(theme.get("title", "")).strip()
        intent = str(theme.get("intent", theme.get("key", ""))).strip()
        examples = [str(example).strip() for example in theme.get("examples", []) if str(example).strip()]
        reply_summary = str(theme.get("replySummary", "")).strip()
        bridge_back = str(theme.get("bridgeBack", "")).strip()
        guardrail = str(theme.get("guardrail", "")).strip()
        if not title or not reply_summary:
            continue

        original_content = (
            f"Smalltalk theme card: {title}. "
            f"Common prompts: {'; '.join(examples)}. "
            f"Reply summary: {reply_summary}. "
            f"Bridge back: {bridge_back}. "
            f"Guardrail: {guardrail}."
        ).strip()
        metadata = {
            "subroute": "smalltalk",
            "theme_key": theme.get("key", intent),
            "intent": intent,
            "examples": examples,
            "reply_summary": reply_summary,
            "bridge_back": bridge_back,
            "guardrail": guardrail,
            "chunk_strategy": "theme_card",
            "chunk_index": index,
            "chunk_count": chunk_count,
            "playbook_source": source_key,
            "source_status": source_status,
        }
        records.append(
            {
                "source_key": source_key,
                "source_name": source_name,
                "source_type": source_type,
                "source_status": source_status,
                "source_notes": source_notes,
                "company_name": "DFDS",
                "route_key": route_key,
                "evidence_kind": "smalltalk_theme_card",
                "original_content": original_content,
                "translated_content": reply_summary,
                "rating": None,
                "review_count": None,
                "published_at": collected_at,
                "url": "",
                "keywords": derive_keywords(
                    "smalltalk_theme_card",
                    original_content,
                    reply_summary,
                    metadata,
                    route=route_key,
                    source=source_key,
                    extra=[title, intent, "smalltalk"],
                    limit=24,
                ),
                "metadata": metadata,
            }
        )

    return records


def slugify(value):
    slug = re.sub(r"[^a-z0-9]+", "-", str(value).lower()).strip("-")
    return slug or "unknown"


def parse_rating(value):
    match = re.search(r"(\d+(?:\.\d+)?)", str(value or ""))
    return float(match.group(1)) if match else None


def parse_review_count(value):
    text = str(value or "").replace(",", "")
    patterns = [
        r"from\s+(\d+)\s+(?:google\s+)?reviews?",
        r"(\d+)\s+(?:google\s+)?reviews?",
        r"(\d+)\s+ratings?",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            return int(match.group(1))
    match = re.search(r"(\d+)", text)
    return int(match.group(1)) if match else None


def route_key(value):
    value = str(value or "all").lower()
    if "dover" in value and "calais" in value:
        return "dover-calais"
    if "newhaven" in value and "dieppe" in value:
        return "newhaven-dieppe"
    if "newcastle" in value or "ijmuiden" in value:
        return "newcastle-ijmuiden"
    if "jersey" in value or "channel islands" in value:
        return "jersey"
    if value in {"all", ""}:
        return "all"
    return slugify(value)


KEYWORD_STOPWORDS = {
    "the", "and", "for", "are", "was", "were", "with", "from", "this", "that",
    "into", "does", "have", "show", "shows", "page", "view", "all", "can",
    "you", "tell", "about", "around", "because", "passenger", "passengers",
    "dfds", "ferry", "ferries", "public", "review", "reviews", "rating",
    "ratings", "source", "signal", "signals", "includes", "included",
}

KEYWORD_PHRASES = [
    (r"\bmobile\s+check[\s-]?in\b", "mobile-check-in"),
    (r"\bcheck[\s-]?in\b", "check-in"),
    (r"\blog[\s-]?in\b|\blogin\b", "login"),
    (r"\bbooking\b|\bbookings\b", "booking"),
    (r"\bborder\s+control\b|\bpassport\s+control\b", "border-control"),
    (r"\bdelay\b|\bdelays\b|\bdelayed\b", "delay"),
    (r"\bdisruption\b|\bdisrupted\b", "disruption"),
    (r"\bcustomer\s+service\b", "customer-service"),
    (r"\bonboard\b|\bon-board\b", "onboard"),
    (r"\bcabin\b|\bcabins\b", "cabins"),
    (r"\bfood\b|\bcatering\b", "food"),
    (r"\bstaff\b|\bcrew\b", "staff"),
    (r"\bterminal\b|\bboarding\b", "terminal-boarding"),
    (r"\bcompensation\b|\brefund\b|\brefunds\b", "compensation"),
    (r"\bticket\b|\btickets\b|\bfare\b|\bfares\b", "ticketing"),
]


def normalize_keyword(value):
    text = str(value or "").lower().replace("_", "-")
    text = re.sub(r"[^a-z0-9\u4e00-\u9fff-]+", "-", text).strip("-")
    text = re.sub(r"-{2,}", "-", text)
    if not text or text in KEYWORD_STOPWORDS or len(text) < 2:
        return ""
    return text


def derive_keywords(kind, original, translated=None, metadata=None, route="", source="", extra=None, limit=18):
    metadata = metadata or {}
    parts = [
        kind or "",
        original or "",
        translated or "",
        json.dumps(metadata, ensure_ascii=False),
        route or "",
        source or "",
        " ".join(extra or []),
    ]
    text = " ".join(parts).lower().replace("_", "-")

    keywords = []

    def add(value):
        keyword = normalize_keyword(value)
        if keyword and keyword not in keywords:
            keywords.append(keyword)

    add(route)
    add(source)
    add(kind)
    for key in ("theme", "title", "sentiment", "journey_stage", "source_label", "platform"):
        add(metadata.get(key))

    for pattern, keyword in KEYWORD_PHRASES:
        if re.search(pattern, text, re.IGNORECASE):
            add(keyword)

    for token in re.findall(r"[a-z0-9\u4e00-\u9fff][a-z0-9\u4e00-\u9fff-]*", text):
        add(token)
        if len(keywords) >= limit:
            break

    return keywords[:limit]


ROUTE_KEYWORDS = {"dover-calais", "newhaven-dieppe", "newcastle-ijmuiden", "jersey"}
SENTIMENT_KEYWORDS = {"positive", "negative", "mixed", "neutral", "mixed-positive"}
SOURCE_KEYWORD_CATEGORIES = {
    "apple-app-store": "app_store",
    "google-play": "app_store",
    "google-reviews": "location_review",
    "trustpilot": "review_platform",
    "reddit": "discussion",
    "company-news": "news",
    "dfds-official": "operator_public_page",
    "jersey-government": "public_authority",
}
TOPIC_KEYWORD_CATEGORIES = {
    "login": "digital_experience",
    "booking": "digital_experience",
    "check-in": "digital_experience",
    "mobile-check-in": "digital_experience",
    "digital-pre-trip": "digital_experience",
    "delay": "disruption",
    "disruption": "disruption",
    "border-control": "disruption",
    "compensation": "disruption",
    "ticketing": "commercial",
    "customer-service": "service",
    "staff": "service",
    "terminal-boarding": "terminal",
    "onboard": "onboard",
    "cabins": "onboard",
    "food": "onboard",
}


def keyword_display_label(keyword_key):
    special_labels = {
        "dfds": "DFDS",
        "ios": "iOS",
        "bbc": "BBC",
        "dover-calais": "Dover-Calais",
        "newhaven-dieppe": "Newhaven-Dieppe",
        "newcastle-ijmuiden": "Newcastle-IJmuiden",
        "apple-app-store": "Apple App Store",
        "google-play": "Google Play",
        "google-reviews": "Google Reviews",
        "company-news": "Company News",
        "dfds-official": "DFDS Official",
        "jersey-government": "Jersey Government",
        "mobile-check-in": "Mobile check-in",
        "check-in": "Check-in",
        "border-control": "Border control",
        "customer-service": "Customer service",
        "digital-pre-trip": "Digital pre-trip",
        "terminal-boarding": "Terminal boarding",
    }
    if keyword_key in special_labels:
        return special_labels[keyword_key]
    return keyword_key.replace("-", " ").capitalize()


def keyword_profile(keyword):
    keyword_key = normalize_keyword(keyword)
    if not keyword_key:
        raise ValueError("keyword cannot be empty")

    keyword_type = "topic"
    keyword_category = TOPIC_KEYWORD_CATEGORIES.get(keyword_key, "general")
    language = "en"

    if keyword_key in ROUTE_KEYWORDS:
        keyword_type = "route"
        keyword_category = "route_focus"
    elif keyword_key in SOURCE_KEYWORD_CATEGORIES:
        keyword_type = "source"
        keyword_category = SOURCE_KEYWORD_CATEGORIES[keyword_key]
    elif keyword_key in SENTIMENT_KEYWORDS:
        keyword_type = "sentiment"
        keyword_category = "sentiment"
    elif re.search(r"[\u4e00-\u9fff]", keyword_key):
        keyword_type = "topic"
        keyword_category = "localized"
        language = "zh"

    return {
        "keyword_key": keyword_key,
        "display_label": keyword_display_label(keyword_key),
        "keyword_type": keyword_type,
        "keyword_category": keyword_category,
        "language": language,
        "description": f"Derived DFDS evidence keyword for {keyword_category}.",
    }


def keyword_weight(profile, position):
    if profile["keyword_type"] in {"route", "source"}:
        return 2.0
    if profile["keyword_type"] == "sentiment":
        return 1.5
    if position <= 6:
        return 1.25
    return 1.0


def build_evidence_keyword_links(evidence_id, keywords, metadata=None):
    metadata = metadata or {}
    links = []
    seen = set()
    for position, keyword in enumerate(keywords or [], start=1):
        profile = keyword_profile(keyword)
        key = profile["keyword_key"]
        if key in seen:
            continue
        seen.add(key)
        links.append(
            {
                **profile,
                "evidence_id": evidence_id,
                "position": position,
                "weight": keyword_weight(profile, position),
                "source": metadata.get("source", "derived"),
                "metadata": metadata,
            }
        )
    return links


def canonical_source_key(value):
    label = str(value or "").lower()
    if "google play" in label:
        return "google-play"
    if "apple" in label or "ios" in label:
        return "apple-app-store"
    if "google review" in label or "wanderlog" in label:
        return "google-reviews"
    if "reddit" in label:
        return "reddit"
    if "bbc" in label or "guardian" in label or "news" in label or "local public" in label:
        return "company-news"
    if "survey" in label:
        return "survey-csv"
    if "trustpilot" in label:
        return "trustpilot"
    return slugify(value)


def vectorize_text(text):
    vector = [0.0] * VECTOR_DIMS
    tokens = re.findall(r"[a-z0-9\u4e00-\u9fff]+", str(text).lower())
    for token in tokens:
        index = zlib.crc32(token.encode("utf-8")) % VECTOR_DIMS
        weight = 1.0 + min(len(token), 12) / 12
        vector[index] += weight
    norm = math.sqrt(sum(value * value for value in vector)) or 1.0
    normalized = [round(value / norm, 6) for value in vector]
    return Vector(normalized) if Vector else normalized


def query_terms(*parts):
    text = " ".join(str(part or "") for part in parts).lower()
    tokens = re.findall(r"[a-z0-9\u4e00-\u9fff]+", text)
    stopwords = {
        "the", "and", "for", "are", "what", "why", "how", "main", "about",
        "with", "from", "this", "that", "into", "does", "have", "show",
        "page", "view", "all", "can", "you", "tell", "use", "public",
        "source", "record", "records", "dataset", "status", "mention",
        "evidence", "synthetic",
    }
    terms = []

    def add(value):
        normalized = normalize_keyword(value)
        if normalized and normalized not in terms:
            terms.append(normalized)

    for pattern, keyword in KEYWORD_PHRASES:
        if re.search(pattern, text, re.IGNORECASE):
            add(keyword)

    for route_key in ["dover-calais", "newhaven-dieppe", "newcastle-ijmuiden", "jersey"]:
        compact_route = route_key.replace("-", " ")
        if route_key in text.replace(" ", "-") or compact_route in text:
            add(route_key)
            for part in route_key.split("-"):
                add(part)

    for token in tokens:
        if len(token) > 2 and token not in stopwords:
            add(token)

    synonym_map = {
        "app": ["app", "mobile", "login", "booking", "check", "android", "apple", "ios", "google", "play"],
        "reviews": ["review", "reviews", "rating", "ratings", "trustpilot", "google", "apple"],
        "competitor": ["competitor", "competitors", "benchmark", "comparison", "overlap", "route"],
        "recommendation": ["recommendation", "recommend", "priority", "action", "marketing"],
        "root": ["root", "cause", "reason", "risk", "issue", "problem"],
    }
    for key, synonyms in synonym_map.items():
        if key in text or any(word in text for word in synonyms):
            for synonym in synonyms:
                add(synonym)

    if any(word in text for word in ["应用", "登录", "评分", "评论", "苹果", "安卓"]):
        for synonym in synonym_map["app"] + synonym_map["reviews"]:
            add(synonym)
    if any(word in text for word in ["竞品", "竞争", "对比"]):
        for synonym in synonym_map["competitor"]:
            add(synonym)
    if any(word in text for word in ["推荐", "建议", "下一步"]):
        for synonym in synonym_map["recommendation"]:
            add(synonym)
    if any(word in text for word in ["根因", "原因", "风险", "问题"]):
        for synonym in synonym_map["root"]:
            add(synonym)

    return terms[:18]


def json_compact(value, max_chars=3000):
    text = json.dumps(value or {}, ensure_ascii=False, indent=2)
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n... [truncated]"


def synthetic_archetype_key(item):
    record_id = str((item.get("metadata") or {}).get("synthetic_record_id") or "")
    match = re.match(r"(.+)-v\d+$", record_id)
    if match:
        return match.group(1)
    if record_id:
        return record_id
    return ""


def dedupe_evidence_for_context(items, limit=6):
    selected = []
    seen = set()
    for item in items:
        archetype = synthetic_archetype_key(item)
        key = f"synthetic:{archetype}" if archetype else f"evidence:{item.get('id')}"
        if key in seen:
            continue
        seen.add(key)
        selected.append(item)
        if len(selected) >= limit:
            break
    return selected


def evidence_candidate_limit(limit):
    return max(limit * 40, 1000)


def connect_db(register=True):
    database_url = os.environ.get("DATABASE_URL")
    if not database_url or psycopg is None:
        return None
    conn = psycopg.connect(database_url, row_factory=dict_row)
    if register and register_vector is not None:
        register_vector(conn)
    return conn


def execute_schema(conn):
    schema = SCHEMA_PATH.read_text(encoding="utf-8")
    with conn.cursor() as cur:
        cur.execute(schema)
    conn.commit()


def fetch_id(cur, table, key_column, key_value):
    cur.execute(f"SELECT id FROM {table} WHERE {key_column} = %s", (key_value,))
    row = cur.fetchone()
    return row["id"] if row else None


def insert_company(cur, name, is_baseline=False):
    slug = slugify(name)
    cur.execute(
        """
        INSERT INTO companies (name, slug, is_baseline)
        VALUES (%s, %s, %s)
        ON CONFLICT (slug) DO UPDATE SET name = EXCLUDED.name
        RETURNING id
        """,
        (name, slug, is_baseline),
    )
    return cur.fetchone()["id"]


def insert_route(cur, key, display_name, notes="", region=""):
    cur.execute(
        """
        INSERT INTO routes (route_key, display_name, notes, region)
        VALUES (%s, %s, %s, %s)
        ON CONFLICT (route_key) DO UPDATE
        SET display_name = EXCLUDED.display_name,
            notes = COALESCE(NULLIF(EXCLUDED.notes, ''), routes.notes)
        RETURNING id
        """,
        (key, display_name, notes, region),
    )
    return cur.fetchone()["id"]


def insert_source(cur, key, name, source_type, status, logo_url="", source_url="", notes="", last_checked_at=None):
    cur.execute(
        """
        INSERT INTO sources (source_key, name, source_type, status, logo_url, source_url, notes, last_checked_at)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (source_key) DO UPDATE
        SET name = EXCLUDED.name,
            status = EXCLUDED.status,
            logo_url = COALESCE(NULLIF(EXCLUDED.logo_url, ''), sources.logo_url),
            source_url = COALESCE(NULLIF(EXCLUDED.source_url, ''), sources.source_url),
            notes = COALESCE(NULLIF(EXCLUDED.notes, ''), sources.notes)
        RETURNING id
        """,
        (key, name, source_type, status, logo_url, source_url, notes, last_checked_at),
    )
    return cur.fetchone()["id"]


def upsert_keyword(cur, profile):
    cur.execute(
        """
        INSERT INTO keywords (
          keyword_key, display_label, keyword_type, keyword_category,
          language, description, updated_at
        )
        VALUES (%s, %s, %s, %s, %s, %s, NOW())
        ON CONFLICT (keyword_key) DO UPDATE
        SET display_label = EXCLUDED.display_label,
            keyword_type = EXCLUDED.keyword_type,
            keyword_category = EXCLUDED.keyword_category,
            language = EXCLUDED.language,
            description = EXCLUDED.description,
            updated_at = NOW()
        RETURNING id
        """,
        (
            profile["keyword_key"],
            profile["display_label"],
            profile["keyword_type"],
            profile["keyword_category"],
            profile["language"],
            profile["description"],
        ),
    )
    return cur.fetchone()["id"]


def sync_evidence_keywords(cur, evidence_id, keywords, metadata=None):
    links = build_evidence_keyword_links(evidence_id, keywords, metadata or {})
    for link in links:
        keyword_id = upsert_keyword(cur, link)
        cur.execute(
            """
            INSERT INTO evidence_keywords (
              evidence_id, keyword_id, position, weight, source, metadata
            )
            VALUES (%s, %s, %s, %s, %s, %s::jsonb)
            ON CONFLICT (evidence_id, keyword_id) DO UPDATE
            SET position = EXCLUDED.position,
                weight = EXCLUDED.weight,
                source = EXCLUDED.source,
                metadata = EXCLUDED.metadata
            """,
            (
                evidence_id,
                keyword_id,
                link["position"],
                link["weight"],
                link["source"],
                json.dumps(link["metadata"], ensure_ascii=False),
            ),
        )
    return len(links)


def replace_evidence_keywords(cur, evidence_id, keywords, metadata=None):
    cur.execute("DELETE FROM evidence_keywords WHERE evidence_id = %s", (evidence_id,))
    return sync_evidence_keywords(cur, evidence_id, keywords, metadata or {})


def insert_evidence(cur, source_id, company_id, route_id, kind, original, translated=None, rating=None, review_count=None, published_at=None, url="", metadata=None, keywords=None):
    metadata = metadata or {}
    if keywords is None:
        keywords = derive_keywords(kind, original, translated, metadata)
    text_for_vector = " ".join(
        [
            kind,
            original or "",
            translated or "",
            json.dumps(metadata, ensure_ascii=False),
            " ".join(keywords),
        ]
    )
    cur.execute(
        """
        INSERT INTO evidence_items (
          source_id, company_id, route_id, evidence_kind, original_content, translated_content,
          rating, review_count, published_at, url, metadata, keywords, embedding
        )
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb, %s::text[], %s)
        ON CONFLICT (source_id, original_content) DO NOTHING
        RETURNING id
        """,
        (
            source_id,
            company_id,
            route_id,
            kind,
            original,
            translated,
            rating,
            review_count,
            published_at,
            url,
            json.dumps(metadata, ensure_ascii=False),
            keywords,
            vectorize_text(text_for_vector),
        ),
    )
    row = cur.fetchone()
    inserted_count = 1 if row else 0
    if row:
        evidence_id = row["id"]
    else:
        cur.execute(
            "SELECT id FROM evidence_items WHERE source_id = %s AND original_content = %s",
            (source_id, original),
        )
        evidence_id = cur.fetchone()["id"]
    replace_evidence_keywords(
        cur,
        evidence_id,
        keywords,
        {
            "source": "evidence_items.keywords",
            "evidence_kind": kind,
            "published_at": str(published_at or ""),
        },
    )
    return inserted_count


def insert_insight(cur, company_id, route_id, insight_type, title, summary, confidence="", impact="", period="", metadata=None):
    metadata = metadata or {}
    text_for_vector = f"{insight_type} {title} {summary} {json.dumps(metadata, ensure_ascii=False)}"
    cur.execute(
        """
        INSERT INTO insight_items (
          company_id, route_id, insight_type, title, summary, confidence,
          business_impact, time_period, metadata, embedding
        )
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb, %s)
        """,
        (
            company_id,
            route_id,
            insight_type,
            title,
            summary,
            confidence,
            impact,
            period,
            json.dumps(metadata, ensure_ascii=False),
            vectorize_text(text_for_vector),
        ),
    )


def insert_project_memory(cur, memory):
    text_for_vector = " ".join(
        [
            memory.get("memory_layer", ""),
            memory.get("memory_type", ""),
            memory.get("title", ""),
            memory.get("memory_text", ""),
            json.dumps(memory.get("metadata", {}), ensure_ascii=False),
        ]
    )
    cur.execute(
        """
        INSERT INTO project_memories (
          memory_key, memory_layer, memory_type, title, memory_text, source,
          priority, evidence_refs, metadata, embedding, updated_at
        )
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s::jsonb, %s::jsonb, %s, NOW())
        ON CONFLICT (memory_key) DO UPDATE
        SET memory_layer = EXCLUDED.memory_layer,
            memory_type = EXCLUDED.memory_type,
            title = EXCLUDED.title,
            memory_text = EXCLUDED.memory_text,
            source = EXCLUDED.source,
            priority = EXCLUDED.priority,
            evidence_refs = EXCLUDED.evidence_refs,
            metadata = EXCLUDED.metadata,
            embedding = EXCLUDED.embedding,
            updated_at = NOW()
        """,
        (
            memory["memory_key"],
            memory["memory_layer"],
            memory["memory_type"],
            memory["title"],
            memory["memory_text"],
            memory.get("source", "system"),
            memory.get("priority", 3),
            json.dumps(memory.get("evidence_refs", []), ensure_ascii=False),
            json.dumps(memory.get("metadata", {}), ensure_ascii=False),
            vectorize_text(text_for_vector),
        ),
    )


def seed_project_memories(conn):
    with conn.cursor() as cur:
        for memory in PROJECT_MEMORY_SEEDS:
            insert_project_memory(cur, memory)
    conn.commit()


def analyze_database(conn):
    with conn.cursor() as cur:
        cur.execute("ANALYZE companies")
        cur.execute("ANALYZE routes")
        cur.execute("ANALYZE sources")
        cur.execute("ANALYZE keywords")
        cur.execute("ANALYZE evidence_keywords")
        cur.execute("ANALYZE evidence_items")
        cur.execute("ANALYZE insight_items")
        cur.execute("ANALYZE project_memories")
    conn.commit()


def seed_database(conn):
    with conn.cursor() as cur:
        cur.execute("SELECT COUNT(*) AS count FROM evidence_items")
        if cur.fetchone()["count"] > 0:
            return

        seed = json.loads(SEED_PATH.read_text(encoding="utf-8"))
        collected_at = seed.get("collectedAt")

        company_ids = {}
        company_ids["DFDS"] = insert_company(cur, "DFDS", is_baseline=True)
        for competitor in seed.get("competitors", []):
            name = competitor.get("company")
            if name:
                company_ids[name] = insert_company(cur, name, is_baseline=name == "DFDS")

        route_ids = {}
        for key, route in seed.get("routeInsights", {}).items():
            route_ids[key] = insert_route(
                cur,
                key,
                route.get("title", key),
                route.get("reason", ""),
            )
        route_ids.setdefault("all", insert_route(cur, "all", "All signals", "Cross-route passenger ferry signals."))
        route_ids.setdefault(
            "smalltalk",
            insert_route(
                cur,
                "smalltalk",
                "Smalltalk",
                "Bridge replies for greetings, thanks, weather, jokes, and other light prompts.",
            ),
        )

        source_types = {
            "trustpilot": "review",
            "google-play": "app_store",
            "apple-app-store": "app_store",
            "google-reviews": "location_review",
            "reddit": "discussion",
            "company-news": "news",
            "survey-csv": "future_upload",
            "mia-smalltalk-playbook": "conversation_playbook",
        }
        source_names = {
            "trustpilot": "Trustpilot",
            "google-play": "Google Play Reviews",
            "apple-app-store": "Apple App Store Reviews",
            "google-reviews": "Google Reviews",
            "reddit": "Reddit",
            "company-news": "Company News",
            "survey-csv": "Future Survey CSV",
            "mia-smalltalk-playbook": "Mia Small Talk Playbook",
        }
        source_ids = {}
        for key, name in source_names.items():
            source_ids[key] = insert_source(
                cur,
                key,
                name,
                source_types[key],
                "Ready" if key != "survey-csv" else "Upload later",
                last_checked_at=collected_at if key != "survey-csv" else None,
            )
        for coverage in seed.get("sourceCoverage", []):
            key = canonical_source_key(coverage.get("source"))
            source_ids[key] = insert_source(
                cur,
                key,
                source_names.get(key, coverage.get("source", key)),
                source_types.get(key, "public_source"),
                coverage.get("status", "Collected"),
                coverage.get("logo", ""),
                notes=f"{coverage.get('records', '')}. {coverage.get('insight', '')}".strip(),
                last_checked_at=collected_at if key != "survey-csv" else None,
            )

        dfds_id = company_ids["DFDS"]
        all_route_id = route_ids["all"]

        for app in seed.get("appStores", []):
            source_key = canonical_source_key(app.get("platform"))
            insert_evidence(
                cur,
                source_ids[source_key],
                dfds_id,
                all_route_id,
                "app_store_rating",
                f"{app.get('app')} on {app.get('platform')}: {app.get('rating')} from {app.get('reviewCount')}. {app.get('summary')}",
                app.get("summary"),
                parse_rating(app.get("rating")),
                parse_review_count(app.get("reviewCount")),
                collected_at,
                app.get("url", ""),
                {
                    "platform": app.get("platform"),
                    "app": app.get("app"),
                    "package": app.get("package"),
                    "status": app.get("status"),
                    "normalized": app.get("normalized"),
                },
            )

        for review in seed.get("appReviewEvidence", []):
            key = "newcastle-ijmuiden" if "newcastle" in json.dumps(review).lower() or "ijmuiden" in json.dumps(review).lower() else "all"
            insert_evidence(
                cur,
                source_ids[canonical_source_key(review.get("platform"))],
                dfds_id,
                route_ids.get(key, all_route_id),
                "app_review",
                review.get("evidence", ""),
                review.get("businessMeaning", ""),
                parse_rating(review.get("rating")),
                None,
                review.get("date"),
                "",
                {
                    "title": review.get("title"),
                    "theme": review.get("theme"),
                    "platform": review.get("platform"),
                },
            )

        for location in seed.get("googleReviewLocations", []):
            insert_evidence(
                cur,
                source_ids["google-reviews"],
                dfds_id,
                route_ids.get(location.get("route"), all_route_id),
                "google_location_review",
                f"{location.get('name')}: {location.get('signal')}",
                location.get("signal", ""),
                parse_rating(location.get("rating")),
                parse_review_count(location.get("reviews")),
                collected_at,
                location.get("url", ""),
                {
                    "location": location.get("name"),
                    "route": location.get("route"),
                    "source_detail": location.get("source"),
                    "rating_text": location.get("rating"),
                    "review_count_text": location.get("reviews"),
                },
            )

        for signal in seed.get("signals", []):
            key = canonical_source_key(signal.get("source"))
            insert_evidence(
                cur,
                source_ids.get(key, source_ids["company-news"]),
                dfds_id,
                route_ids.get(signal.get("route"), all_route_id),
                "public_feedback_signal",
                signal.get("evidence", ""),
                signal.get("businessMeaning", ""),
                parse_rating(signal.get("evidence")),
                parse_review_count(signal.get("evidence")),
                collected_at,
                "",
                {
                    "theme": signal.get("theme"),
                    "sentiment": signal.get("sentiment"),
                    "journey_stage": signal.get("journeyStage"),
                    "source_label": signal.get("source"),
                },
            )

        for record in build_smalltalk_records():
            source_key = record.get("source_key", "mia-smalltalk-playbook")
            if source_key not in source_ids:
                source_ids[source_key] = insert_source(
                    cur,
                    source_key,
                    record.get("source_name", "Mia Small Talk Playbook"),
                    record.get("source_type", "conversation_playbook"),
                    record.get("source_status", "Curated"),
                    source_url=record.get("url", ""),
                    notes=record.get("source_notes", ""),
                    last_checked_at=collected_at,
                )
            insert_evidence(
                cur,
                source_ids[source_key],
                dfds_id,
                route_ids.get(record.get("route_key"), all_route_id),
                record.get("evidence_kind", "smalltalk_theme_card"),
                record.get("original_content", ""),
                record.get("translated_content", ""),
                record.get("rating"),
                record.get("review_count"),
                record.get("published_at"),
                record.get("url", ""),
                record.get("metadata", {}),
                record.get("keywords", []),
            )

        for root_cause in seed.get("rootCauses", []):
            insert_insight(
                cur,
                dfds_id,
                all_route_id,
                "root_cause",
                root_cause.get("title", ""),
                root_cause.get("summary", ""),
                root_cause.get("confidence", ""),
                root_cause.get("impact", ""),
                collected_at,
            )

        for theme in seed.get("voiceThemes", []):
            insert_insight(
                cur,
                dfds_id,
                all_route_id,
                "customer_insight",
                theme.get("theme", ""),
                f"{theme.get('customerMeaning', '')} {theme.get('marketingAction', '')}",
                "",
                theme.get("sentiment", ""),
                collected_at,
                {
                    "routes": theme.get("routes", []),
                    "sources": theme.get("sources", []),
                    "evidence": theme.get("evidence", []),
                },
            )

        for recommendation in seed.get("recommendations", []):
            insert_insight(
                cur,
                dfds_id,
                all_route_id,
                "recommendation",
                recommendation.get("title", ""),
                recommendation.get("reasoning", ""),
                recommendation.get("priority", ""),
                recommendation.get("expectedImprovement", ""),
                collected_at,
                recommendation,
            )

        for competitor in seed.get("competitors", []):
            company_id = company_ids.get(competitor.get("company"), dfds_id)
            insert_insight(
                cur,
                company_id,
                all_route_id,
                "competitor_comparison",
                competitor.get("company", ""),
                competitor.get("userView", ""),
                "",
                competitor.get("learning", ""),
                collected_at,
                competitor,
            )

    conn.commit()


def initialize_database():
    if not os.environ.get("DATABASE_URL") or psycopg is None:
        return False

    last_error = None
    for _ in range(20):
        conn = None
        try:
            conn = connect_db(register=False)
            execute_schema(conn)
            if register_vector is not None:
                register_vector(conn)
            seed_database(conn)
            seed_project_memories(conn)
            seed_agent_monitoring_demo(conn)
            analyze_database(conn)
            return True
        except Exception as exc:
            last_error = exc
            time.sleep(1)
        finally:
            if conn is not None:
                conn.close()

    raise RuntimeError(f"Database initialization failed after waiting for PostgreSQL: {last_error}")


def seed_agent_monitoring_demo(conn):
    """Seed the administrator demo using the production monitoring tables, idempotently."""
    events = build_demo_events()
    user_ids = {}
    with conn.cursor() as cur:
        for event in events:
            user_key = event["user_key"]
            if user_key in user_ids:
                continue
            user_id = str(uuid.uuid5(uuid.UUID("e4c8dc2a-8ed7-4d39-a96d-6f13e580c4cf"), user_key))
            cur.execute(
                """INSERT INTO app_users (id, username, password_hash, role)
                   VALUES (%s, %s, %s, 'project_user')
                   ON CONFLICT (id) DO UPDATE SET username = EXCLUDED.username
                   RETURNING id""",
                (user_id, event["username"], hash_password("246810")),
            )
            user_ids[user_key] = str(cur.fetchone()["id"])
        for event in events:
            cur.execute(
                """INSERT INTO chat_sessions (id, user_id, user_language, active_view, route_focus, created_at, updated_at)
                   VALUES (%s, %s, 'English', 'agent-control', %s, %s, %s)
                   ON CONFLICT (id) DO UPDATE SET
                     user_id = EXCLUDED.user_id,
                     user_language = EXCLUDED.user_language,
                     active_view = EXCLUDED.active_view,
                     route_focus = EXCLUDED.route_focus,
                     created_at = EXCLUDED.created_at,
                     updated_at = EXCLUDED.updated_at""",
                (event["session_id"], user_ids[event["user_key"]], event["route_key"], event["created_at"], event["created_at"]),
            )
            cur.execute(
                """INSERT INTO chat_messages (id, session_id, role, message_text, evidence_context, page_context, created_at)
                   VALUES (%s, %s, 'user', %s, '{}'::jsonb, %s::jsonb, %s)
                   ON CONFLICT (id) DO UPDATE SET
                     session_id = EXCLUDED.session_id,
                     role = EXCLUDED.role,
                     message_text = EXCLUDED.message_text,
                     evidence_context = EXCLUDED.evidence_context,
                     page_context = EXCLUDED.page_context,
                     created_at = EXCLUDED.created_at""",
                (str(uuid.uuid5(uuid.UUID("e4c8dc2a-8ed7-4d39-a96d-6f13e580c4cf"), f"message:{event['id']}")), event["session_id"], event["message"], json.dumps({"demo_agent_monitoring": True, "route_key": event["route_key"]}), event["created_at"]),
            )
            cur.execute(
                """INSERT INTO rag_usage_events (id, user_id, chat_session_id, retrieval_trace, input_tokens, output_tokens, total_tokens, created_at)
                   VALUES (%s, %s, %s, %s::jsonb, %s, %s, %s, %s)
                   ON CONFLICT (id) DO UPDATE SET
                     user_id = EXCLUDED.user_id,
                     chat_session_id = EXCLUDED.chat_session_id,
                     retrieval_trace = EXCLUDED.retrieval_trace,
                     input_tokens = EXCLUDED.input_tokens,
                     output_tokens = EXCLUDED.output_tokens,
                     total_tokens = EXCLUDED.total_tokens,
                     created_at = EXCLUDED.created_at""",
                (event["id"], user_ids[event["user_key"]], event["session_id"], json.dumps(event["retrieval_trace"]), event["input_tokens"], event["output_tokens"], event["total_tokens"], event["created_at"]),
            )
    conn.commit()


def retrieve_evidence_from_db(payload, limit=6):
    conn = connect_db()
    if conn is None:
        return {"items": [], "semanticAvailable": False}

    message = str(payload.get("message", ""))
    active_view = str(payload.get("activeView", ""))
    route = route_key(payload.get("routeKey") or "all")
    terms = query_terms(active_view, route, message)
    semantic_available = False
    semantic_embedding = None
    settings = rag_embeddings.load_embedding_settings()
    if settings is not None:
        try:
            semantic_embedding = rag_embeddings.embed_text(
                f"{active_view} {route} {message}", settings
            )
            semantic_available = True
        except rag_embeddings.EmbeddingUnavailable:
            semantic_embedding = None

    candidate_limit = evidence_candidate_limit(limit)
    sql, params = hybrid_retrieval.build_hybrid_evidence_query(
        terms, route, semantic_available, candidate_limit
    )
    if semantic_available:
        params["semantic_embedding"] = rag_embeddings.semantic_vector(semantic_embedding)

    try:
        with conn.cursor() as cur:
            cur.execute(sql, params)
            rows = cur.fetchall()
    finally:
        conn.close()

    ranked_items = hybrid_retrieval.rank_evidence_rows(rows, route, limit=candidate_limit)
    return {
        "items": dedupe_evidence_for_context(ranked_items, limit=limit),
        "semanticAvailable": semantic_available,
    }


def retrieve_insights_from_db(payload, limit=5):
    conn = connect_db()
    if conn is None:
        return []

    message = str(payload.get("message", ""))
    active_view = str(payload.get("activeView", ""))
    route = str(payload.get("routeKey") or "all")
    query_vector = vectorize_text(f"{active_view} {route} {message}")
    terms = query_terms(active_view, route, message)
    params = {"embedding": query_vector, "limit": limit}

    search_blob = (
        "lower(concat_ws(' ', i.insight_type, i.title, i.summary, i.confidence, "
        "i.business_impact, i.time_period, c.name, r.display_name, i.metadata::text))"
    )
    score_parts = []
    for index, term in enumerate(terms):
        key = f"insight_term_{index}"
        params[key] = f"%{term.lower()}%"
        score_parts.append(f"CASE WHEN {search_blob} LIKE %({key})s THEN 1 ELSE 0 END")
    text_score_sql = " + ".join(score_parts) if score_parts else "0"

    sql = f"""
        SELECT
          i.id,
          i.insight_type,
          i.title,
          i.summary,
          i.supporting_evidence_ids,
          i.confidence,
          i.business_impact,
          i.time_period,
          i.metadata,
          c.name AS company_name,
          r.display_name AS route_name,
          r.route_key,
          ({text_score_sql}) AS text_score,
          i.embedding <=> %(embedding)s AS distance
        FROM insight_items i
        LEFT JOIN companies c ON c.id = i.company_id
        LEFT JOIN routes r ON r.id = i.route_id
        WHERE i.embedding IS NOT NULL
        ORDER BY text_score DESC, i.embedding <=> %(embedding)s
        LIMIT %(limit)s
    """

    try:
        with conn.cursor() as cur:
            cur.execute(sql, params)
            rows = cur.fetchall()
    finally:
        conn.close()

    return [
        {
            "id": row["id"],
            "type": row["insight_type"],
            "title": row["title"],
            "summary": row["summary"],
            "supportingEvidenceIds": row["supporting_evidence_ids"] or [],
            "confidence": row["confidence"],
            "businessImpact": row["business_impact"],
            "timePeriod": row["time_period"],
            "company": row["company_name"],
            "route": row["route_name"] or "All signals",
            "routeKey": row["route_key"] or "all",
            "metadata": row.get("metadata") or {},
            "textScore": int(row.get("text_score") or 0),
            "distance": float(row["distance"]) if row["distance"] is not None else None,
        }
        for row in rows
    ]


def retrieve_project_memories_from_db(payload, limit=4):
    conn = connect_db()
    if conn is None:
        return []

    message = str(payload.get("message", ""))
    active_view = str(payload.get("activeView", ""))
    route = str(payload.get("routeKey") or "all")
    query_vector = vectorize_text(f"{active_view} {route} {message}")

    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT
                  id, memory_key, memory_layer, memory_type, title, memory_text,
                  source, priority, evidence_refs, metadata,
                  embedding <=> %(embedding)s AS distance
                FROM project_memories
                WHERE status = 'active'
                  AND (expires_at IS NULL OR expires_at > NOW())
                  AND embedding IS NOT NULL
                ORDER BY priority DESC, embedding <=> %(embedding)s
                LIMIT %(limit)s
                """,
                {"embedding": query_vector, "limit": limit},
            )
            rows = cur.fetchall()
            if rows:
                cur.execute(
                    "UPDATE project_memories SET last_used_at = NOW() WHERE id = ANY(%s)",
                    ([row["id"] for row in rows],),
                )
                conn.commit()
    finally:
        conn.close()

    return [
        {
            "id": row["id"],
            "key": row["memory_key"],
            "layer": row["memory_layer"],
            "type": row["memory_type"],
            "title": row["title"],
            "text": row["memory_text"],
            "source": row["source"],
            "priority": row["priority"],
            "evidenceRefs": row["evidence_refs"] or [],
            "metadata": row.get("metadata") or {},
            "distance": float(row["distance"]) if row["distance"] is not None else None,
        }
        for row in rows
    ]


def retrieve_recent_chat_history(payload, limit=6):
    conn = connect_db()
    if conn is None:
        return []

    session_id = valid_uuid(payload.get("sessionId"))
    identity = payload.get("_identity") or {}
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT m.role, m.message_text, m.created_at
                FROM chat_messages m
                JOIN chat_sessions s ON s.id = m.session_id
                WHERE m.session_id = %s AND s.user_id = %s
                ORDER BY m.created_at DESC
                LIMIT %s
                """,
                (session_id, identity.get("id"), limit),
            )
            rows = cur.fetchall()
    finally:
        conn.close()

    return [
        {
            "role": row["role"],
            "message": str(row["message_text"])[:700],
            "createdAt": row["created_at"].isoformat() if row["created_at"] else "",
        }
        for row in reversed(rows)
    ]


def compact_evidence(evidence_items):
    compacted = []
    for item in evidence_items[:6]:
        compacted.append(
            {
                "evidence_id": item.get("id", ""),
                "type": item.get("type", ""),
                "title": item.get("title", ""),
                "content": str(item.get("body", ""))[:700],
                "business_meaning": str(item.get("translatedContent", ""))[:500],
                "source": item.get("source", ""),
                "company": item.get("company", ""),
                "route": item.get("route", ""),
                "rating": item.get("rating"),
                "review_count": item.get("reviewCount"),
                "timestamp": item.get("timestamp", ""),
                "url": item.get("url", ""),
                "keywords": item.get("keywords", []),
                "sourceTier": item.get("sourceTier", "public_snapshot"),
                "isSynthetic": bool(item.get("isSynthetic")),
                "sourceVersion": item.get("sourceVersion", "legacy-v1"),
                "scoreComponents": item.get("scoreComponents", {}),
                "metadata": {
                    key: item.get("metadata", {}).get(key)
                    for key in [
                        "synthetic_source",
                        "basis",
                        "confidence",
                        "source_family",
                        "synthetic_record_id",
                        "journey_stage",
                        "sentiment",
                        "reliability_note",
                    ]
                    if key in item.get("metadata", {})
                },
            }
        )
    return compacted


def compact_insights(insight_items):
    compacted = []
    for item in insight_items[:5]:
        compacted.append(
            {
                "insight_id": item.get("id", ""),
                "type": item.get("type", ""),
                "title": item.get("title", ""),
                "summary": item.get("summary", ""),
                "confidence": item.get("confidence", ""),
                "business_impact": item.get("businessImpact", ""),
                "time_period": item.get("timePeriod", ""),
                "company": item.get("company", ""),
                "route": item.get("route", ""),
                "evidence_ids": item.get("supportingEvidenceIds", []),
            }
        )
    return compacted


def compact_memories(memory_items):
    compacted = []
    for item in memory_items[:4]:
        compacted.append(
            {
                "memory_id": item.get("id", ""),
                "key": item.get("key", ""),
                "layer": item.get("layer", ""),
                "type": item.get("type", ""),
                "title": item.get("title", ""),
                "text": item.get("text", ""),
                "priority": item.get("priority", 0),
            }
        )
    return compacted


def compact_history(history_items):
    compacted = []
    for item in history_items[:6]:
        compacted.append(
            {
                "role": item.get("role", ""),
                "message": item.get("message", ""),
                "created_at": item.get("createdAt", ""),
            }
        )
    return compacted


def compact_page_context(payload):
    page_context = payload.get("pageContext") or {}
    return page_context if isinstance(page_context, dict) else {"raw": str(page_context)}


def build_rag_context(payload):
    evidence_started = time.perf_counter()
    retrieval_result = retrieve_evidence_from_db(payload)
    if isinstance(retrieval_result, dict):
        evidence_items = retrieval_result.get("items") or []
        semantic_available = bool(retrieval_result.get("semanticAvailable"))
    else:
        # Keep callers that still provide the previous list shape working.
        evidence_items = retrieval_result or []
        semantic_available = False
    if not evidence_items:
        evidence_items = fallback_evidence_from_payload(payload)
        semantic_available = False

    insights_started = time.perf_counter()
    insight_items = retrieve_insights_from_db(payload)
    memories_started = time.perf_counter()
    memory_items = retrieve_project_memories_from_db(payload)
    history_started = time.perf_counter()
    history_items = retrieve_recent_chat_history(payload)
    upload_context = payload.get("_uploadedBatchContext")
    page_context = compact_page_context(payload)
    if isinstance(upload_context, dict):
        page_context["uploadedBatch"] = upload_context
        overview = upload_context.get("datasetOverview") or {}
        file_overview = upload_context.get("fileOverview") or {}
        source_names = ", ".join(str(item) for item in overview.get("dataSources") or []) or "No profiled worksheets"
        quality = upload_context.get("qualitySummary") or {}
        dimension_evidence = []
        for index, profile in enumerate((upload_context.get("dimensionProfiles") or [])[:16], start=1):
            if not isinstance(profile, dict):
                continue
            field = str(profile.get("field") or "").strip()
            if not field or RESTRICTED_SOURCE_FIELD.search(field):
                continue
            top_values = [
                f"{item.get('value', '')}: {int(item.get('count') or 0)}"
                for item in (profile.get("topValues") or [])[:8]
                if isinstance(item, dict) and str(item.get("value", "")).strip()
            ]
            if not top_values:
                continue
            dimension_evidence.append({
                "id": f"uploaded-batch:{upload_context.get('batchId', '')}:field:{index}",
                "title": f"Uploaded data: {field}",
                "body": (
                    f"In {profile.get('reference', 'the uploaded workbook')}, {field} has "
                    f"{int(profile.get('observedValues') or 0)} observed values. "
                    f"Most common values: {'; '.join(top_values)}."
                ),
                "source": "Authenticated IT Data Flow upload",
                "route": "Current uploaded batch",
                "rating": None,
                "reviewCount": None,
            })
        evidence_items = [
            {
                "id": f"uploaded-batch:{upload_context.get('batchId', '')}",
                "title": "Current uploaded batch",
                "body": (
                    f"Version {upload_context.get('version') or 'current'} contains "
                    f"{int(file_overview.get('profiledRows') or 0):,} profiled rows across "
                    f"{int(file_overview.get('worksheetCount') or 0)} worksheets. Sources: {source_names}."
                ),
                "source": "Authenticated IT Data Flow upload",
                "route": "Current uploaded batch",
                "rating": None,
                "reviewCount": None,
            },
            {
                "id": f"uploaded-batch:{upload_context.get('batchId', '')}:cleaning",
                "title": "Uploaded data cleaning",
                "body": (
                    f"The batch had {int(quality.get('receivedRows') or 0):,} received rows and "
                    f"{int(quality.get('cleanedRows') or 0):,} cleaned rows; "
                    f"{int(quality.get('blankRowsRemoved') or 0):,} blank rows removed; "
                    f"{int(quality.get('duplicateRowsRemoved') or 0):,} duplicate rows removed; "
                    f"{int(quality.get('missingValuesStandardized') or 0):,} missing values standardized."
                ),
                "source": "Authenticated IT Data Flow upload",
                "route": "Current uploaded batch",
                "rating": None,
                "reviewCount": None,
            },
            *dimension_evidence,
            *evidence_items,
        ]

    return {
        "pageContext": page_context,
        "evidenceItems": evidence_items,
        "insightItems": insight_items,
        "memoryItems": memory_items,
        "historyItems": history_items,
        "retrieval": {
            "mode": "hybrid" if semantic_available else "bm25" if evidence_items else "unavailable",
            "candidateCount": len(evidence_items),
            "semanticAvailable": semantic_available,
            "layers": [
                {"name": "evidence_retrieval", "durationMs": round((insights_started - evidence_started) * 1000, 1), "records": len(evidence_items)},
                {"name": "insight_retrieval", "durationMs": round((memories_started - insights_started) * 1000, 1), "records": len(insight_items)},
                {"name": "memory_retrieval", "durationMs": round((history_started - memories_started) * 1000, 1), "records": len(memory_items)},
                {"name": "chat_history", "durationMs": round((time.perf_counter() - history_started) * 1000, 1), "records": len(history_items)},
                {"name": "uploaded_batch_context", "durationMs": 0, "records": 1 if isinstance(upload_context, dict) else 0},
            ],
        },
    }


def fallback_evidence_from_payload(payload):
    fallback_items = []
    for item in (payload.get("evidenceItems") or [])[:6]:
        fallback_items.append(
            {
                "id": item.get("id", "frontend-fallback"),
                "type": item.get("type", "frontend_evidence"),
                "title": item.get("title", "Frontend evidence"),
                "body": item.get("body", ""),
                "translatedContent": item.get("businessMeaning", ""),
                "source": item.get("source", "Frontend dashboard dataset"),
                "company": "DFDS",
                "route": item.get("route", "All signals"),
                "routeKey": item.get("route", "all"),
                "rating": parse_rating(item.get("scoreText")),
                "reviewCount": parse_review_count(item.get("scoreText")),
                "timestamp": "",
                "url": item.get("url", ""),
                "metadata": {"fallback": "frontend_embedded_dataset"},
                "distance": None,
                "sourceTier": item.get("sourceTier", "public_snapshot"),
                "isSynthetic": bool(item.get("isSynthetic")),
                "sourceVersion": item.get("sourceVersion", "frontend-fallback"),
                "scoreComponents": item.get("scoreComponents", {}),
            }
        )
    return fallback_items


def localized_evidence_summary(item, language):
    source = item.get("source") or "dashboard"
    route = item.get("route") or "All signals"
    rating = item.get("rating")
    review_count = item.get("reviewCount")
    score = []
    if rating is not None:
        score.append(f"{rating}/5")
    if review_count:
        score.append(f"{review_count}")
    score_text = ", ".join(score)

    summaries = {
        "Chinese": f"来自 {source} 的证据摘要，适用于 {route}。它用于判断客户体验、航线沟通、竞品对比或行动建议是否有公开证据支撑。",
        "Danish": f"Evidensresumé fra {source} for {route}. Det bruges til at vurdere kundeoplevelse, rutekommunikation, konkurrentbenchmark eller anbefalede handlinger.",
        "Spanish": f"Resumen de evidencia de {source} para {route}. Sirve para evaluar experiencia del cliente, comunicación de ruta, benchmark competitivo o acciones recomendadas.",
        "French": f"Résumé de preuve provenant de {source} pour {route}. Il sert à évaluer l’expérience client, la communication par route, le benchmark concurrentiel ou les actions recommandées.",
        "German": f"Evidenz-Zusammenfassung aus {source} für {route}. Sie hilft bei der Bewertung von Kundenerlebnis, Routenkommunikation, Wettbewerbsbenchmark oder empfohlenen Maßnahmen.",
        "Japanese": f"{route} に関する {source} からのエビデンス要約です。顧客体験、ルート別コミュニケーション、競合ベンチマーク、推奨アクションの判断に使います。",
        "Korean": f"{route}에 대한 {source} 근거 요약입니다. 고객 경험, 노선 커뮤니케이션, 경쟁사 벤치마크 또는 권장 조치를 판단하는 데 사용됩니다."
    }
    summary = summaries.get(language)
    if summary:
        return f"{summary} {score_text}".strip()

    score_suffix = f" ({score_text})" if score_text else ""
    return f"{item.get('title')} - {item.get('body')}{score_suffix}"


def should_show_evidence_details(message):
    text = str(message or "").strip().lower()
    if not text:
        return False
    if re.search(
        r"\b(evidence|source|sources|citation|citations|proof|supporting data|supporting evidence|show me why|where is this from|what is this based on)\b",
        text,
        re.I,
    ):
        return True
    return bool(re.search(r"(证据|佐证|来源|出处|引用|根据|支撑|原文|哪条|哪些数据|数据依据)", str(message)))


def plain_evidence_takeaway(item, language):
    title = str(item.get("title") or item.get("type") or "DFDS signal").strip()
    body = str(item.get("translatedContent") or item.get("body") or "").strip()
    route = str(item.get("route") or "").strip()
    if language == "Chinese":
        route_text = f"（{route}）" if route else ""
        if body:
            return f"{title}{route_text}：{body}"
        return f"{title}{route_text} 是当前判断里的主要信号。"
    if body:
        return f"{title}: {body}"
    if route:
        return f"{title}: this is a relevant signal for {route}."
    return f"{title}: this is a relevant DFDS signal."


def evidence_detail_heading(language):
    headings = {
        "Chinese": "对应佐证：",
        "Danish": "Relevant dokumentation:",
        "Spanish": "Evidencia de apoyo:",
        "French": "Éléments à l’appui :",
        "German": "Unterstützende Belege:",
        "Japanese": "補足根拠:",
        "Korean": "관련 근거:",
    }
    return headings.get(language, "Supporting evidence:")


def strip_answer_labels(answer, message):
    cleaned = str(answer or "").strip()
    if not cleaned:
        return cleaned

    cleaned = re.sub(r"^\s*(Short answer|Answer)\s*:\s*", "", cleaned, flags=re.I)
    cleaned = re.sub(r"^\s*简短结论\s*[:：]\s*", "", cleaned)
    cleaned = re.sub(r"Evidence shown in this answer\s*:?", "", cleaned, flags=re.I)

    if should_show_evidence_details(message):
        cleaned = re.sub(
            r"(^|\n)\s*Evidence used\s*:\s*",
            f"\\1{evidence_detail_heading('English')}\n",
            cleaned,
            flags=re.I,
        )
        return cleaned.strip()

    section_patterns = [
        r"(^|\n)\s*Evidence used\s*:.*$",
        r"(^|\n)\s*Sources\s*:.*$",
        r"(^|\n)\s*使用到的证据\s*[:：].*$",
        r"(^|\n)\s*来源\s*[:：].*$",
        r"(^|\n)\s*Kilder brugt\s*:.*$",
        r"(^|\n)\s*Fuentes utilizadas\s*:.*$",
        r"(^|\n)\s*Sources utilisées\s*:.*$",
        r"(^|\n)\s*Genutzte Quellen\s*:.*$",
        r"(^|\n)\s*参照した情報源\s*:.*$",
        r"(^|\n)\s*사용한 출처\s*:.*$",
    ]
    for pattern in section_patterns:
        cleaned = re.sub(pattern, "", cleaned, flags=re.I | re.S)
    return cleaned.strip()


def evidence_only_answer(payload, evidence_items, reason):
    language = str(payload.get("language", "English"))
    message = str(payload.get("message", ""))
    evidence_requested = should_show_evidence_details(message)
    sources = sorted({item.get("source", "") for item in evidence_items if item.get("source")})
    plain_bullets = []
    detail_bullets = []
    for index, item in enumerate(evidence_items[:4], start=1):
        plain_bullets.append(f"{index}. {plain_evidence_takeaway(item, language)}")
        detail_bullets.append(f"{index}. {localized_evidence_summary(item, language)}")

    if not plain_bullets:
        return bridge_answer(payload)

    if evidence_requested:
        heading = evidence_detail_heading(language)
        if language == "Chinese":
            return (
                f"{reason} 下面是对应佐证，便于你检查这个判断来自哪里。\n\n"
                f"{heading}\n"
                + "\n".join(detail_bullets)
            )
        return (
            f"{reason} Here is the supporting context behind that answer.\n\n"
            f"{heading}\n"
            + "\n".join(detail_bullets)
            + (f"\n\nSource names included above: {', '.join(sources)}." if sources else "")
        )

    if language == "Chinese":
        return (
            f"{reason} 当前信号已经能形成一个方向性判断：先看最影响客户信任和下一步行动的主题，而不是只看单条评论。\n\n"
            "这意味着：\n"
            "- 可以先用它判断 marketing / CX 下一步该关注什么。\n"
            "- 把它当成内部工作结论，而不是直接对外承诺。\n"
            "- 如果要对外承诺或做高风险决策，还需要继续补真实公开来源或内部客户数据。\n\n"
            "主要信号：\n"
            + "\n".join(plain_bullets[:3])
        )

    if language == "Danish":
        return (
            f"{reason} De aktuelle DFDS-signaler peger på en retning for næste beslutning.\n\n"
            "Det betyder:\n"
            "- Brug svaret som en intern arbejdsopsummering.\n"
            "- Prioriter det tema, der påvirker kundernes tillid mest.\n"
            "- Indsaml mere offentlig eller intern kundedata før større eksterne claims."
        )

    if language == "Spanish":
        return (
            f"{reason} Las señales actuales de DFDS ya apuntan a una lectura direccional para la siguiente decisión.\n\n"
            "Qué implica:\n"
            "- Úsalo como resumen interno de trabajo.\n"
            "- Prioriza el tema que más afecte la confianza del cliente.\n"
            "- Añade más datos públicos o internos antes de hacer afirmaciones externas fuertes."
        )

    if language == "French":
        return (
            f"{reason} Les signaux DFDS disponibles donnent déjà une lecture directionnelle pour la prochaine décision.\n\n"
            "Ce que cela implique :\n"
            "- Utilisez-le comme synthèse de travail interne.\n"
            "- Priorisez le thème qui influence le plus la confiance client.\n"
            "- Ajoutez davantage de données publiques ou internes avant toute affirmation externe forte."
        )

    if language == "German":
        return (
            f"{reason} Die aktuellen DFDS-Signale reichen für eine richtungsweisende Einschätzung der nächsten Entscheidung.\n\n"
            "Das bedeutet:\n"
            "- Nutze dies als interne Arbeitszusammenfassung.\n"
            "- Priorisiere das Thema, das das Kundenvertrauen am stärksten beeinflusst.\n"
            "- Ergänze weitere öffentliche oder interne Kundendaten vor starken externen Aussagen."
        )

    if language == "Japanese":
        return (
            f"{reason} 現在のDFDSシグナルから、次の判断に使える方向性は見えています。\n\n"
            "意味すること:\n"
            "- 内部向けの作業サマリーとして使えます。\n"
            "- 顧客信頼に最も影響するテーマを優先してください。\n"
            "- 強い外部発信の前には、追加の公開データや内部顧客データを確認してください。"
        )

    if language == "Korean":
        return (
            f"{reason} 현재 DFDS 신호만으로도 다음 의사결정을 위한 방향성은 볼 수 있습니다.\n\n"
            "의미:\n"
            "- 내부 작업 요약으로 활용하세요.\n"
            "- 고객 신뢰에 가장 큰 영향을 주는 주제를 우선순위로 두세요.\n"
            "- 강한 외부 메시지를 내기 전에는 공개 데이터나 내부 고객 데이터를 더 보강하세요."
        )

    return (
        f"{reason} The current DFDS signals are enough for a directional working read: focus first on the customer trust issue, then decide the next Marketing or CX action.\n\n"
        "What it means:\n"
        "- Use this as an internal decision summary, not an external claim.\n"
        "- Prioritize the theme most likely to affect confidence before or during travel.\n"
        "- Add more real public or internal customer data before making high-stakes claims.\n\n"
        "Main signals:\n"
        + "\n".join(plain_bullets[:3])
    )


def bridge_answer(payload):
    language = str(payload.get("language", "English"))
    intent = classify_conversation_intent(str(payload.get("message", "")))
    if intent == "liveExternal":
        responses = {
            "Chinese": "我这里不能查看实时天气。不过如果你想评估天气或延误对客户体验的影响，我可以帮你看 DFDS 报告里的 route signals、delay themes 和 customer expectations。你想看 Dover-Calais，还是其他航线？",
            "Japanese": "このレポート内ではリアルタイムの天気は確認できません。ただし、天候や遅延が顧客体験にどう影響しているかは、route signals や delay themes から一緒に見られます。どのルートを見ますか？",
            "Korean": "이 보고서에서는 실시간 날씨를 확인할 수는 없어요. 대신 날씨나 지연이 고객 경험에 어떤 영향을 주는지는 route signals 와 delay themes 로 함께 볼 수 있습니다. 어떤 노선을 볼까요?",
            "Spanish": "No puedo consultar el clima en tiempo real desde este informe. Pero sí puedo ayudarte a ver cómo las rutas, retrasos y expectativas del cliente aparecen en la evidencia de DFDS. ¿Quieres revisar Dover-Calais u otra ruta?",
            "French": "Je ne peux pas vérifier la météo en temps réel depuis ce rapport. Mais je peux vous aider à voir comment les routes, les retards et les attentes clients apparaissent dans les preuves DFDS. Voulez-vous regarder Dover-Calais ou une autre route ?",
            "German": "Live-Wetter kann ich in diesem Bericht nicht prüfen. Ich kann aber zeigen, wie Routen, Verspätungen und Kundenerwartungen in den DFDS-Signalen auftauchen. Möchtest du Dover-Calais oder eine andere Route ansehen?",
            "Danish": "Jeg kan ikke tjekke live-vejret i denne rapport. Men jeg kan hjælpe med at se, hvordan ruter, forsinkelser og kundeforventninger viser sig i DFDS-signalerne. Vil du se Dover-Calais eller en anden rute?"
        }
        return responses.get(language, "I can’t check live weather from this report. But I can help you understand how route experience, delays, and customer expectations show up in the DFDS evidence. Would you like to look at Dover-Calais or another route?")

    if intent == "goodbye":
        responses = {
            "Chinese": "再见，我是 Mia。下次你可以直接问我 routes、app reviews、competitors 或 recommendations。",
            "Japanese": "またね、Miaです。次は routes、app reviews、competitors、recommendations をそのまま聞いてください。",
            "Korean": "안녕히 가세요, Mia입니다. 다음에는 routes, app reviews, competitors, recommendations를 바로 물어보세요.",
            "Spanish": "Hasta luego, soy Mia. La próxima vez puedes preguntarme por routes, app reviews, competitors o recommendations.",
            "French": "À bientôt, je suis Mia. La prochaine fois, demandez-moi directement les routes, app reviews, competitors ou recommendations.",
            "German": "Bis dann, ich bin Mia. Frag mich beim nächsten Mal direkt nach routes, app reviews, competitors oder recommendations.",
            "Danish": "Farvel, jeg er Mia. Næste gang kan du bare spørge mig om routes, app reviews, competitors eller recommendations."
        }
        return responses.get(language, "Goodbye, I’m Mia. Next time you can ask me about routes, app reviews, competitors, or recommendations.")

    if intent == "capability":
        responses = {
            "Chinese": "我是 Mia，DFDS Customer Intelligence 里的报告助手。我可以帮你看 routes、app reviews、competitors、recommendations，还有证据和更新日志。",
            "Japanese": "私は Mia です。DFDS Customer Intelligence のレポートアシスタントです。routes、app reviews、competitors、recommendations、証拠、更新ログをお手伝いできます。",
            "Korean": "저는 Mia예요. DFDS Customer Intelligence의 보고서 도우미입니다. routes, app reviews, competitors, recommendations, 근거, 업데이트 로그를 도와드릴 수 있어요.",
            "Spanish": "Soy Mia, la asistente del informe de DFDS Customer Intelligence. Puedo ayudarte con routes, app reviews, competitors, recommendations, evidencia y el registro de cambios.",
            "French": "Je suis Mia, l’assistante du rapport DFDS Customer Intelligence. Je peux aider avec les routes, app reviews, competitors, recommendations, les preuves et le journal des changements.",
            "German": "Ich bin Mia, die Assistentin des DFDS Customer Intelligence Reports. Ich kann bei routes, app reviews, competitors, recommendations, Belegen und dem Änderungsprotokoll helfen.",
            "Danish": "Jeg er Mia, assistenten i DFDS Customer Intelligence. Jeg kan hjælpe med routes, app reviews, competitors, recommendations, evidens og opdateringsloggen."
        }
        return responses.get(language, "I’m Mia, the DFDS Customer Intelligence assistant. I can help with routes, app reviews, competitors, recommendations, evidence, and update logs.")

    if intent == "wellbeing":
        responses = {
            "Chinese": "我状态很好，谢谢你。现在我可以继续帮你看 DFDS 报告里的 routes、app reviews 或 competitors。",
            "Japanese": "元気です、ありがとう。DFDSレポートの routes、app reviews、competitors を続けて見られます。",
            "Korean": "저는 잘 지내고 있어요, 감사합니다. 이제 DFDS 보고서의 routes, app reviews, competitors를 계속 볼 수 있어요.",
            "Spanish": "Estoy bien, gracias. Puedo seguir ayudándote con las routes, app reviews o competitors del informe DFDS.",
            "French": "Je vais bien, merci. Je peux continuer à vous aider avec les routes, les app reviews ou les competitors du rapport DFDS.",
            "German": "Mir geht es gut, danke. Ich kann dir weiter bei den routes, app reviews oder competitors im DFDS-Bericht helfen.",
            "Danish": "Jeg har det godt, tak. Jeg kan fortsætte med at hjælpe dig med routes, app reviews eller competitors i DFDS-rapporten."
        }
        return responses.get(language, "I’m doing well, thanks. I can keep helping with DFDS routes, app reviews, or competitors.")

    if intent == "thanks":
        responses = {
            "Chinese": "不客气，我是 Mia。你还想看 routes、app reviews、competitors 还是 recommendations？",
            "Japanese": "どういたしまして、Miaです。routes、app reviews、competitors、recommendations のどれを見ますか？",
            "Korean": "천만에요, Mia입니다. routes, app reviews, competitors, recommendations 중 무엇을 볼까요?",
            "Spanish": "De nada, soy Mia. ¿Quieres ver routes, app reviews, competitors o recommendations?",
            "French": "Avec plaisir, je suis Mia. Voulez-vous regarder les routes, les app reviews, les competitors ou les recommendations ?",
            "German": "Gern geschehen, ich bin Mia. Möchtest du routes, app reviews, competitors oder recommendations ansehen?",
            "Danish": "Selv tak, jeg er Mia. Vil du se routes, app reviews, competitors eller recommendations?"
        }
        return responses.get(language, "You’re welcome. I’m Mia. Want to look at routes, app reviews, competitors, or recommendations?")

    if intent == "apology":
        responses = {
            "Chinese": "没关系。我们可以直接回到 DFDS 报告里，你想看 routes、app reviews 还是 recommendations？",
            "Japanese": "大丈夫です。DFDSレポートに戻りましょう。routes、app reviews、recommendations のどれを見ますか？",
            "Korean": "괜찮아요. DFDS 보고서로 바로 돌아가죠. routes, app reviews, recommendations 중 무엇을 볼까요?",
            "Spanish": "No pasa nada. Volvamos al informe DFDS. ¿Quieres ver routes, app reviews o recommendations?",
            "French": "Ce n’est rien. Revenons au rapport DFDS. Voulez-vous voir les routes, les app reviews ou les recommendations ?",
            "German": "Kein Problem. Gehen wir zurück zum DFDS-Bericht. Möchtest du routes, app reviews oder recommendations ansehen?",
            "Danish": "Det er helt fint. Lad os gå tilbage til DFDS-rapporten. Vil du se routes, app reviews eller recommendations?"
        }
        return responses.get(language, "No problem. Let’s get back to the DFDS report. Want routes, app reviews, or recommendations?")

    if intent == "joke":
        responses = {
            "Chinese": "我不太会讲段子，但我可以把 DFDS 的证据讲得更清楚。想看路线、评论还是竞品？",
            "Japanese": "おもしろネタは少し苦手ですが、DFDSの証拠はわかりやすくできます。ルート、レビュー、競合のどれを見ますか？",
            "Korean": "농담은 조금 약하지만 DFDS 근거는 깔끔하게 설명할 수 있어요. 노선, 리뷰, 경쟁사 중 무엇을 볼까요?",
            "Spanish": "No soy la mejor contando chistes, pero sí puedo aclarar la evidencia de DFDS. ¿Ruta, reviews o competidores?",
            "French": "Je ne suis pas la meilleure pour les blagues, mais je peux clarifier les preuves DFDS. Route, avis ou concurrents ?",
            "German": "Witze sind nicht meine Stärke, aber ich kann die DFDS-Belege klar erklären. Route, Bewertungen oder Wettbewerber?",
            "Danish": "Jeg er ikke bedst til jokes, men jeg kan gøre DFDS-evidensen klar. Rute, anmeldelser eller konkurrenter?"
        }
        return responses.get(language, "I’m better at DFDS evidence than jokes, but I can help with routes, reviews, or competitors.")

    if intent == "confusion":
        responses = {
            "Chinese": "我可以换一种说法。你想先看 routes、app reviews、competitors，还是让我直接总结这一页？",
            "Japanese": "別の言い方にできます。routes、app reviews、competitors のどれから見ますか？それともこのページを要約しましょうか？",
            "Korean": "다른 방식으로 설명할게요. routes, app reviews, competitors 중 무엇부터 볼까요? 아니면 이 페이지를 바로 요약할까요?",
            "Spanish": "Puedo decirlo de otra manera. ¿Quieres empezar por routes, app reviews o competitors, o prefieres un resumen de esta página?",
            "French": "Je peux le reformuler. Voulez-vous commencer par les routes, les app reviews ou les competitors, ou un résumé de cette page ?",
            "German": "Ich kann es anders formulieren. Willst du mit routes, app reviews oder competitors anfangen oder lieber eine Zusammenfassung dieser Seite?",
            "Danish": "Jeg kan sige det på en anden måde. Vil du starte med routes, app reviews eller competitors, eller skal jeg bare opsummere siden?"
        }
        return responses.get(language, "I can rephrase that. Want routes, app reviews, competitors, or a quick summary of this page?")

    responses = {
        "Chinese": "嗨，我是 Mia。我可以帮你看 DFDS 报告里的 routes、app reviews、competitors 和 recommendations。你想先看哪一块？",
        "Japanese": "こんにちは、Miaです。DFDSレポートの routes、app reviews、competitors、recommendations をお手伝いできます。どこから見ましょうか？",
        "Korean": "안녕하세요, Mia입니다. DFDS 보고서의 routes, app reviews, competitors, recommendations 를 도와드릴 수 있어요. 어디부터 볼까요?",
        "Spanish": "Hola, soy Mia. Puedo ayudarte con las routes, app reviews, competitors y recommendations del informe de DFDS. ¿Por dónde quieres empezar?",
        "French": "Bonjour, je suis Mia. Je peux vous aider avec les routes, les app reviews, les competitors et les recommendations du rapport DFDS. Par quoi voulez-vous commencer ?",
        "German": "Hallo, ich bin Mia. Ich kann dir bei den routes, app reviews, competitors und recommendations im DFDS-Bericht helfen. Womit sollen wir anfangen?",
        "Danish": "Hej, jeg er Mia. Jeg kan hjælpe med routes, app reviews, competitors og recommendations i DFDS-rapporten. Hvad vil du kigge på først?"
    }
    return responses.get(language, "Hi, I’m Mia. I can help with DFDS routes, app reviews, competitors, and recommendations. What would you like to look at first?")


def classify_conversation_intent(message):
    text = message.strip().lower()
    normalized = re.sub(r"^(mia|assistant|ai)[,:\s-]+", "", text, flags=re.I)
    normalized = re.sub(r"[,:\s-]+(mia|assistant|ai)$", "", normalized, flags=re.I).strip()
    normalized = re.sub(r"[,.!?。！？]+", " ", normalized).strip()
    normalized = re.sub(r"\s+", " ", normalized)
    if re.search(r"(weather|temperature|forecast|rain|raining|sunny|snow|wind|windy|storm|天气|气温|下雨|降雨|晴天|刮风|天气怎么样)", message, re.I):
        return "liveExternal"
    if re.search(r"^(bye|goodbye|see you|see ya|talk soon|farewell|later|good night|gn|再见|拜拜|回头见|晚安|明天见|下次聊)[!.。！\s]*$", normalized, re.I):
        return "goodbye"
    if re.search(r"^(hi|hello|hey|hiya|morning|good morning|good afternoon|good evening|你好|您好|嗨|早上好|下午好|晚上好)[!.。！\s]*$", normalized, re.I):
        return "smallTalk"
    if re.search(r"^(who are you|what are you|what can you do|what do you do|who is mia|help|can you help|what is this|what's this|who am i talking to|你是谁|你是什么|你能做什么|你会什么|这是什么|帮我一下)[?.!。\s]*$", normalized, re.I):
        return "capability"
    if re.search(r"^(how are you|how are you doing|how's it going|how is it going|are you okay|你好吗|最近怎么样|还好吗)[?.!。\s]*$", normalized, re.I):
        return "wellbeing"
    if re.search(r"^(thank you|thanks|thx|many thanks|appreciate it|谢谢|謝謝|多谢|多謝|感谢|感謝)[!.。！\s]*$", normalized, re.I):
        return "thanks"
    if re.search(
        r"^(ok|okay|k|got it|i got it|i understand|understood|makes sense|sounds good|cool|alright|all right|fine|yes|yep|yeah|sure|thanks|thank you|thx|ok thanks|okay thanks|got it thanks|i got it thanks|ok i got it thanks)(\s+(mia|assistant|ai))?$",
        normalized,
        re.I,
    ):
        return "thanks"
    if re.search(r"(好的|好|可以|明白|明白了|懂了|了解|收到|知道了|没问题|谢谢|謝謝|感谢|感謝)", message, re.I):
        return "thanks"
    if re.search(r"^(sorry|oops|my bad|apologies|pardon|excuse me|抱歉|不好意思|对不起|對不起)[!.。！\s]*$", normalized, re.I):
        return "apology"
    if re.search(r"^(tell me a joke|joke|make me laugh|say something funny|讲个笑话|说个笑话|來個笑話|來點幽默|說個笑話)[?.!。\s]*$", normalized, re.I):
        return "joke"
    if re.search(r"^(i don't understand|i do not understand|confused|what do you mean|can you explain|not sure i follow|什么意思|我不太懂|看不懂|你在说什么|你在說什麼)[?.!。\s]*$", normalized, re.I):
        return "confusion"
    if re.search(r"\b(dfds|report|dashboard|route|routes|dover|calais|newhaven|dieppe|newcastle|ijmuiden|jersey|app|review|reviews|competitor|competitors|recommendation|recommendations|source|sources|evidence|trustpilot|google play|apple|sentiment|customer|marketing|cx|ferry|ferries)\b", text, re.I):
        return "report"
    if re.search(r"(航线|路线|报告|仪表盘|竞品|竞争|建议|推荐|证据|来源|评分|评论|客户|渡轮|应用|登录|延误|根因)", message):
        return "report"
    return "offTopic"


def classify_conversation_mode(message):
    return "vertical" if classify_conversation_intent(message) == "report" else "smalltalk"


def build_messages(payload, rag_context):
    message = str(payload.get("message", ""))[:1200]
    language = str(payload.get("language", "English"))[:80]
    active_view = str(payload.get("activeView", "dashboard"))[:80]
    route_name = str(payload.get("routeName", "All signals"))[:120]
    conversation_mode = classify_conversation_mode(message)
    evidence_detail_requested = should_show_evidence_details(message)
    page_context = json_compact(rag_context.get("pageContext"), max_chars=2500)
    evidence_text = json.dumps(compact_evidence(rag_context.get("evidenceItems", [])), ensure_ascii=False, indent=2)
    insight_text = json.dumps(compact_insights(rag_context.get("insightItems", [])), ensure_ascii=False, indent=2)
    memory_text = json.dumps(compact_memories(rag_context.get("memoryItems", [])), ensure_ascii=False, indent=2)
    history_text = json.dumps(compact_history(rag_context.get("historyItems", [])), ensure_ascii=False, indent=2)

    system = (
        "You are the user-facing assistant for a DFDS Passenger Ferry Customer "
        "Intelligence Platform. Use the supplied dashboard page context, PostgreSQL "
        "Evidence Knowledge Base rows, historical insight items, recent chat history, "
        "and project memories together. Do not invent facts, scores, routes, or "
        "sources. If evidence is insufficient, say what is missing and what should be "
        "collected next. Keep the answer concise, business-friendly, and useful for a "
        "DFDS Marketing, CX, product, or leadership user. If the user's message is a greeting, small talk, or "
        "not about the DFDS report, respond with a warm bridge back to the report in "
        "the user's input language instead of forcing evidence. Reply only in the user's input language: "
        f"{language}. Do not add an English-first section, translation section, or any "
        f"language labels. Conversation mode for this request: {conversation_mode}. Use "
        "vertical mode for report questions and smalltalk mode for brief bridge replies. "
        "For report answers, start directly with the result in 1-2 plain-language sentences, "
        "then add 2-4 practical bullets only if they make the answer clearer. Do not prefix "
        "the answer with a canned summary label. Do not include visible internal labels for "
        "evidence cards, evidence trails, or source lists. "
        "Do not include a separate evidence section unless the user asks for it. If the user "
        "explicitly asks for evidence, sources, proof, citations, or supporting data, include a "
        "brief localized supporting-evidence section with source names, evidence IDs, and ratings "
        "or review counts when available."
    )
    user = (
        f"Dashboard page: {active_view}\n"
        f"Route focus: {route_name}\n"
        f"Evidence detail requested: {'yes' if evidence_detail_requested else 'no'}\n"
        f"User question: {message}\n\n"
        f"Current dashboard page context from the frontend:\n{page_context}\n\n"
        f"Recent short-term chat memory:\n{history_text}\n\n"
        f"Retrieved project memories:\n{memory_text}\n\n"
        f"Retrieved historical business insights:\n{insight_text}\n\n"
        f"Retrieved public evidence from PostgreSQL + pgvector:\n{evidence_text}"
    )
    return [{"role": "system", "content": system}, {"role": "user", "content": user}]


def call_deepseek(payload, rag_context):
    api_key = os.environ.get("DEEPSEEK_API_KEY")
    evidence_items = rag_context.get("evidenceItems", [])
    intent = classify_conversation_intent(str(payload.get("message", "")))
    mode = classify_conversation_mode(str(payload.get("message", "")))
    evidence_detail_requested = should_show_evidence_details(str(payload.get("message", "")))
    if mode != "vertical":
        return 200, {
            "mode": mode,
            "intent": intent,
            "answer": bridge_answer(payload),
            "evidence": [],
            "insights": compact_insights(rag_context.get("insightItems", [])),
            "memories": compact_memories(rag_context.get("memoryItems", [])),
        }
    if not evidence_items:
        return 200, {
            "mode": mode,
            "intent": intent,
            "answer": bridge_answer(payload),
            "evidence": [],
            "insights": compact_insights(rag_context.get("insightItems", [])),
            "memories": compact_memories(rag_context.get("memoryItems", [])),
        }
    if not api_key:
        answer = evidence_only_answer(
            payload,
            evidence_items,
            "DeepSeek is not connected because DEEPSEEK_API_KEY is not set.",
        )
        return 200, {
            "mode": mode,
            "intent": intent,
            "answer": answer,
            "evidence": compact_evidence(evidence_items),
            "insights": compact_insights(rag_context.get("insightItems", [])),
            "memories": compact_memories(rag_context.get("memoryItems", [])),
        }

    request_body = {
        "model": os.environ.get("DEEPSEEK_MODEL", "deepseek-chat"),
        "messages": build_messages(payload, rag_context),
        "temperature": 0.2,
        "max_tokens": 900,
    }
    data = json.dumps(request_body).encode("utf-8")
    request = urllib.request.Request(
        DEEPSEEK_URL,
        data=data,
        method="POST",
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
    )

    try:
        with urllib.request.urlopen(request, timeout=35, context=HTTPS_CONTEXT) as response:
            response_data = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")[:500]
        return exc.code, {"error": f"DeepSeek API error: {detail}"}
    except Exception as exc:
        return 502, {"error": f"DeepSeek request failed: {exc}"}

    answer = (
        response_data.get("choices", [{}])[0]
        .get("message", {})
        .get("content", "")
        .strip()
    )
    answer = strip_answer_labels(answer, str(payload.get("message", "")))
    return 200, {
        "mode": mode,
        "intent": intent,
        "answer": answer or "No answer returned by DeepSeek.",
        "evidence": compact_evidence(evidence_items),
        "insights": compact_insights(rag_context.get("insightItems", [])),
        "memories": compact_memories(rag_context.get("memoryItems", [])),
        "providerUsage": response_data.get("usage", {}),
    }


class ServerConversationServices:
    """Adapter which exposes existing deterministic server capabilities as graph nodes."""

    def __init__(self, payload, identity):
        self.payload = payload
        self.identity = identity

    def route_request(self, state):
        return {"route": classify_conversation_mode(state["message"])}

    def retrieve_evidence(self, state):
        rag_context = build_rag_context(self.payload)
        safe_evidence = []
        for item in rag_context.get("evidenceItems", []):
            # Evidence is untrusted data: suspicious instruction-like content is excluded.
            if not assess_untrusted_text(json.dumps(item, ensure_ascii=False)):
                safe_evidence.append(item)
        rag_context["evidenceItems"] = safe_evidence
        return {"rag_context": rag_context, "evidence_ids": [str(item.get("id")) for item in safe_evidence if item.get("id") is not None]}

    def generate_answer(self, state):
        status, result = call_deepseek(self.payload, state["rag_context"])
        answer = {
            "text": result.get("answer", ""),
            "evidence_ids": state.get("evidence_ids", []),
            "confidence": 0.96 if state.get("evidence_ids") else 0.70,
            "recommendation": result.get("recommendation", ""),
            "suspected_injection": False,
        }
        return {"status": status, "result": result, "answer": answer}

    def validate_draft(self, state):
        return {}

    def queue_review(self, state):
        return {"publication_state": "pending_human_review"}


class ServerEvaluationServices:
    """Adapter for the existing retrieval and answer path; labels never enter these methods."""

    def retrieve_case(self, case):
        started = time.perf_counter()
        payload = {"message": case["question"], "language": case.get("language", "English"), "routeKey": case.get("route", "all"), "activeView": "evaluation"}
        rag_context = build_rag_context(payload)
        return {
            "retrieved_evidence_ids": [str(item.get("id")) for item in rag_context.get("evidenceItems", []) if item.get("id") is not None][:5],
            "rag_context": rag_context,
            "latency_ms": round((time.perf_counter() - started) * 1000, 1),
        }

    def answer_case(self, case, retrieval):
        started = time.perf_counter()
        payload = {"message": case["question"], "language": case.get("language", "English"), "routeKey": case.get("route", "all"), "activeView": "evaluation"}
        status, response = call_deepseek(payload, retrieval.get("rag_context") or {})
        return {
            "answer": response.get("answer", ""),
            "failed": status != 200,
            "error_code": None if status == 200 else str(status),
            "model_usage": response.get("providerUsage") or {},
            "latency_ms": round((time.perf_counter() - started) * 1000, 1),
        }

    def judge_case(self, result, labels):
        api_key = os.environ.get("DEEPSEEK_API_KEY")
        if not api_key:
            raise RuntimeError("DEEPSEEK_API_KEY is required for the evaluation LLM judge")
        expected = {str(value) for value in labels.get("expected_evidence_ids") or []}
        retrieved = {str(value) for value in result.get("retrieved_evidence_ids") or []}
        request_body = {
            "model": os.environ.get("MIA_EVALUATION_JUDGE_MODEL") or os.environ.get("DEEPSEEK_MODEL", "deepseek-chat"),
            "temperature": 0,
            "max_tokens": 120,
            "messages": [
                {"role": "system", "content": "You are a strict evaluation judge. Return only JSON with score (0..1), citation_correct (boolean), and rubric_version answer-quality-v1."},
                {"role": "user", "content": json.dumps({"answer": result.get("answer", ""), "retrieved_evidence_ids": sorted(retrieved), "expected_answer": labels.get("expected_answer", ""), "expected_evidence_ids": sorted(expected), "expects_abstention": bool(labels.get("expects_abstention"))}, ensure_ascii=False)},
            ],
        }
        request = urllib.request.Request(DEEPSEEK_URL, data=json.dumps(request_body).encode("utf-8"), method="POST", headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"})
        with urllib.request.urlopen(request, timeout=35, context=HTTPS_CONTEXT) as response:
            payload = json.loads(response.read().decode("utf-8"))
        content = payload.get("choices", [{}])[0].get("message", {}).get("content", "")
        verdict = parse_judge_verdict(content)
        verdict["model_usage"] = payload.get("usage") or {}
        return verdict


def build_it_data_flow_review_context(manifest, cleaned_sheets=None):
    fields = []
    evidence = []
    dimension_profiles = []
    dataset_sources = []
    dimension_values = {"countries": [], "routes": [], "languages": []}
    observed_dates = []
    worksheet_count = 0
    profiled_rows = 0
    cleaned_sheets = cleaned_sheets or [
        {"fileName": file_entry.get("name"), "sheetName": sheet.get("name"), **sheet}
        for file_entry in manifest.get("files", []) if file_entry.get("parserState") == "profiled"
        for sheet in file_entry.get("sheets", [])
    ]
    for sheet in cleaned_sheets:
        worksheet_count += 1
        profiled_rows += int(sheet.get("cleaning", {}).get("cleanedRows", 0))
        sheet_name = str(sheet.get("sheetName", "")).strip()
        if sheet_name and sheet_name not in dataset_sources:
            dataset_sources.append(sheet_name)
        columns = sheet.get("columns", [])[:30]
        fields.extend(column for column in columns if column not in fields)
        reference = f"{sheet.get('fileName', 'upload')} / {sheet.get('sheetName', 'sheet')}"
        for column in columns:
            if RESTRICTED_SOURCE_FIELD.search(column):
                continue
            values = [
                str(row.get(column)).strip()
                for row in sheet.get("previewRows", [])
                if row.get(column) is not None and str(row.get(column)).strip()
            ]
            column_key = re.sub(r"[^a-z]", "", column.lower())
            if re.search(r"date|time|travel", column_key):
                for value in values:
                    iso_match = re.search(r"\b(20\d{2})[-/](\d{1,2})[-/](\d{1,2})\b", value)
                    dmy_match = re.search(r"\b(\d{1,2})[-/](\d{1,2})[-/](20\d{2})\b", value)
                    if iso_match:
                        observed_dates.append((int(iso_match.group(1)), int(iso_match.group(2)), int(iso_match.group(3))))
                    elif dmy_match:
                        observed_dates.append((int(dmy_match.group(3)), int(dmy_match.group(2)), int(dmy_match.group(1))))
            frequencies = {}
            for value in values:
                frequencies[value] = frequencies.get(value, 0) + 1
            if not frequencies or len(frequencies) > 50:
                continue
            dimension_profiles.append({
                "field": column,
                "reference": reference,
                "observedValues": len(frequencies),
                "topValues": [{"value": value, "count": count} for value, count in sorted(frequencies.items(), key=lambda item: (-item[1], item[0]))[:8]],
            })
            target_dimension = "countries" if re.search(r"country|market", column_key) else "routes" if "route" in column_key else "languages" if re.search(r"language|locale|lang", column_key) else None
            if target_dimension:
                for value in sorted(frequencies, key=lambda item: (-frequencies[item], item)):
                    if value not in dimension_values[target_dimension] and len(dimension_values[target_dimension]) < 8:
                        dimension_values[target_dimension].append(value)
        evidence.append({"reference": reference, "columns": columns, "rowCount": int(sheet.get("cleaning", {}).get("cleanedRows", 0))})
    valid_dates = sorted({date for date in observed_dates if 1 <= date[1] <= 12 and 1 <= date[2] <= 31})
    time_period = "Not identified"
    if valid_dates:
        start, end = valid_dates[0], valid_dates[-1]
        format_date = lambda value: f"{value[0]:04d}-{value[1]:02d}-{value[2]:02d}"
        time_period = format_date(start) if start == end else f"{format_date(start)} to {format_date(end)}"
    return {
        "fileOverview": {
            "fileCount": int(manifest.get("batch", {}).get("fileCount", 0)),
            "worksheetCount": worksheet_count,
            "profiledRows": profiled_rows,
            "fields": fields[:40],
        },
        "evidence": evidence[:12],
        "dimensionProfiles": dimension_profiles[:40],
        "datasetOverview": {
            "recordCount": profiled_rows,
            "dataSources": dataset_sources[:8],
            "timePeriod": time_period,
            "countries": dimension_values["countries"],
            "routes": dimension_values["routes"],
            "languages": dimension_values["languages"],
        },
    }


def uploaded_batch_chat_context(batch_id, identity):
    manifest = require_owned_upload_batch(batch_id, identity)
    _, cleaned_sheets = _cleaned_batch_sheets(UPLOAD_STORAGE_DIR, batch_id)
    context = build_it_data_flow_review_context(manifest, cleaned_sheets)
    quality_summary = {
        key: sum(int((sheet.get("cleaning") or {}).get(key) or 0) for sheet in cleaned_sheets)
        for key in ("receivedRows", "cleanedRows", "blankRowsRemoved", "duplicateRowsRemoved", "missingValuesStandardized")
    }
    return {
        "batchId": str(manifest["batch"]["id"]),
        "version": str(manifest["batch"].get("version") or ""),
        "fileOverview": {
            "fileCount": int(context["fileOverview"]["fileCount"]),
            "worksheetCount": int(context["fileOverview"]["worksheetCount"]),
            "profiledRows": int(context["fileOverview"]["profiledRows"]),
        },
        "datasetOverview": context["datasetOverview"],
        "availableFields": [field for field in context["fileOverview"]["fields"] if not RESTRICTED_SOURCE_FIELD.search(field)][:24],
        "dimensionProfiles": [
            profile for profile in context["dimensionProfiles"]
            if not RESTRICTED_SOURCE_FIELD.search(str(profile.get("field") or ""))
        ][:16],
        "qualitySummary": quality_summary,
    }


def normalize_it_data_flow_review(payload):
    payload = payload if isinstance(payload, dict) else {}
    topics = payload.get("mainBusinessTopics", {}) if isinstance(payload.get("mainBusinessTopics"), dict) else {}
    def topic_items(key):
        values = topics.get(key, []) if isinstance(topics.get(key), list) else []
        return [str(value).strip()[:120] for value in values[:6] if str(value).strip()]
    insights = []
    for item in payload.get("keyInsights", [])[:6]:
        if not isinstance(item, dict) or not str(item.get("title", "")).strip() or not str(item.get("detail", "")).strip():
            continue
        insights.append({
            "title": str(item["title"]).strip()[:120],
            "detail": str(item["detail"]).strip()[:360],
            "category": str(item.get("category", "business")).strip()[:40] or "business",
        })
    return {
        "datasetSummary": str(payload.get("datasetSummary", "No AI summary was returned.")).strip()[:900],
        "mainBusinessTopics": {
            "customerThemes": topic_items("customerThemes"),
            "serviceAreas": topic_items("serviceAreas"),
            "operations": topic_items("operations"),
        },
        "keyInsights": insights,
    }


IT_DATA_FLOW_REVIEW_SYSTEM_PROMPT = """You create an executive business briefing from one newly uploaded and cleaned dataset. You receive only facts and aggregate distributions calculated for that current batch. Write datasetSummary as 3 to 5 concise sentences, no more than 120 words. Explain the business processes represented, notable coverage or key insight, the overall data quality using only supplied aggregate counts, and end with a clear readiness recommendation. Use plain business language for commercial, operations, CX, and marketing leaders. Do not list columns, discuss parsing, PII, system architecture, or invent facts. If a quality fact is not supplied, do not claim it. Return JSON only with this exact structure: {\"datasetSummary\": \"3-5 sentence executive briefing, max 120 words\", \"mainBusinessTopics\": {\"customerThemes\": [\"up to 4 concise themes\"], \"serviceAreas\": [\"up to 4 concise service areas\"], \"operations\": [\"up to 4 concise operations\"]}, \"keyInsights\": [{\"title\": \"short business insight\", \"detail\": \"one concise evidence-based sentence\", \"category\": \"customer|route|sentiment|service|trend|anomaly\"}]}. Return 4 to 6 keyInsights, but omit unsupported insights."""


def run_it_data_flow_review(batch_id):
    manifest, cleaned_sheets = _cleaned_batch_sheets(UPLOAD_STORAGE_DIR, batch_id)
    context = build_it_data_flow_review_context(manifest, cleaned_sheets)
    if not context["evidence"]:
        return 422, {"error": "This batch has no profiled spreadsheet data for AI review"}
    api_key = os.environ.get("DEEPSEEK_API_KEY")
    if not api_key:
        return 200, {"status": "unavailable", "fileOverview": context["fileOverview"], "datasetOverview": context["datasetOverview"], "review": normalize_it_data_flow_review({"datasetSummary": "AI business summary is unavailable for this upload."})}
    request_body = {"model": os.environ.get("DEEPSEEK_MODEL", "deepseek-chat"), "messages": [{"role": "system", "content": IT_DATA_FLOW_REVIEW_SYSTEM_PROMPT}, {"role": "user", "content": json.dumps(context, ensure_ascii=False)}], "temperature": 0.1, "max_tokens": 1000, "response_format": {"type": "json_object"}}
    request = urllib.request.Request(DEEPSEEK_URL, data=json.dumps(request_body).encode("utf-8"), method="POST", headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(request, timeout=35, context=HTTPS_CONTEXT) as response:
            response_data = json.loads(response.read().decode("utf-8"))
        content = response_data.get("choices", [{}])[0].get("message", {}).get("content", "{}")
        review = normalize_it_data_flow_review(json.loads(content))
    except (urllib.error.HTTPError, urllib.error.URLError, json.JSONDecodeError) as exc:
        return 502, {"error": f"DeepSeek review failed: {exc}"}
    return 200, {"status": "complete", "fileOverview": context["fileOverview"], "datasetOverview": context["datasetOverview"], "review": review}


def valid_uuid(value):
    try:
        return str(uuid.UUID(str(value)))
    except (ValueError, TypeError):
        return str(uuid.uuid4())


def access_token_hash(token):
    return hashlib.sha256(str(token).encode("utf-8")).hexdigest()


def create_account(username, password, role="project_user"):
    normalized = validate_registration_username(username)
    if role not in {"project_user", "administrator"}:
        raise ValueError("Invalid role")
    conn = connect_db()
    if conn is None:
        raise RuntimeError("Database is unavailable")
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT id, username, role FROM app_users WHERE username = %s", (normalized,))
            existing = cur.fetchone()
            if existing:
                return None
            user_id = str(uuid.uuid4())
            cur.execute(
                """
                INSERT INTO app_users (id, username, password_hash, role)
                VALUES (%s, %s, %s, %s)
                """,
                (user_id, normalized, hash_password(password), role),
            )
        conn.commit()
        return {"id": user_id, "username": normalized, "role": role}
    finally:
        conn.close()


def authenticate_account(username, password):
    try:
        normalized = normalize_username(username)
    except ValueError:
        return None
    conn = connect_db()
    if conn is None:
        return None
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT id, username, password_hash, role FROM app_users WHERE username = %s", (normalized,))
            user = cur.fetchone()
    finally:
        conn.close()
    if not user or not verify_password(str(password or ""), user["password_hash"]):
        return None
    return {"id": str(user["id"]), "username": user["username"], "role": user["role"]}


def create_access_session(user):
    token = secrets.token_urlsafe(32)
    conn = connect_db()
    if conn is None:
        raise RuntimeError("Database is unavailable")
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO user_access_sessions (token_hash, user_id, expires_at)
                VALUES (%s, %s, NOW() + INTERVAL '12 hours')
                """,
                (access_token_hash(token), user["id"]),
            )
        conn.commit()
    finally:
        conn.close()
    return token


def identity_from_authorization(value):
    prefix = "Bearer "
    if not str(value or "").startswith(prefix):
        return None
    token = str(value)[len(prefix):]
    public_workspace_identities = {
        "public-project-user": {"id": "public-project-user", "username": "Project user", "role": "project_user"},
        "public-administrator": {"id": "public-administrator", "username": "Administrator", "role": "administrator"},
    }
    if token in public_workspace_identities:
        return public_workspace_identities[token]
    conn = connect_db()
    if conn is None:
        return None
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT u.id, u.username, u.role
                FROM user_access_sessions s
                JOIN app_users u ON u.id = s.user_id
                WHERE s.token_hash = %s AND s.expires_at > NOW()
                """,
                (access_token_hash(token),),
            )
            user = cur.fetchone()
            if user:
                cur.execute("UPDATE user_access_sessions SET last_seen_at = NOW() WHERE token_hash = %s", (access_token_hash(token),))
        conn.commit()
    finally:
        conn.close()
    if not user:
        return None
    return {"id": str(user["id"]), "username": user["username"], "role": user["role"]}


def has_capability(identity, capability):
    """Keep authorization server-side; the current administrator role owns governance capabilities."""
    if not identity:
        return False
    if identity.get("role") == "administrator":
        return capability in {"governance_admin", "data_admin", "reviewer", "evaluation_admin"}
    return capability == "project_user"


def require_production_ready():
    if PRODUCTION_MODE_ENABLED and readiness_report().get("status") != "ready":
        raise PermissionError("Production governance readiness is blocked")


def production_readiness():
    return readiness_report()


def enable_production_mode(identity):
    global PRODUCTION_MODE_ENABLED
    if not has_capability(identity, "governance_admin"):
        raise PermissionError("Governance administrator access is required")
    report = production_readiness()
    if report["status"] != "ready":
        raise RuntimeError("Production governance readiness is blocked")
    if not PRODUCTION_MODE_ENABLED:
        PRODUCTION_MODE_ENABLED = True
        GOVERNANCE_AUDIT_EVENTS.append({
            "actor": pseudonymous_actor_id(identity.get("id")),
            "action": "production_mode_enabled",
            "governance_version": report["governance_version"],
            "declaration_hash": report["declaration_hash"],
            "status": "enabled",
        })
        record_governance_audit(GOVERNANCE_AUDIT_EVENTS[-1])
    return {"status": "enabled", "idempotent": True, "governance": report}


def measure_usage(payload, answer, rag_context, provider_usage=None):
    provider_usage = provider_usage or {}
    input_tokens = int(provider_usage.get("prompt_tokens") or max(1, math.ceil(len(json.dumps(payload, ensure_ascii=False)) / 4)))
    output_tokens = int(provider_usage.get("completion_tokens") or max(1, math.ceil(len(str(answer or "")) / 4)))
    return {
        "inputTokens": input_tokens,
        "outputTokens": output_tokens,
        "totalTokens": int(provider_usage.get("total_tokens") or (input_tokens + output_tokens)),
        "retrieval": rag_context.get("retrieval", {}),
    }


def record_rag_usage(identity, session_id, answer, rag_context, usage):
    conn = connect_db()
    if conn is None:
        return
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO rag_usage_events (id, user_id, chat_session_id, retrieval_trace, input_tokens, output_tokens, total_tokens)
                VALUES (%s, %s, %s, %s::jsonb, %s, %s, %s)
                """,
                (str(uuid.uuid4()), identity["id"], session_id, json.dumps(rag_context.get("retrieval", {})), usage["inputTokens"], usage["outputTokens"], usage["totalTokens"]),
            )
            cur.execute(
                """
                INSERT INTO user_summary_history (id, user_id, chat_session_id, summary_text)
                VALUES (%s, %s, %s, %s)
                """,
                (str(uuid.uuid4()), identity["id"], session_id, str(answer or "")[:900]),
            )
        conn.commit()
    finally:
        conn.close()


def record_graph_audit(record):
    """Persist only the allow-listed redacted audit record; failure never exposes request data."""
    conn = connect_db()
    if conn is None:
        return
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO graph_run_audits (graph_name, graph_version, idempotency_key, status, trace_id, metadata)
                VALUES (%s, %s, %s, %s, %s, %s::jsonb)
                """,
                (
                    record.get("graph", "unknown"), record.get("graph_version", "unknown"), record.get("idempotency_key"),
                    record.get("status", "unknown"), record.get("trace_id"), json.dumps(record),
                ),
            )
        conn.commit()
    except Exception:
        conn.rollback()
    finally:
        conn.close()


def record_governance_audit(record):
    """Persist only minimum governance metadata; dev mode remains usable without PostgreSQL."""
    conn = connect_db()
    if conn is None:
        return
    try:
        with conn.cursor() as cur:
            cur.execute(
                """INSERT INTO governance_audit_events
                   (actor_hash, action, governance_version, declaration_hash, status)
                   VALUES (%s, %s, %s, %s, %s)""",
                (record["actor"], record["action"], record["governance_version"], record["declaration_hash"], record["status"]),
            )
        conn.commit()
    except Exception:
        conn.rollback()
    finally:
        conn.close()


def _evaluation_label_hash(case):
    labels = {
        "expected_answer": str(case.get("expected_answer") or ""),
        "expected_evidence_ids": list(case.get("expected_evidence_ids") or []),
        "expects_abstention": bool(case.get("expects_abstention")),
    }
    return hashlib.sha256(json.dumps(labels, sort_keys=True).encode("utf-8")).hexdigest()


def freeze_evaluation_set(payload, identity):
    """Create an immutable, restricted evaluation-set version from administrator-supplied cases."""
    cases = payload.get("cases")
    version = str(payload.get("version") or "").strip()
    if not version or not isinstance(cases, list) or not cases:
        return 400, {"error": "version and at least one evaluation case are required"}
    if any(not isinstance(case, dict) or not str(case.get("question") or "").strip() for case in cases):
        return 400, {"error": "every evaluation case requires a question"}
    conn = connect_db()
    if conn is None:
        return 503, {"error": "Database is required for frozen evaluation sets"}
    source_hash = hashlib.sha256(json.dumps(cases, sort_keys=True).encode("utf-8")).hexdigest()
    try:
        with conn.cursor() as cur:
            cur.execute("""INSERT INTO mia_evaluation_set_versions (version, status, source_snapshot_hash, configuration, created_by, frozen_at)
                           VALUES (%s, 'frozen', %s, %s::jsonb, %s, NOW()) RETURNING id""",
                        (version, source_hash, json.dumps(payload.get("configuration") or {}), identity["id"]))
            set_id = str(cur.fetchone()["id"])
            for index, case in enumerate(cases, start=1):
                cur.execute("""INSERT INTO mia_evaluation_cases
                    (evaluation_set_id, case_key, question, language, route_key, source_type, expected_answer, expected_evidence_ids, expects_abstention, label_hash)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s::jsonb, %s, %s)""",
                    (set_id, str(case.get("case_key") or f"case-{index}"), str(case["question"]), str(case.get("language") or "English"),
                     str(case.get("route") or "all"), str(case.get("source_type") or "evidence"), str(case.get("expected_answer") or ""),
                     json.dumps(case.get("expected_evidence_ids") or []), bool(case.get("expects_abstention")), _evaluation_label_hash(case)))
        conn.commit()
        return 201, {"evaluationSetId": set_id, "version": version, "caseCount": len(cases), "status": "frozen"}
    except Exception as exc:
        conn.rollback()
        return 409, {"error": f"Unable to freeze evaluation set: {exc}"}
    finally:
        conn.close()


def load_frozen_evaluation_set(set_id=None):
    conn = connect_db()
    if conn is None:
        raise RuntimeError("Database is required for evaluation runs")
    run_id = None
    try:
        with conn.cursor() as cur:
            if set_id:
                cur.execute("SELECT id, version FROM mia_evaluation_set_versions WHERE id = %s AND status = 'frozen'", (set_id,))
            else:
                cur.execute("SELECT id, version FROM mia_evaluation_set_versions WHERE status = 'frozen' ORDER BY frozen_at DESC LIMIT 1")
            evaluation_set = cur.fetchone()
            if evaluation_set is None:
                raise ValueError("No frozen evaluation set is available")
            cur.execute("""SELECT id, question, language, route_key, source_type, expected_answer, expected_evidence_ids, expects_abstention, label_hash
                           FROM mia_evaluation_cases WHERE evaluation_set_id = %s ORDER BY case_key""", (evaluation_set["id"],))
            return dict(evaluation_set), [dict(row) for row in cur.fetchall()]
    finally:
        conn.close()


def run_frozen_evaluation(set_id, identity):
    """Execute one frozen set and persist only redacted outcomes and aggregate metrics."""
    evaluation_set, stored_cases = load_frozen_evaluation_set(set_id)
    cases = [{
        "question": row["question"], "language": row["language"], "route": row["route_key"], "source_type": row["source_type"],
        "expected_answer": row["expected_answer"], "expected_evidence_ids": row["expected_evidence_ids"], "expects_abstention": row["expects_abstention"],
    } for row in stored_cases]
    conn = connect_db()
    if conn is None:
        raise RuntimeError("Database is required for evaluation runs")
    try:
        with conn.cursor() as cur:
            cur.execute("""SELECT r.id, r.metrics FROM mia_evaluation_baselines b
                           JOIN mia_evaluation_runs r ON r.id = b.run_id ORDER BY b.created_at DESC LIMIT 1""")
            baseline = cur.fetchone()
            baseline_metrics = dict(baseline["metrics"]) if baseline else {}
            cur.execute("""INSERT INTO mia_evaluation_runs (evaluation_set_id, graph_version, configuration, status, baseline_run_id, started_by)
                           VALUES (%s, %s, %s::jsonb, 'running', %s, %s) RETURNING id""",
                        (evaluation_set["id"], "mia-graph-v1", json.dumps({"judge_rubric_version": "citation-grounded-v1"}), baseline["id"] if baseline else None, identity["id"]))
            run_id = str(cur.fetchone()["id"])
        conn.commit()
        prepared_evaluation = build_evaluation_run(cases)
        graph = build_evaluation_graph(checkpointer=MIA_GRAPH_CHECKPOINTER, services=ServerEvaluationServices(), label_vault=prepared_evaluation["label_vault"])
        trace_state = {
            "graph": "evaluation", "graph_version": "mia-graph-v1", "request_id": run_id,
            "retrieval_version": "hybrid-retrieval-v1", "judge_rubric_version": "citation-grounded-v1",
            "prompt_version": "evaluation-judge-v1", "config_version": "mia-config-v1",
        }
        with RedactedTracer(configured_langfuse_client()).span("mia.evaluation.run", trace_state) as span:
            state = graph.invoke({"cases": prepared_evaluation["live_retrieval_payload"]["cases"], "baseline_metrics": baseline_metrics, "trace_id": span.trace_id}, config={"configurable": {"thread_id": f"evaluation:{run_id}"}})
        trace_state.update({
            "trace_id": span.trace_id,
            "status": state["promotion_state"],
            "scores": state["scored_metrics"],
            "verdict": state["reasons"],
            "validation_verdict": state["reasons"],
            "final_outcome": state["promotion_state"],
        })
        record_graph_result_trace(configured_langfuse_client(), name="mia.evaluation.result", state=trace_state, trace_id=span.trace_id)
        trace = graph_audit_record({"graph": "evaluation", "graph_version": "mia-graph-v1", "request_id": run_id, "trace_id": span.trace_id, "status": state["promotion_state"], "scores": state["scored_metrics"], "verdict": state["reasons"], "judge_rubric_version": "citation-grounded-v1"})
        record_graph_audit(trace)
        with conn.cursor() as cur:
            for source, result in zip(stored_cases, state.get("case_results") or []):
                cur.execute("""INSERT INTO mia_evaluation_case_results
                  (run_id, case_id, label_hash, retrieved_evidence_ids, citation_correct, judge_score, abstained, latency_ms, failed, error_code, model_usage, trace_id)
                  VALUES (%s, %s, %s, %s::jsonb, %s, %s, %s, %s, %s, %s, %s::jsonb, %s)""",
                  (run_id, source["id"], source["label_hash"], json.dumps(result.get("retrieved_evidence_ids") or []), result.get("citation_correct"), result.get("judge_score"),
                   not bool(result.get("retrieved_evidence_ids")), result.get("latency_ms"), bool(result.get("failed")), result.get("error_code"), json.dumps(result.get("model_usage") or {}), trace["trace_id"] if trace.get("trace_id") else None))
            cur.execute("""INSERT INTO mia_evaluation_metric_segments (run_id, segment_type, segment_value, metrics) VALUES (%s, 'overall', 'all', %s::jsonb)""", (run_id, json.dumps(state["scored_metrics"])))
            for key, metrics in (state.get("metric_segments") or {}).items():
                segment_type, segment_value = key.split(":", 1)
                cur.execute("""INSERT INTO mia_evaluation_metric_segments (run_id, segment_type, segment_value, metrics) VALUES (%s, %s, %s, %s::jsonb)""", (run_id, segment_type, segment_value, json.dumps(metrics)))
            cur.execute("""UPDATE mia_evaluation_runs SET status = %s, promotion_state = %s, reasons = %s::jsonb, trace_id = %s, metrics = %s::jsonb, completed_at = NOW() WHERE id = %s""",
                        (state["promotion_state"], state["promotion_state"], json.dumps(state["reasons"]), trace.get("trace_id"), json.dumps(state["scored_metrics"]), run_id))
        conn.commit()
        return {"id": run_id, "evaluationSetVersion": evaluation_set["version"], "status": state["promotion_state"], "promotionState": state["promotion_state"], "reasons": state["reasons"], "metrics": state["scored_metrics"], "segments": state.get("metric_segments") or {}, "traceId": trace.get("trace_id")}
    except Exception:
        conn.rollback()
        if run_id:
            with conn.cursor() as cur:
                cur.execute("UPDATE mia_evaluation_runs SET status = 'failed', promotion_state = 'blocked', reasons = %s::jsonb, completed_at = NOW() WHERE id = %s", (json.dumps(["execution_failed"]), run_id))
            conn.commit()
        raise
    finally:
        conn.close()


def _scheduled_dataset_graph(*, snapshot, run_id, graph_version, evidence_ids, trace_id=None, **_ignored):
    return run_dataset_graph(snapshot, run_id=run_id, graph_version=graph_version, evidence_ids=evidence_ids, trace_id=trace_id, checkpointer=MIA_GRAPH_CHECKPOINTER)


def _scheduled_evaluation_graph(*, trigger_event, **_ignored):
    set_id = trigger_event.get("evaluationSetId") or trigger_event.get("evaluation_set_id")
    if not set_id:
        raise ValueError("version change event requires evaluationSetId")
    return run_frozen_evaluation(set_id, {"id": "system-version-trigger"})


def mia_trigger_scheduler():
    global MIA_TRIGGER_SCHEDULER
    if MIA_TRIGGER_SCHEDULER is None:
        MIA_TRIGGER_SCHEDULER = MiaTriggerScheduler(run_dataset=_scheduled_dataset_graph, run_evaluation=_scheduled_evaluation_graph)
    return MIA_TRIGGER_SCHEDULER


def process_external_refresh_event(event):
    """Production worker entry point after an approved external refresh."""
    return mia_trigger_scheduler().poll_external_refresh(dict(event))


def process_component_version_event(event):
    """Production release-hook entry point for a governed version change."""
    return mia_trigger_scheduler().poll_version_change(dict(event))


def set_evaluation_baseline(payload, identity):
    run_id = str(payload.get("runId") or "").strip()
    name = str(payload.get("name") or "promotion-default")[:120]
    if not run_id:
        return 400, {"error": "runId is required"}
    conn = connect_db()
    if conn is None:
        return 503, {"error": "Database is required for evaluation baselines"}
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM mia_evaluation_runs WHERE id = %s AND status IN ('blocked', 'accepted')", (run_id,))
            if cur.fetchone() is None:
                return 404, {"error": "Completed evaluation run not found"}
            cur.execute("""INSERT INTO mia_evaluation_baselines (name, run_id, threshold_version, created_by)
                           VALUES (%s, %s, 'evaluation-gate-v1', %s)
                           ON CONFLICT (name) DO UPDATE SET run_id = EXCLUDED.run_id, threshold_version = EXCLUDED.threshold_version, created_by = EXCLUDED.created_by, created_at = NOW()""",
                        (name, run_id, identity["id"]))
        conn.commit()
        return 200, {"name": name, "runId": run_id, "thresholdVersion": "evaluation-gate-v1"}
    finally:
        conn.close()


def admin_evaluation_detail(run_id):
    """Return one evaluation run with redacted case, segment, and trace details."""
    conn = connect_db()
    if conn is None:
        return 503, {"error": "Database is required for evaluation details"}
    try:
        with conn.cursor() as cur:
            cur.execute("""SELECT r.id, s.version AS evaluation_set_version, r.status, r.promotion_state, r.reasons,
                                  r.metrics, r.trace_id, r.graph_version, r.configuration, r.started_at, r.completed_at
                           FROM mia_evaluation_runs r JOIN mia_evaluation_set_versions s ON s.id = r.evaluation_set_id
                           WHERE r.id = %s""", (run_id,))
            run = cur.fetchone()
            if run is None:
                return 404, {"error": "Evaluation run not found"}
            cur.execute("""SELECT result.case_id, evaluation_case.case_key, result.retrieved_evidence_ids,
                                  result.citation_correct, result.judge_score, result.abstained, result.latency_ms,
                                  result.failed, result.error_code, result.model_usage, result.trace_id,
                                  result.human_scores, result.human_overall_score, result.human_citation_correct,
                                  result.human_should_abstain, result.human_policy_issue, result.human_notes,
                                  result.human_reviewer_id, result.human_reviewed_at
                           FROM mia_evaluation_case_results result
                           JOIN mia_evaluation_cases evaluation_case ON evaluation_case.id = result.case_id
                           WHERE result.run_id = %s ORDER BY evaluation_case.case_key""", (run_id,))
            cases = [dict(row) for row in cur.fetchall()]
            cur.execute("""SELECT segment_type, segment_value, metrics
                           FROM mia_evaluation_metric_segments WHERE run_id = %s
                           ORDER BY segment_type, segment_value""", (run_id,))
            segments = [dict(row) for row in cur.fetchall()]
        return 200, {"run": dict(run), "caseResults": cases, "segments": segments, "trace": {"traceId": run.get("trace_id"), "graphVersion": run.get("graph_version"), "configuration": run.get("configuration") or {}}}
    finally:
        conn.close()


def record_human_calibration(run_id, payload, identity):
    """Persist reviewer scores against a completed run without exposing frozen labels."""
    records = payload.get("records") if isinstance(payload, dict) else None
    if not isinstance(records, list) or not records:
        return 400, {"error": "records must be a non-empty list"}
    try:
        normalized = [validate_human_calibration_record(dict(record)) for record in records]
    except (TypeError, ValueError) as exc:
        return 400, {"error": str(exc)}
    reviewer_id = str(identity.get("id") or "")
    if not reviewer_id:
        return 403, {"error": "Reviewer identity is required"}
    conn = connect_db()
    if conn is None:
        return 503, {"error": "Database is required for human calibration"}
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT id, status FROM mia_evaluation_runs WHERE id = %s AND status IN ('accepted', 'blocked')", (run_id,))
            if cur.fetchone() is None:
                return 404, {"error": "Completed evaluation run not found"}
            for record in normalized:
                cur.execute("""UPDATE mia_evaluation_case_results result
                               SET human_scores = %s::jsonb,
                                   human_overall_score = %s,
                                   human_citation_correct = %s,
                                   human_should_abstain = %s,
                                   human_policy_issue = %s,
                                   human_notes = %s,
                                   human_reviewer_id = %s,
                                   human_reviewed_at = %s
                               FROM mia_evaluation_cases evaluation_case
                               WHERE result.run_id = %s
                                 AND result.case_id = evaluation_case.id
                                 AND evaluation_case.case_key = %s
                               RETURNING result.case_id""",
                            (json.dumps(record["dimension_scores"]), record["overall_score"], record["citation_correct"],
                             record["should_abstain"], record["policy_issue"], record["notes"], reviewer_id,
                             record["reviewed_at"], run_id, record["case_id"]))
                if cur.fetchone() is None:
                    raise ValueError(f"Calibration case not found in run: {record['case_id']}")
            # Recompute the promotion-facing aggregate after calibration. Human
            # outcomes remain separate fields, but the gate consumes their
            # normalized metrics so a later review cannot leave stale approval.
            cur.execute("""SELECT human_scores, human_overall_score, human_citation_correct,
                                  human_should_abstain, human_policy_issue, judge_score
                           FROM mia_evaluation_case_results WHERE run_id = %s
                             AND human_overall_score IS NOT NULL""", (run_id,))
            calibrated = []
            for row in cur.fetchall():
                calibrated.append({
                    "case_id": str(row.get("case_id") or "calibration"),
                    "dimension_scores": row.get("human_scores") or {},
                    "overall_score": row["human_overall_score"],
                    "citation_correct": row["human_citation_correct"],
                    "should_abstain": row["human_should_abstain"],
                    "policy_issue": row["human_policy_issue"],
                    "notes": "persisted calibration",
                    "reviewer_id": reviewer_id,
                    "reviewed_at": "persisted",
                    "judge_score": row.get("judge_score"),
                })
            human_report = aggregate_human_calibration(calibrated)
            cur.execute("SELECT metrics FROM mia_evaluation_runs WHERE id = %s", (run_id,))
            run_row = cur.fetchone() or {}
            metrics = dict(run_row.get("metrics") or {})
            metrics.update({
                "human_overall_score": round(float(human_report.get("overall_score_mean") or 0) / 5.0, 3),
                "human_citation_correctness": human_report.get("citation_correctness_human", 0.0),
                "human_abstention_accuracy": human_report.get("abstention_accuracy_human", 0.0),
                "human_policy_issue_rate": human_report.get("policy_issue_rate", 0.0),
            })
            cur.execute("SELECT metrics FROM mia_evaluation_runs WHERE id = (SELECT run_id FROM mia_evaluation_baselines ORDER BY created_at DESC LIMIT 1)")
            baseline_row = cur.fetchone() or {}
            gate = promotion_decision(metrics, dict(baseline_row.get("metrics") or {}))
            cur.execute("""UPDATE mia_evaluation_runs SET metrics = %s::jsonb,
                              promotion_state = %s, status = %s,
                              reasons = %s::jsonb WHERE id = %s""",
                        (json.dumps(metrics), gate["promotion_state"], gate["promotion_state"], json.dumps(gate["reasons"]), run_id))
            cur.execute("""UPDATE mia_evaluation_metric_segments SET metrics = %s::jsonb
                           WHERE run_id = %s AND segment_type = 'overall' AND segment_value = 'all'""",
                        (json.dumps(metrics), run_id))
        conn.commit()
        report = aggregate_human_calibration(normalized)
        return 200, {"runId": str(run_id), "calibration": report, "promotion": gate, "reviewedBy": pseudonymous_actor_id(reviewer_id)}
    except ValueError as exc:
        conn.rollback()
        return 404, {"error": str(exc)}
    except Exception:
        conn.rollback()
        return 500, {"error": "Unable to persist human calibration"}
    finally:
        conn.close()


def conversation_preview(row):
    """Return a bounded, JSON-safe audit record without authentication secrets."""
    created_at = row.get("created_at")
    return {
        "role": str(row.get("role") or "unknown"),
        "message": str(row.get("message_text") or "")[:700],
        "username": str(row.get("username") or ""),
        "routeKey": str(row.get("route_focus") or "all"),
        "userId": str(row.get("user_id") or ""),
        "sessionId": str(row.get("session_id") or ""),
        "createdAt": created_at.isoformat() if hasattr(created_at, "isoformat") else str(created_at or ""),
    }


def may_use_chat_session(existing_user_id, authenticated_user_id):
    return str(existing_user_id) == str(authenticated_user_id)


def batch_owned_by(manifest, identity):
    return bool(identity and manifest and str((manifest.get("batch") or {}).get("ownerId") or "") == str(identity.get("id") or ""))


def require_owned_upload_batch(batch_id, identity):
    assert_batch_active(batch_id, UPLOAD_STORAGE_DIR)
    manifest = load_batch_manifest(UPLOAD_STORAGE_DIR, batch_id)
    if not batch_owned_by(manifest, identity):
        raise PermissionError("This upload batch belongs to another user")
    return manifest


def delete_upload_batch(batch_id, identity, reason="User deletion request"):
    """Tombstone and purge one owned batch; retries are idempotent."""
    try:
        batch_id = str(uuid.UUID(str(batch_id)))
        assert_batch_active(batch_id, UPLOAD_STORAGE_DIR)
        manifest = load_batch_manifest(UPLOAD_STORAGE_DIR, batch_id)
    except ValueError:
        return 400, {"error": "Invalid batch identifier"}
    except BatchDeletedError:
        return 200, {"batchId": str(batch_id), "status": "completed", "idempotent": True}
    except FileNotFoundError:
        return 404, {"error": "Uploaded batch was not found"}
    owner_id = str((manifest.get("batch") or {}).get("ownerId") or "")
    if identity.get("role") != "administrator" and owner_id != str(identity.get("id") or ""):
        return 403, {"error": "This upload batch belongs to another user"}
    connection = connect_db()
    try:
        result = propagate_batch_deletion(
            batch_id,
            UPLOAD_STORAGE_DIR,
            actor_id=identity.get("id", "unknown"),
            reason=reason,
            db_connection=connection,
            trace_sink=langfuse_batch_deletion_sink(),
        )
        return 200, {
            "batchId": result.batch_id,
            "status": result.status,
            "completedLayers": list(result.completed_layers),
            "deletionRecorded": True,
        }
    except Exception:
        if connection is not None:
            connection.rollback()
        return 503, {"error": "Deletion propagation is pending; retry the request"}
    finally:
        if connection is not None:
            connection.close()


def _write_upload_batch_manifest(manifest):
    batch_id = str((manifest.get("batch") or {}).get("id") or "")
    batch_path = UPLOAD_STORAGE_DIR / batch_id / "manifest.json"
    batch_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def persist_upload_batch_lineage(manifest):
    """Persist source-object lineage only; the upload payload stays on disk."""
    connection = connect_db()
    if connection is None:
        return
    batch = manifest.get("batch") or {}
    batch_id = str(batch.get("id") or "")
    if not batch_id:
        return
    try:
        with connection.cursor() as cur:
            cur.execute(
                """INSERT INTO upload_batches
                   (id, owner_scope, state, file_count, profiled_row_count, total_bytes, manifest_path)
                   VALUES (%s, %s, %s, %s, %s, %s, %s)
                   ON CONFLICT (id) DO UPDATE SET updated_at = NOW(), state = EXCLUDED.state""",
                (
                    batch_id, str(batch.get("ownerId") or ""), str(batch.get("state") or "ready_for_validation"),
                    int(batch.get("fileCount") or 0), int(batch.get("profiledRowCount") or 0), int(batch.get("totalBytes") or 1),
                    f"upload_batches/{batch_id}/manifest.json",
                ),
            )
        register_lineage(connection, batch_id=batch_id, layer="raw_files", subject_type="upload_batch", subject_id=batch_id)
        for file_entry in manifest.get("files") or []:
            register_lineage(connection, batch_id=batch_id, layer="raw_files", subject_type="uploaded_file", subject_id=str(file_entry.get("id") or ""))
            for sheet in file_entry.get("sheets") or []:
                register_lineage(connection, batch_id=batch_id, layer="cleaned_data", subject_type="uploaded_sheet", subject_id=f"{file_entry.get('id')}:{sheet.get('name')}")
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def backfill_upload_batch_lineage():
    """Register legacy filesystem batches without reopening their source files."""
    if not UPLOAD_STORAGE_DIR.is_dir():
        return
    for manifest_path in UPLOAD_STORAGE_DIR.glob("*/manifest.json"):
        try:
            persist_upload_batch_lineage(json.loads(manifest_path.read_text(encoding="utf-8")))
            batch_id = manifest_path.parent.name
            snapshot_path = manifest_path.parent / "analytics.json"
            if snapshot_path.is_file():
                connection = connect_db()
                if connection is not None:
                    try:
                        register_lineage(connection, batch_id=batch_id, layer="analysis_snapshot", subject_type="analytics_snapshot", subject_id=snapshot_path.name)
                        connection.commit()
                    finally:
                        connection.close()
        except (OSError, ValueError, json.JSONDecodeError):
            continue
        except Exception:
            # A later startup retries the non-content lineage backfill.
            continue


def _dataset_run_quality(sheets):
    cleaning = [sheet.get("cleaning") or {} for sheet in sheets]
    return {
        "cleanedRows": sum(int(item.get("cleanedRows") or 0) for item in cleaning),
        "mappingExceptions": sum(int(item.get("mappingExceptions") or 0) for item in cleaning),
        "invalidRatingValues": sum(int(item.get("invalidRatingValues") or 0) for item in cleaning),
        "invalidDateValues": sum(int(item.get("invalidDateValues") or 0) for item in cleaning),
        "invalidAmountValues": sum(int(item.get("invalidAmountValues") or 0) for item in cleaning),
    }


def _analysis_sheets(sheets):
    result = []
    for sheet in sheets:
        copied = dict(sheet)
        # Preserve an explicit full-row source; otherwise ingest provides all cleaned rows here.
        copied["rows"] = list(sheet.get("rows") or sheet.get("previewRows") or [])
        result.append(copied)
    return result


def publish_dataset_run(batch_id, identity):
    try:
        manifest = require_owned_upload_batch(batch_id, identity)
        _, sheets = _cleaned_batch_sheets(UPLOAD_STORAGE_DIR, batch_id)
        published = publish_run(manifest, identity["id"], _dataset_run_quality(sheets))
    except FileNotFoundError:
        return 404, {"error": "Uploaded batch was not found"}
    except PermissionError as exc:
        return 403, {"error": str(exc)}
    except (DatasetRunError, UploadValidationError, ValueError) as exc:
        return 422, {"error": str(exc)}

    run = published["batch"]["datasetRun"]
    run["analysisState"] = "processing"
    event = enqueue_outbox(UPLOAD_STORAGE_DIR / str(batch_id))
    _write_upload_batch_manifest(published)
    DATASET_RUN_WORKER.submit(_process_dataset_run, str(batch_id), event["eventId"])
    return 202, {"run": run, "analysisState": "processing"}


def _process_dataset_run(batch_id, event_id):
    batch_dir = UPLOAD_STORAGE_DIR / batch_id
    event = claim_outbox(batch_dir)
    if not event or event["eventId"] != event_id:
        return
    published = None
    run = None
    try:
        manifest = load_batch_manifest(UPLOAD_STORAGE_DIR, batch_id)
        _, sheets = _cleaned_batch_sheets(UPLOAD_STORAGE_DIR, batch_id)
        run = manifest["batch"]["datasetRun"]
        published = manifest
        snapshot = build_analysis_snapshot(published, _analysis_sheets(sheets))
        comparative_evidence = retrieve_dataset_comparative_evidence(snapshot)
        trace_state = {
            "graph": "dataset", "graph_version": "mia-graph-v1",
            "request_id": str(run.get("publishedAt") or batch_id),
            "batch_id": batch_id, "retrieval_version": "hybrid-retrieval-v1",
            "prompt_version": "dataset-insight-v1", "config_version": "mia-config-v1",
        }
        with RedactedTracer(configured_langfuse_client()).span("mia.dataset.run", trace_state) as span:
            snapshot = run_dataset_graph(
                snapshot,
                run_id=str(run.get("publishedAt") or batch_id),
                graph_version="mia-graph-v1",
                evidence_ids=comparative_evidence,
                checkpointer=MIA_GRAPH_CHECKPOINTER,
                trace_id=span.trace_id,
            )
        trace_state.update({
            "trace_id": snapshot.get("trace_id"),
            "status": snapshot.get("publication_state", "failed"),
            "publication_state": snapshot.get("publication_state"),
            "evidence_ids": snapshot.get("evidence_ids", []),
            "validation_verdict": (snapshot.get("review") or {}).get("reasons", []),
            "final_outcome": snapshot.get("publication_state", "failed"),
        })
        record_graph_result_trace(configured_langfuse_client(), name="mia.dataset.result", state=trace_state, trace_id=snapshot.get("trace_id"))
        if snapshot.get("publication_state") == "pending_human_review":
            graph_record = graph_audit_record({
                "graph": "dataset", "graph_version": "mia-graph-v1",
                "request_id": snapshot.get("graph_run_id"), "trace_id": snapshot.get("trace_id"), "status": "pending_human_review",
                "evidence_ids": snapshot.get("evidence_ids", []),
                "verdict": (snapshot.get("review") or {}).get("reasons", []),
            })
            record_graph_audit(graph_record)
            recommendations = snapshot.get("recommendations") or []
            recommendation = " ".join(str(item.get("type") or item.get("title") or "") for item in recommendations if isinstance(item, dict))
            snapshot["reviewItemId"] = queue_graph_review_item(
                graph_name="dataset", graph_run_id=str(snapshot.get("graph_run_id") or batch_id),
                trace_id=str(graph_record.get("trace_id") or ""),
                title="Dataset insight requires human review",
                reason="; ".join((snapshot.get("review") or {}).get("reasons") or ["Review policy requires approval"]),
                recommendation=recommendation,
                evidence_ids=[str(item) for item in snapshot.get("evidence_ids") or []],
            )
        snapshot_path = batch_dir / "analytics.json"
        snapshot_path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
        connection = connect_db()
        if connection is not None:
            try:
                register_lineage(connection, batch_id=batch_id, layer="analysis_snapshot", subject_type="analytics_snapshot", subject_id=str(snapshot_path.name))
                register_lineage(connection, batch_id=batch_id, layer="audit_trace", subject_type="dataset_graph_run", subject_id=str(run.get("publishedAt") or batch_id))
                connection.commit()
            finally:
                connection.close()
        record_graph_audit(graph_audit_record({
            "graph": "dataset", "graph_version": "mia-graph-v1", "request_id": str(run.get("publishedAt") or batch_id), "trace_id": snapshot.get("trace_id"),
            "status": snapshot.get("publication_state", "pending_human_review"), "evidence_ids": snapshot.get("evidence_ids", []),
            "batch_id": batch_id,
        }))
        run["analysisState"] = "ready"
        run["analysisGeneratedAt"] = snapshot["generatedAt"]
        run.setdefault("analysisAttempts", []).append({"attempt": snapshot.get("attempt", 1), "publicationState": snapshot.get("publication_state"), "evidenceCount": len(snapshot.get("comparativeEvidence") or [])})
        run.setdefault("events", []).extend([
            event_envelope("dataset-run.published", published),
            event_envelope("dataset-run.analysis-ready", published),
        ])
        _write_upload_batch_manifest(published)
        complete_outbox(batch_dir, event_id)
    except Exception:
        if published is None:
            try:
                published = load_batch_manifest(UPLOAD_STORAGE_DIR, batch_id)
                run = (published.get("batch") or {}).get("datasetRun") or None
            except (FileNotFoundError, ValueError, json.JSONDecodeError):
                published = None
                run = None
        if run is not None and published is not None:
            run["analysisState"] = "failed"
            _write_upload_batch_manifest(published)
        complete_outbox(batch_dir, event_id, "failed")


def resume_dataset_run_outbox():
    """Submit durable pending work after a server restart."""
    if not UPLOAD_STORAGE_DIR.is_dir():
        return
    for outbox_path in UPLOAD_STORAGE_DIR.glob("*/outbox.json"):
        try:
            events = json.loads(outbox_path.read_text(encoding="utf-8"))
            changed = False
            for event in events:
                if event.get("status") == "processing":
                    try:
                        age = time.time() - __import__("datetime").datetime.fromisoformat(event["updatedAt"].replace("Z", "+00:00")).timestamp()
                        if age > 300:
                            event["status"] = "pending"; changed = True
                    except (KeyError, ValueError):
                        continue
                if event.get("status") == "pending":
                    DATASET_RUN_WORKER.submit(_process_dataset_run, outbox_path.parent.name, event["eventId"])
            if changed:
                outbox_path.write_text(json.dumps(events, indent=2), encoding="utf-8")
        except (OSError, ValueError, json.JSONDecodeError):
            continue


def dataset_run_analytics(batch_id, identity):
    try:
        manifest = require_owned_upload_batch(batch_id, identity)
    except FileNotFoundError:
        return 404, {"error": "Uploaded batch was not found"}
    except PermissionError as exc:
        return 403, {"error": str(exc)}
    batch = manifest.get("batch") or {}
    run = batch.get("datasetRun") or {}
    if run.get("state") != "published":
        return 409, {"error": "Dataset run has not been published", "analysisState": run.get("analysisState", "not_started")}
    snapshot_path = UPLOAD_STORAGE_DIR / str(batch_id) / "analytics.json"
    if run.get("analysisState") in {"processing", "not_started", "failed"}:
        status = 500 if run.get("analysisState") == "failed" else 202
        return status, {
            "batchId": str(batch_id),
            "version": batch.get("version"),
            "publishedAt": run.get("publishedAt"),
            "analysisState": run.get("analysisState", "processing"),
            "schemaAvailability": {"availableConcepts": list((run.get("lineage") or {}).get("availableConcepts") or [])},
        }
    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    snapshot.setdefault("batchId", str(batch_id))
    snapshot.setdefault("version", batch.get("version"))
    snapshot.setdefault("publishedAt", run.get("publishedAt"))
    snapshot.setdefault("analysisState", run.get("analysisState", "ready"))
    return 200, snapshot


def compare_dataset_runs(current_id, baseline_id, identity):
    current_status, current = dataset_run_analytics(current_id, identity)
    baseline_status, baseline = dataset_run_analytics(baseline_id, identity)
    if current_status != 200:
        return current_status, current
    if baseline_status != 200:
        return baseline_status, baseline
    if current.get("schemaAvailability", {}).get("schemaFingerprint") != baseline.get("schemaAvailability", {}).get("schemaFingerprint"):
        return 422, {"error": "Dataset versions have incompatible schemas"}
    changes = []
    for key, value in current.get("metrics", {}).items():
        baseline_metric = baseline.get("metrics", {}).get(key)
        if not baseline_metric:
            continue
        current_value = value.get("average", value.get("sum"))
        baseline_value = baseline_metric.get("average", baseline_metric.get("sum"))
        if current_value is not None and baseline_value is not None:
            changes.append({"metric": key, "current": current_value, "baseline": baseline_value, "change": round(current_value - baseline_value, 2), "sampleSize": value.get("sampleSize", 0)})
    return 200, {"currentBatchId": current_id, "baselineBatchId": baseline_id, "changes": changes}


def dataset_run_status(batch_id, identity):
    try:
        manifest = require_owned_upload_batch(batch_id, identity)
    except FileNotFoundError:
        return 404, {"error": "Uploaded batch was not found"}
    except PermissionError as exc:
        return 403, {"error": str(exc)}
    batch = manifest.get("batch") or {}
    return 200, {"batchId": batch.get("id"), "version": batch.get("version"), "datasetRun": batch.get("datasetRun") or {"state": "draft", "analysisState": "not_started"}}


def retrieve_dataset_comparative_evidence(snapshot, limit=6):
    """Retrieve approved public evidence for aggregate dataset signals."""
    metrics = snapshot.get("metrics") or {}
    distributions = snapshot.get("distributions") or {}
    route_counts = distributions.get("route") or {}
    requested_route = "all"
    if route_counts:
        requested_route = max(route_counts, key=lambda key: int(route_counts.get(key) or 0))
    metric_terms = " ".join(sorted(str(key) for key in metrics))
    query_payload = {
        "message": f"customer experience comparison {metric_terms}".strip(),
        "activeView": "dataset-insight",
        "routeKey": route_key(requested_route),
    }
    retrieval = retrieve_evidence_from_db(query_payload, limit=limit)
    items = retrieval.get("items", []) if isinstance(retrieval, dict) else []
    comparative = []
    for item in items:
        if bool(item.get("isSynthetic", item.get("is_synthetic", False))):
            continue
        source_tier = str(item.get("sourceTier") or item.get("source_tier") or "")
        if source_tier != "public_snapshot":
            continue
        metadata = item.get("metadata") if isinstance(item.get("metadata"), dict) else {}
        comparative.append({
            "id": str(item.get("id") or ""),
            "sourceType": str(item.get("source") or item.get("type") or "external_public"),
            "title": str(item.get("title") or "External comparison evidence"),
            "excerpt": str(item.get("body") or item.get("translatedContent") or "")[:500],
            "collectedAt": str(item.get("timestamp") or ""),
            "reliability": str(metadata.get("reliability") or "directional"),
            "coverage": str(item.get("route") or "public source coverage"),
        })
    return comparative


def review_dataset_run(batch_id, identity, payload):
    """Apply a traceable Review Console decision to a dataset snapshot."""
    decision = dict(payload or {})
    status = str(decision.get("status") or "").lower()
    if status not in {"approved", "corrected", "rejected", "withdrawn", "reanalyse"}:
        return 400, {"error": "Invalid dataset review decision"}
    decision["reviewedAt"] = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    try:
        manifest = require_owned_upload_batch(batch_id, identity)
    except FileNotFoundError:
        return 404, {"error": "Uploaded batch was not found"}
    except PermissionError as exc:
        return 403, {"error": str(exc)}

    run = (manifest.get("batch") or {}).get("datasetRun") or {}
    if run.get("state") != "published":
        return 409, {"error": "Dataset run has not been published"}
    snapshot_path = UPLOAD_STORAGE_DIR / str(batch_id) / "analytics.json"
    if not snapshot_path.is_file():
        return 409, {"error": "Dataset analysis snapshot is not ready"}
    try:
        snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    except (OSError, ValueError, json.JSONDecodeError):
        return 500, {"error": "Dataset analysis snapshot is invalid"}

    evidence = decision.pop("comparativeEvidence", None)
    if evidence is None:
        evidence = snapshot.get("comparativeEvidence") or snapshot.get("evidence_ids") or retrieve_dataset_comparative_evidence(snapshot)
    result = run_dataset_graph(
        snapshot,
        run_id=str(run.get("publishedAt") or batch_id),
        graph_version=str(snapshot.get("graphVersion") or "mia-graph-v1"),
        evidence_ids=evidence,
        review_decision=decision,
        retry=True,
        checkpointer=MIA_GRAPH_CHECKPOINTER,
    )
    snapshot_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    run["analysisState"] = "ready"
    run["reviewState"] = status
    run["publicationState"] = result["publication_state"]
    run.setdefault("analysisAttempts", []).append({"attempt": result.get("attempt", 1), "publicationState": result.get("publication_state"), "reviewStatus": status, "reviewerId": decision.get("reviewerId")})
    run.setdefault("events", []).append(event_envelope("dataset-run.reviewed", manifest))
    _write_upload_batch_manifest(manifest)
    return 200, {"snapshot": result, "datasetRun": run}


def list_owned_dataset_runs(identity):
    runs = []
    if not UPLOAD_STORAGE_DIR.is_dir():
        return {"runs": runs}
    for manifest_path in UPLOAD_STORAGE_DIR.glob("*/manifest.json"):
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            batch = manifest.get("batch") or {}
            run = batch.get("datasetRun") or {}
            if str(batch.get("ownerId") or "") == str(identity.get("id") or "") and run.get("state") == "published":
                runs.append({"batchId": batch.get("id"), "version": batch.get("version"), "sourceType": batch.get("sourceType"), "datasetRun": run})
        except (OSError, ValueError, json.JSONDecodeError):
            continue
    return {"runs": sorted(runs, key=lambda item: str(item["datasetRun"].get("publishedAt") or ""), reverse=True)}


def admin_conversation_logs(limit=100):
    limit = max(1, min(int(limit), 500))
    conn = connect_db()
    if conn is None:
        return {"connected": False, "conversations": []}
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT m.role, m.message_text, m.created_at, m.session_id, s.user_id, s.route_focus, u.username
                FROM chat_messages m
                JOIN chat_sessions s ON s.id = m.session_id
                JOIN app_users u ON u.id = s.user_id
                ORDER BY m.created_at DESC
                LIMIT %s
                """,
                (limit,),
            )
            conversations = [conversation_preview(row) for row in cur.fetchall()]
    finally:
        conn.close()
    return {"connected": True, "conversations": conversations}


def user_history(identity, limit=100):
    """Return only the authenticated user's own chat history."""
    limit = max(1, min(int(limit), 500))
    conn = connect_db()
    if conn is None:
        return {"connected": False, "history": []}
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT s.id AS session_id, s.active_view, s.route_focus, s.created_at,
                       s.updated_at, COUNT(m.id) AS message_count,
                       MAX(m.message_text) FILTER (WHERE m.role = 'user') AS last_message
                FROM chat_sessions s
                LEFT JOIN chat_messages m ON m.session_id = s.id
                WHERE s.user_id = %s
                GROUP BY s.id
                ORDER BY s.updated_at DESC
                LIMIT %s
                """,
                (identity["id"], limit),
            )
            rows = cur.fetchall()
    finally:
        conn.close()
    return {"connected": True, "history": [
        {"sessionId": str(row["session_id"]), "activeView": row["active_view"], "routeFocus": row["route_focus"],
         "createdAt": row["created_at"].isoformat(), "updatedAt": row["updated_at"].isoformat(),
         "messageCount": int(row["message_count"]), "lastMessage": row["last_message"] or ""}
        for row in rows
    ]}


def admin_observability():
    conn = connect_db()
    if conn is None:
        return {"connected": False, "layers": [], "routeMonitoring": [], "tokens": {"input": 0, "output": 0, "total": 0, "requests": 0}, "users": []}
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT COALESCE(SUM(input_tokens), 0) AS input_tokens, COALESCE(SUM(output_tokens), 0) AS output_tokens,
                       COALESCE(SUM(total_tokens), 0) AS total_tokens, COUNT(*) AS requests
                FROM rag_usage_events
                WHERE created_at >= NOW() - INTERVAL '24 hours'
                """
            )
            tokens = dict(cur.fetchone())
            cur.execute(
                """
                SELECT u.username, COUNT(e.id) AS requests, COALESCE(SUM(e.total_tokens), 0) AS total_tokens,
                       MAX(e.created_at) AS last_seen_at
                FROM app_users u
                LEFT JOIN rag_usage_events e ON e.user_id = u.id AND e.created_at >= NOW() - INTERVAL '24 hours'
                GROUP BY u.id, u.username
                ORDER BY total_tokens DESC, u.username
                """
            )
            users = [dict(row) for row in cur.fetchall()]
            cur.execute(
                """
                SELECT layer->>'name' AS name, COUNT(*) AS requests,
                       ROUND(AVG((layer->>'durationMs')::numeric), 1) AS latency_ms,
                       COALESCE(SUM((layer->>'records')::integer), 0) AS records
                FROM rag_usage_events e
                CROSS JOIN LATERAL jsonb_array_elements(e.retrieval_trace->'layers') AS layer
                WHERE e.created_at >= NOW() - INTERVAL '24 hours'
                GROUP BY layer->>'name'
                ORDER BY name
                """
            )
            layers = [dict(row) for row in cur.fetchall()]
            cur.execute(
                """
                SELECT s.route_focus AS route_key,
                       CASE s.route_focus
                         WHEN 'dover-calais' THEN 'Dover-Calais'
                         WHEN 'newhaven-dieppe' THEN 'Newhaven-Dieppe'
                         WHEN 'newcastle-ijmuiden' THEN 'Newcastle-IJmuiden'
                         WHEN 'jersey' THEN 'Jersey / Channel Islands'
                         ELSE 'All signals'
                       END AS route_name,
                       COUNT(event.id) AS requests,
                       COALESCE(SUM(event.total_tokens), 0) AS total_tokens,
                       ROUND(AVG(event.latency_ms), 1) AS avg_latency_ms,
                       COALESCE(SUM(event.records), 0) AS records,
                       MAX(event.created_at) AS last_seen_at
                FROM (
                    SELECT e.id, e.chat_session_id, e.total_tokens, e.created_at,
                           COALESCE(SUM((layer->>'durationMs')::numeric), 0) AS latency_ms,
                           COALESCE(SUM((layer->>'records')::integer), 0) AS records
                    FROM rag_usage_events e
                    CROSS JOIN LATERAL jsonb_array_elements(e.retrieval_trace->'layers') AS layer
                    WHERE e.created_at >= NOW() - INTERVAL '24 hours'
                    GROUP BY e.id, e.chat_session_id, e.total_tokens, e.created_at
                ) AS event
                JOIN chat_sessions s ON s.id = event.chat_session_id
                GROUP BY s.route_focus
                ORDER BY requests DESC, route_name
                """
            )
            route_monitoring = summarize_route_monitoring([dict(row) for row in cur.fetchall()])
    finally:
        conn.close()
    return {"connected": True, "demo": {"batch": DEMO_BATCH, "events": DEMO_EVENT_COUNT}, "tokens": {"input": int(tokens["input_tokens"]), "output": int(tokens["output_tokens"]), "total": int(tokens["total_tokens"]), "requests": int(tokens["requests"])}, "layers": layers, "users": users, "routeMonitoring": route_monitoring}


def store_chat_turn(payload, answer, rag_context, identity=None):
    conn = connect_db()
    if conn is None:
        return valid_uuid(payload.get("sessionId"))
    if not identity:
        raise ValueError("Authenticated identity is required to store chat history")

    session_id = valid_uuid(payload.get("sessionId"))
    language = str(payload.get("language", "English"))[:80]
    active_view = str(payload.get("activeView", "dashboard"))[:80]
    route_focus = str(payload.get("routeKey") or payload.get("routeName") or "all")[:120]
    message = str(payload.get("message", ""))
    evidence_context = json.dumps(
        {
            "evidence": compact_evidence(rag_context.get("evidenceItems", [])),
            "insights": compact_insights(rag_context.get("insightItems", [])),
            "memories": compact_memories(rag_context.get("memoryItems", [])),
        },
        ensure_ascii=False,
    )
    page_context = json.dumps(rag_context.get("pageContext") or {}, ensure_ascii=False)

    try:
        with conn.cursor() as cur:
            cur.execute("SELECT user_id FROM chat_sessions WHERE id = %s FOR UPDATE", (session_id,))
            existing_session = cur.fetchone()
            if existing_session and not may_use_chat_session(existing_session["user_id"], identity["id"]):
                session_id = str(uuid.uuid4())
            cur.execute(
                """
                INSERT INTO chat_sessions (id, user_id, user_language, active_view, route_focus)
                VALUES (%s, %s, %s, %s, %s)
                ON CONFLICT (id) DO UPDATE
                SET user_language = EXCLUDED.user_language,
                    active_view = EXCLUDED.active_view,
                    route_focus = EXCLUDED.route_focus,
                    updated_at = NOW()
                """,
                (session_id, identity["id"], language, active_view, route_focus),
            )
            cur.execute(
                """
                INSERT INTO chat_messages (id, session_id, role, message_text, evidence_context, page_context)
                VALUES (%s, %s, 'user', %s, %s::jsonb, %s::jsonb)
                """,
                (str(uuid.uuid4()), session_id, message, evidence_context, page_context),
            )
            cur.execute(
                """
                INSERT INTO chat_messages (id, session_id, role, message_text, evidence_context, page_context)
                VALUES (%s, %s, 'assistant', %s, %s::jsonb, %s::jsonb)
                """,
                (str(uuid.uuid4()), session_id, answer, evidence_context, page_context),
            )
        conn.commit()
    finally:
        conn.close()

    return session_id


def is_valid_review_status(value):
    return str(value or "") in VALID_REVIEW_STATUSES


def filter_review_seed_items(items, params):
    params = params or {}

    def first(name):
        value = params.get(name, [""])
        if isinstance(value, list):
            value = value[0] if value else ""
        return str(value or "").strip().lower()

    query = first("q")
    layer = first("layer")
    status = first("status")
    severity = first("severity")
    route = first("route")

    filtered = []
    for item in items:
        haystack = (
            f"{item.get('title', '')} "
            f"{item.get('reason', '')} "
            f"{item.get('recommendation', '')}"
        ).lower()
        if query and query not in haystack:
            continue
        if layer and item.get("layer") != layer:
            continue
        if status and item.get("status") != status:
            continue
        if severity and item.get("severity") != severity:
            continue
        if route and item.get("route_key") != route:
            continue
        filtered.append(item)
    return filtered


def normalize_review_row(row):
    item = dict(row)
    for key in ["metadata", "evidence_chain", "artifacts", "action_history"]:
        value = item.get(key)
        if isinstance(value, str):
            try:
                item[key] = json.loads(value)
            except json.JSONDecodeError:
                item[key] = value
    if "source" not in item:
        item["source"] = item.get("source_key") or "internal review"
    if "evidence_chain" not in item or item.get("evidence_chain") is None:
        item["evidence_chain"] = item.get("metadata", {}).get("evidence_chain", [])
    if "artifacts" not in item or item.get("artifacts") is None:
        item["artifacts"] = item.get("metadata", {}).get("artifacts", [])
    return item


def list_review_items(params=None):
    params = params or {}
    conn = connect_db()
    if conn is None:
        return {
            "connected": False,
            "items": filter_review_seed_items(REVIEW_CONSOLE_SEED_ITEMS, params),
            "source": "seed",
        }

    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id, subject_type, subject_id, subject_label, layer, status, severity,
                       title, reason, recommendation, publish_state, route_key, source_key,
                       language, original_output, supervisor_verdict, suggested_fix,
                       downstream_impact, metadata, created_at, updated_at,
                       COALESCE((SELECT jsonb_agg(jsonb_build_object(
                         'id', a.id, 'action', a.action_type, 'actor', a.actor_name,
                         'actorId', a.actor_id, 'rationale', a.notes, 'correction', a.correction,
                         'previousStatus', a.previous_status, 'nextStatus', a.next_status,
                         'graphRunId', a.graph_run_id, 'traceId', a.trace_id, 'createdAt', a.created_at
                       ) ORDER BY a.created_at DESC) FROM review_item_actions a
                       WHERE a.review_item_id = review_items.id), '[]'::jsonb) AS action_history
                FROM review_items
                ORDER BY created_at DESC, id DESC
                LIMIT 200
                """
            )
            rows = [normalize_review_row(row) for row in cur.fetchall()]
            items = filter_review_seed_items(rows, params)
            if not items:
                items = filter_review_seed_items(REVIEW_CONSOLE_SEED_ITEMS, params)
            return {"connected": True, "items": items, "source": "database"}
    finally:
        conn.close()


def update_review_item_status(item_id, payload):
    status = str(payload.get("status", "")).strip()
    notes = str(payload.get("notes", "")).strip()
    actor_name = str(payload.get("actorName") or "Internal reviewer")[:120]
    if not is_valid_review_status(status):
        return 400, {"error": "Invalid review status"}

    conn = connect_db()
    if conn is None:
        return 503, {"error": "Database is not connected; seed review items are read-only"}

    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                UPDATE review_items
                SET status = %s, updated_at = NOW()
                WHERE id = %s
                RETURNING id, status
                """,
                (status, item_id),
            )
            row = cur.fetchone()
            if row is None:
                conn.rollback()
                return 404, {"error": "Review item not found"}
            cur.execute(
                """
                INSERT INTO review_item_actions (review_item_id, actor_type, actor_name, action_type, notes)
                VALUES (%s, 'human', %s, %s, %s)
                """,
                (item_id, actor_name, status, notes),
            )
        conn.commit()
        return 200, {"item": dict(row)}
    finally:
        conn.close()


def apply_review_action(item_id, payload, identity):
    """Apply one auditable human review action, safely replayable by idempotency key."""
    action = str(payload.get("action") or "").strip().lower()
    rationale = str(payload.get("rationale") or payload.get("notes") or "").strip()[:4000]
    correction = str(payload.get("correction") or "").strip()[:8000]
    idempotency_key = str(payload.get("idempotencyKey") or uuid.uuid4())[:180]
    actor_id = str((identity or {}).get("id") or "")
    actor_name = str((identity or {}).get("username") or (identity or {}).get("name") or "Internal reviewer")[:120]
    graph_run_id = str(payload.get("graphRunId") or "")[:180]
    trace_id = str(payload.get("traceId") or "")[:180]
    if not actor_id:
        return 401, {"error": "Authenticated reviewer is required"}
    if not rationale:
        return 400, {"error": "A rationale is required for every review action"}
    if action == "correct" and not correction:
        return 400, {"error": "Correction content is required for a correction action"}
    if action not in {"approve", "correct", "reject", "withdraw", "request_reanalysis"}:
        return 400, {"error": f"Unsupported review action: {action}"}

    conn = connect_db()
    if conn is None:
        return 503, {"error": "Database is not connected; seed review items are read-only"}
    try:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT id, status, publish_state, metadata FROM review_items WHERE id = %s FOR UPDATE",
                (item_id,),
            )
            item = cur.fetchone()
            if item is None:
                conn.rollback()
                return 404, {"error": "Review item not found"}
            cur.execute(
                """SELECT id, action_type, previous_status, next_status, notes,
                          correction, actor_id, actor_name, graph_run_id, trace_id, created_at
                   FROM review_item_actions
                   WHERE review_item_id = %s AND idempotency_key = %s""",
                (item_id, idempotency_key),
            )
            existing = cur.fetchone()
            if existing is not None:
                conn.rollback()
                return 200, {"replayed": True, "action": dict(existing)}
            current_status = str(item["status"] if isinstance(item, dict) else item[1])
            try:
                transition = transition_for_action(current_status, action)
            except ValueError as exc:
                conn.rollback()
                return 409, {"error": str(exc), "status": current_status}
            metadata = item["metadata"] if isinstance(item, dict) else item[3]
            metadata = dict(metadata or {})
            if graph_run_id:
                metadata["graph_run_id"] = graph_run_id
            if trace_id:
                metadata["trace_id"] = trace_id
            cur.execute(
                """UPDATE review_items
                   SET status = %s, publish_state = %s, metadata = %s::jsonb, updated_at = NOW()
                   WHERE id = %s""",
                (transition["next_status"], transition["publish_state"], json.dumps(metadata), item_id),
            )
            cur.execute(
                """INSERT INTO review_item_actions
                   (review_item_id, actor_type, actor_id, actor_name, action_type, notes,
                    graph_run_id, trace_id, previous_status, next_status, correction, idempotency_key)
                   VALUES (%s, 'human', %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                   RETURNING id, action_type, previous_status, next_status,
                             notes, correction, actor_id, actor_name, graph_run_id, trace_id, created_at""",
                (item_id, actor_id, actor_name, action, rationale, graph_run_id, trace_id,
                 current_status, transition["next_status"], correction, idempotency_key),
            )
            record = dict(cur.fetchone())
        conn.commit()
        return 200, {"replayed": False, "item": {"id": int(item_id), "status": transition["next_status"], "publish_state": transition["publish_state"]}, "action": record}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def queue_graph_review_item(*, graph_name, graph_run_id, trace_id, title, reason,
                            recommendation, evidence_ids, route_key="all"):
    """Persist one pending graph artifact without storing raw prompt or upload content."""
    conn = connect_db()
    if conn is None:
        return None
    item = review_item_payload_for_graph(
        graph_name=graph_name, graph_run_id=graph_run_id, trace_id=trace_id, title=title,
        reason=reason, recommendation=recommendation, evidence_ids=evidence_ids, route_key=route_key,
    )
    try:
        with conn.cursor() as cur:
            cur.execute(
                """INSERT INTO review_runs (run_key, run_type, trigger_source, status, summary, metadata)
                   VALUES (%s, %s, 'graph', 'pending_human_review', %s, %s::jsonb)
                   ON CONFLICT (run_key) DO UPDATE SET status = EXCLUDED.status, metadata = EXCLUDED.metadata
                   RETURNING id""",
                (f"{graph_name}:{graph_run_id}", f"{graph_name}_graph", item["title"], json.dumps(item["metadata"])),
            )
            run_id = cur.fetchone()["id"]
            cur.execute(
                "SELECT id FROM review_items WHERE run_id = %s AND subject_id = %s",
                (run_id, item["subject_id"]),
            )
            existing = cur.fetchone()
            if existing:
                conn.commit()
                return int(existing["id"])
            cur.execute(
                """INSERT INTO review_items
                   (run_id, subject_type, subject_id, subject_label, layer, status, severity,
                    title, reason, recommendation, publish_state, route_key, source_key, metadata)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb)
                   RETURNING id""",
                (run_id, item["subject_type"], item["subject_id"], item["subject_label"], item["layer"],
                 item["status"], item["severity"], item["title"], item["reason"], item["recommendation"],
                 item["publish_state"], item["route_key"], item["source_key"], json.dumps(item["metadata"])),
            )
            review_item_id = cur.fetchone()["id"]
            cur.execute(
                """INSERT INTO review_item_actions
                   (review_item_id, actor_type, actor_name, action_type, notes, graph_run_id, trace_id,
                    previous_status, next_status)
                   VALUES (%s, 'system', 'Mia graph', 'queued', %s, %s, %s, 'draft', 'pending_human_review')""",
                (review_item_id, item["reason"], graph_run_id, trace_id),
            )
        conn.commit()
        return int(review_item_id)
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def run_validation_harness(payload):
    from work.validation_agent import validate_batch
    from work.review_store import persist_validation_result

    cases = payload.get("cases")
    if cases is None:
        case = payload.get("case")
        cases = [case] if case else [payload]
    if not isinstance(cases, list) or not all(isinstance(case, dict) for case in cases):
        return 400, {"error": "Validation cases must be a list of objects"}

    result = validate_batch(cases, run_key=payload.get("run_key"))
    conn = connect_db()
    if conn is None:
        result["persistence"] = {
            "persisted": False,
            "reason": "Database is not connected; validation result returned without PG writeback",
        }
        return 200, result

    try:
        result["persistence"] = {
            "persisted": True,
            **persist_validation_result(result, conn=conn),
        }
        return 200, result
    except Exception as exc:
        return 500, {"error": "Validation run failed during PG writeback", "detail": str(exc)}
    finally:
        conn.close()


def build_chat_validation_case(payload, result, rag_context, session_id=None):
    message = str(payload.get("message", ""))
    answer = str(result.get("answer", ""))
    evidence_cards = result.get("evidence") or []
    language = str(payload.get("language", "English"))[:80]
    route_key = str(payload.get("routeKey") or payload.get("routeName") or "all")[:120]
    route_key = route_key if route_key else "all"
    return {
        "case_id": session_id or payload.get("sessionId") or f"chat-{uuid.uuid4().hex[:12]}",
        "subject_type": "chat_turn",
        "subject_id": session_id or payload.get("sessionId") or "",
        "subject_label": message[:140] or "Chat turn",
        "message": message,
        "response": answer,
        "language": language,
        "route_key": route_key,
        "source_type": "chat",
        "source_label": "Mia /api/chat",
        "evidence_cards": evidence_cards,
        "claims_based_on_reviews": True,
        "downstream_impact": "Controls whether this chat turn is safe to surface in the Review Console and later into dashboard-facing assistant behavior.",
        "artifact": "",
        "metadata": {
            "mode": result.get("mode"),
            "intent": result.get("intent"),
            "route_focus": route_key,
            "evidence_count": len(evidence_cards),
            "insight_count": len(result.get("insights") or []),
            "memory_count": len(result.get("memories") or []),
            "page_context_keys": sorted((rag_context.get("pageContext") or {}).keys()),
        },
    }


def database_summary():
    conn = connect_db()
    if conn is None:
        return {"connected": False, "reason": "DATABASE_URL or psycopg is not available"}
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT
                  (SELECT COUNT(*) FROM companies) AS companies,
                  (SELECT COUNT(*) FROM routes) AS routes,
                  (SELECT COUNT(*) FROM sources) AS sources,
                  (SELECT COUNT(*) FROM evidence_items) AS evidence_items,
                  (SELECT COUNT(*) FROM insight_items) AS insight_items,
                  (SELECT COUNT(*) FROM project_memories) AS project_memories,
                  (SELECT COUNT(*) FROM chat_sessions) AS chat_sessions,
                  (SELECT COUNT(*) FROM chat_messages) AS chat_messages,
                  (SELECT COUNT(*) FROM rag_evaluation_items) AS rag_evaluation_items,
                  (SELECT COUNT(*) FROM review_items) AS review_items
                """
            )
            summary = dict(cur.fetchone())
            summary["connected"] = True
            return summary
    finally:
        conn.close()


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        static_dir = APP_DIR / "dist" if (APP_DIR / "dist").exists() else APP_DIR
        super().__init__(*args, directory=str(static_dir), **kwargs)

    def send_json(self, status, payload):
        body = json_response_body(payload)
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        origin = self.headers.get("Origin")
        if origin in LOGIN_APP_ORIGINS:
            self.send_header("Access-Control-Allow-Origin", origin)
            self.send_header("Vary", "Origin")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_OPTIONS(self):
        self.send_response(204)
        origin = self.headers.get("Origin")
        if origin in LOGIN_APP_ORIGINS:
            self.send_header("Access-Control-Allow-Origin", origin)
            self.send_header("Access-Control-Allow-Headers", "Authorization, Content-Type")
            self.send_header("Access-Control-Allow-Methods", "GET, POST, DELETE, OPTIONS")
            self.send_header("Vary", "Origin")
        self.end_headers()

    def send_download(self, content, content_type, filename):
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def source_workbook_page(self, query):
        sheet_name = (query.get("sheet") or [""])[0]
        try:
            page = max(1, int((query.get("page") or ["1"])[0]))
            page_size = min(100, max(25, int((query.get("page_size") or ["100"])[0])))
        except ValueError:
            self.send_json(400, {"error": "page and page_size must be numbers"})
            return

        workbook_name = next((name for name, sheets in SOURCE_WORKBOOKS.items() if sheet_name in sheets), None)
        if workbook_name is None:
            self.send_json(404, {"error": "Unknown worksheet"})
            return

        workbook_path = SOURCE_WORKBOOK_DIR / workbook_name
        if not workbook_path.exists():
            self.send_json(404, {"error": "Source workbook is unavailable"})
            return

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        worksheet = workbook[sheet_name]
        columns = [str(cell.value or "") for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        total_rows = sum(1 for _ in worksheet.iter_rows(min_row=2, values_only=True))
        offset = (page - 1) * page_size
        rows = []
        if offset < total_rows:
            for values in islice(worksheet.iter_rows(min_row=offset + 2, values_only=True), page_size):
                row = {}
                for column, value in zip(columns, values):
                    if RESTRICTED_SOURCE_FIELD.search(column) and value not in (None, ""):
                        row[column] = "Restricted"
                    elif hasattr(value, "isoformat"):
                        row[column] = value.isoformat().replace("+00:00", "Z")
                    elif value is None:
                        row[column] = None
                    else:
                        row[column] = str(value)
                rows.append(row)
        workbook.close()
        self.send_json(
            200,
            {
                "sheet": sheet_name,
                "columns": columns,
                "rows": rows,
                "page": page,
                "pageSize": page_size,
                "totalRows": total_rows,
                "totalPages": max(1, math.ceil(total_rows / page_size)),
            },
        )

    def uploaded_sheet_page(self, batch_id, query):
        identity = identity_from_authorization(self.headers.get("Authorization"))
        if not identity:
            self.send_json(401, {"error": "Sign in is required"})
            return
        file_id = (query.get("file_id") or [""])[0]
        sheet_name = (query.get("sheet") or [""])[0]
        try:
            page = max(1, int((query.get("page") or ["1"])[0]))
            page_size = min(100, max(25, int((query.get("page_size") or ["100"])[0])))
        except ValueError:
            self.send_json(400, {"error": "page and page_size must be numbers"})
            return
        try:
            require_owned_upload_batch(batch_id, identity)
            exceptions_only = (query.get("exceptions_only") or [""])[0].casefold() == "true"
            self.send_json(200, uploaded_sheet_page(UPLOAD_STORAGE_DIR, batch_id, file_id, sheet_name, page, page_size, exceptions_only=exceptions_only))
        except FileNotFoundError:
            self.send_json(404, {"error": "Uploaded worksheet was not found"})
        except PermissionError as exc:
            self.send_json(403, {"error": str(exc)})
        except UploadValidationError as exc:
            self.send_json(422, {"error": str(exc)})
        except ValueError:
            self.send_json(400, {"error": "Invalid batch or file identifier"})

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/governance/production-readiness":
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not has_capability(identity, "governance_admin"):
                self.send_json(403, {"error": "Governance administrator access is required"})
                return
            self.send_json(200, production_readiness())
            return
        if parsed.path == "/api/history":
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            self.send_json(200, user_history(identity))
            return
        if parsed.path == "/api/health":
            self.send_json(200, {"status": "ok", "database": database_summary()})
            return
        if parsed.path == "/api/review-items":
            self.send_json(200, list_review_items(parse_qs(parsed.query)))
            return
        if parsed.path.startswith("/api/review-items/") and parsed.path.endswith("/actions"):
            self.send_json(405, {"error": "POST is required for review actions"})
            return
        if parsed.path == "/api/admin/observability":
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity or identity["role"] != "administrator":
                self.send_json(403, {"error": "Administrator access is required"})
                return
            self.send_json(200, admin_observability())
            return
        if parsed.path == "/api/admin/conversations":
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity or identity["role"] != "administrator":
                self.send_json(403, {"error": "Administrator access is required"})
                return
            raw_limit = (parse_qs(parsed.query).get("limit") or ["100"])[0]
            try:
                self.send_json(200, admin_conversation_logs(raw_limit))
            except ValueError:
                self.send_json(400, {"error": "limit must be a number"})
            return
        if parsed.path.startswith("/api/admin/evaluations/") and not parsed.path.endswith("/calibration"):
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity or identity["role"] != "administrator":
                self.send_json(403, {"error": "Administrator access is required"})
                return
            status, payload = admin_evaluation_detail(parsed.path.split("/")[4])
            self.send_json(status, payload)
            return
        if parsed.path == "/api/admin/evaluations":
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity or identity["role"] != "administrator":
                self.send_json(403, {"error": "Administrator access is required"})
                return
            conn = connect_db()
            if conn is None:
                self.send_json(503, {"error": "Database is required for evaluation history"})
                return
            try:
                with conn.cursor() as cur:
                    cur.execute("""SELECT r.id, s.version AS evaluation_set_version, r.status, r.promotion_state, r.reasons, r.metrics, r.trace_id, r.started_at, r.completed_at
                                   FROM mia_evaluation_runs r JOIN mia_evaluation_set_versions s ON s.id = r.evaluation_set_id
                                   ORDER BY r.started_at DESC LIMIT 50""")
                    rows = [dict(row) for row in cur.fetchall()]
                self.send_json(200, {"runs": rows})
            finally:
                conn.close()
            return
        if parsed.path == "/api/source-workbook":
            self.source_workbook_page(parse_qs(parsed.query))
            return
        if parsed.path == "/api/upload-batches":
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            self.send_json(200, list_owned_dataset_runs(identity))
            return
        if parsed.path.startswith("/api/upload-batches/") and parsed.path.endswith("/events"):
            query = parse_qs(parsed.query)
            identity = identity_from_authorization(self.headers.get("Authorization") or f"Bearer {(query.get('access_token') or [''])[0]}")
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            batch_id = parsed.path.split("/")[3]
            try:
                require_owned_upload_batch(batch_id, identity)
            except FileNotFoundError:
                self.send_json(404, {"error": "Uploaded batch was not found"})
                return
            except PermissionError as exc:
                self.send_json(403, {"error": str(exc)})
                return
            self.send_response(200)
            self.send_header("Content-Type", "text/event-stream")
            self.send_header("Cache-Control", "no-cache")
            self.send_header("Connection", "keep-alive")
            self.end_headers()
            for _ in range(60):
                status, payload = dataset_run_status(batch_id, identity)
                self.wfile.write((f"event: dataset-run-status\\ndata: {json.dumps(payload, ensure_ascii=False)}\\n\\n").encode())
                self.wfile.flush()
                if payload.get("datasetRun", {}).get("analysisState") in {"ready", "failed"}:
                    break
                time.sleep(1)
            return
        if parsed.path.startswith("/api/upload-batches/") and parsed.path.endswith("/run-status"):
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            status, payload = dataset_run_status(parsed.path.split("/")[3], identity)
            self.send_json(status, payload)
            return
        if parsed.path.startswith("/api/upload-batches/") and parsed.path.endswith("/analytics"):
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            status, payload = dataset_run_analytics(parsed.path.split("/")[3], identity)
            self.send_json(status, payload)
            return
        if parsed.path.startswith("/api/upload-batches/") and parsed.path.endswith("/compare"):
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            baseline_id = (parse_qs(parsed.query).get("baseline") or [""])[0]
            status, payload = compare_dataset_runs(parsed.path.split("/")[3], baseline_id, identity)
            self.send_json(status, payload)
            return
        if parsed.path.startswith("/api/upload-batches/") and parsed.path.endswith("/sheets"):
            self.uploaded_sheet_page(parsed.path.split("/")[3], parse_qs(parsed.query))
            return
        if parsed.path.startswith("/api/upload-batches/") and parsed.path.endswith("/export"):
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            try:
                require_owned_upload_batch(parsed.path.split("/")[3], identity)
                content, content_type, filename = export_cleaned_batch(
                    UPLOAD_STORAGE_DIR,
                    parsed.path.split("/")[3],
                    (parse_qs(parsed.query).get("format") or ["xlsx"])[0],
                    (parse_qs(parsed.query).get("file_id") or [None])[0],
                )
                self.send_download(content, content_type, filename)
            except FileNotFoundError:
                self.send_json(404, {"error": "Uploaded batch was not found"})
            except PermissionError as exc:
                self.send_json(403, {"error": str(exc)})
            except UploadValidationError as exc:
                self.send_json(422, {"error": str(exc)})
            except ValueError:
                self.send_json(400, {"error": "Invalid batch identifier"})
            return
        if not parsed.path.startswith("/api/") and "." not in Path(parsed.path).name:
            self.path = "/index.html"
        super().do_GET()

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/governance/production-mode":
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not has_capability(identity, "governance_admin"):
                self.send_json(403, {"error": "Governance administrator access is required"})
                return
            try:
                self.send_json(200, enable_production_mode(identity))
            except RuntimeError as exc:
                self.send_json(409, {"error": str(exc), "governance": production_readiness()})
            except PermissionError as exc:
                self.send_json(403, {"error": str(exc)})
            return
        if parsed.path in {"/api/auth/register", "/api/auth/login"}:
            length = int(self.headers.get("Content-Length", "0"))
            try:
                payload = json.loads(self.rfile.read(length).decode("utf-8"))
            except (ValueError, json.JSONDecodeError):
                self.send_json(400, {"error": "Invalid JSON"})
                return
            if parsed.path == "/api/auth/register":
                try:
                    account = create_account(payload.get("username"), str(payload.get("password") or ""))
                except (ValueError, RuntimeError) as exc:
                    self.send_json(400, {"error": str(exc)})
                    return
                if account is None:
                    self.send_json(409, {"error": "Username is already registered"})
                    return
                self.send_json(201, {"user": {"id": account["id"], "username": account["username"], "role": account["role"]}})
                return
            identity = authenticate_account(payload.get("username"), payload.get("password"))
            if not identity:
                self.send_json(401, {"error": "Invalid username or password"})
                return
            self.send_json(200, {"accessToken": create_access_session(identity), "user": identity})
            return
        calibration_path = parsed.path.startswith("/api/admin/evaluations/") and parsed.path.endswith("/calibration")
        if parsed.path in {"/api/admin/evaluation-sets/freeze", "/api/admin/evaluations", "/api/admin/evaluation-baselines", "/api/admin/mia-triggers/external-refresh", "/api/admin/mia-triggers/version-change"} or calibration_path:
            try:
                require_production_ready()
            except PermissionError as exc:
                self.send_json(503, {"error": str(exc)})
                return
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity or identity["role"] != "administrator":
                self.send_json(403, {"error": "Administrator access is required"})
                return
            try:
                payload = json.loads(self.rfile.read(int(self.headers.get("Content-Length", "0"))).decode("utf-8"))
            except (ValueError, json.JSONDecodeError):
                self.send_json(400, {"error": "Invalid JSON"})
                return
            if parsed.path.endswith("/freeze"):
                status, result = freeze_evaluation_set(payload, identity)
                self.send_json(status, result)
                return
            if parsed.path.endswith("evaluation-baselines"):
                status, result = set_evaluation_baseline(payload, identity)
                self.send_json(status, result)
                return
            if calibration_path:
                run_id = parsed.path.split("/")[4]
                status, result = record_human_calibration(run_id, payload, identity)
                self.send_json(status, result)
                return
            if parsed.path.endswith("/external-refresh"):
                try:
                    self.send_json(202, process_external_refresh_event(payload))
                except (KeyError, ValueError) as exc:
                    self.send_json(422, {"error": str(exc)})
                except Exception:
                    self.send_json(500, {"error": "External refresh trigger failed"})
                return
            if parsed.path.endswith("/version-change"):
                try:
                    self.send_json(202, process_component_version_event(payload))
                except (KeyError, ValueError) as exc:
                    self.send_json(422, {"error": str(exc)})
                except Exception:
                    self.send_json(500, {"error": "Version change trigger failed"})
                return
            try:
                self.send_json(201, run_frozen_evaluation(payload.get("evaluationSetId"), identity))
            except ValueError as exc:
                self.send_json(409, {"error": str(exc)})
            except RuntimeError as exc:
                self.send_json(503, {"error": str(exc)})
            except Exception:
                self.send_json(500, {"error": "Evaluation execution failed"})
            return
        if parsed.path == "/api/upload-batches":
            try:
                require_production_ready()
            except PermissionError as exc:
                self.send_json(503, {"error": str(exc)})
                return
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            try:
                length = int(self.headers.get("Content-Length", "0"))
            except ValueError:
                self.send_json(400, {"error": "Invalid Content-Length"})
                return
            if length <= 0:
                self.send_json(400, {"error": "Choose one or more files to upload"})
                return
            if length > MAX_BATCH_BYTES + 1024 * 1024:
                self.send_json(413, {"error": "Batch exceeds 30 MB"})
                return
            try:
                source_type = (parse_qs(parsed.query).get("source_type") or [""])[0]
                if source_type not in {"survey", "it_data"}:
                    raise UploadValidationError("source_type must be survey or it_data")
                files = parse_multipart_files(self.headers.get("Content-Type", ""), self.rfile.read(length))
                manifest = prepare_batch(files, UPLOAD_STORAGE_DIR, owner_id=identity["id"])
                manifest["batch"]["sourceType"] = source_type
                _write_upload_batch_manifest(manifest)
                persist_upload_batch_lineage(manifest)
            except UploadValidationError as exc:
                self.send_json(400, {"error": str(exc)})
                return
            except ValueError as exc:
                self.send_json(400, {"error": str(exc)})
                return
            except Exception:
                self.send_json(500, {"error": "The upload could not be processed"})
                return
            self.send_json(201, manifest)
            return

        if parsed.path.startswith("/api/upload-batches/") and parsed.path.endswith("/ai-review"):
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            try:
                batch_id = parsed.path.split("/")[3]
                require_owned_upload_batch(batch_id, identity)
                status, payload = run_it_data_flow_review(batch_id)
                self.send_json(status, payload)
            except FileNotFoundError:
                self.send_json(404, {"error": "Uploaded batch was not found"})
            except PermissionError as exc:
                self.send_json(403, {"error": str(exc)})
            except (UploadValidationError, ValueError) as exc:
                self.send_json(422, {"error": str(exc)})
            return

        if parsed.path.startswith("/api/upload-batches/") and parsed.path.endswith("/analysis/review"):
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            try:
                payload = json.loads(self.rfile.read(int(self.headers.get("Content-Length", "0"))).decode("utf-8"))
            except (ValueError, json.JSONDecodeError):
                self.send_json(400, {"error": "Invalid JSON"})
                return
            status, response = review_dataset_run(parsed.path.split("/")[3], identity, payload)
            self.send_json(status, response)
            return

        if parsed.path.startswith("/api/upload-batches/") and (parsed.path.endswith("/publish") or parsed.path.endswith("/analysis/retry")):
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            status, payload = publish_dataset_run(parsed.path.split("/")[3], identity)
            self.send_json(status, payload)
            return

        if parsed.path.startswith("/api/upload-batches/") and parsed.path.endswith("/save-cleaned"):
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity:
                self.send_json(401, {"error": "Sign in is required"})
                return
            try:
                batch_id = parsed.path.split("/")[3]
                require_owned_upload_batch(batch_id, identity)
                manifest, output_path = save_cleaned_batch(UPLOAD_STORAGE_DIR, batch_id)
                self.send_json(200, {"batch": manifest["batch"], "savedFile": output_path.name})
            except FileNotFoundError:
                self.send_json(404, {"error": "Uploaded batch was not found"})
            except PermissionError as exc:
                self.send_json(403, {"error": str(exc)})
            except UploadValidationError as exc:
                self.send_json(422, {"error": str(exc)})
            except ValueError:
                self.send_json(400, {"error": "Invalid batch identifier"})
            return

        if parsed.path.startswith("/api/review-items/") and parsed.path.endswith("/actions"):
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity or identity.get("role") != "administrator":
                self.send_json(403, {"error": "Administrator access is required"})
                return
            try:
                payload = json.loads(self.rfile.read(int(self.headers.get("Content-Length", "0"))).decode("utf-8"))
            except (ValueError, json.JSONDecodeError):
                self.send_json(400, {"error": "Invalid JSON"})
                return
            item_id = parsed.path.split("/")[-2]
            try:
                status, response = apply_review_action(item_id, payload, identity)
            except (ValueError, KeyError):
                status, response = 400, {"error": "Invalid review action payload"}
            self.send_json(status, response)
            return

        if parsed.path.startswith("/api/review-items/") and parsed.path.endswith("/status"):
            identity = identity_from_authorization(self.headers.get("Authorization"))
            if not identity or identity.get("role") != "administrator":
                self.send_json(403, {"error": "Administrator access is required"})
                return
            length = int(self.headers.get("Content-Length", "0"))
            try:
                payload = json.loads(self.rfile.read(length).decode("utf-8"))
            except json.JSONDecodeError:
                self.send_json(400, {"error": "Invalid JSON"})
                return
            item_id = parsed.path.split("/")[-2]
            status, response = update_review_item_status(item_id, payload)
            self.send_json(status, response)
            return

        if parsed.path == "/api/review-runs/validate-chat":
            length = int(self.headers.get("Content-Length", "0"))
            if length > 50000:
                self.send_json(413, {"error": "Request is too large"})
                return
            try:
                payload = json.loads(self.rfile.read(length).decode("utf-8"))
            except json.JSONDecodeError:
                self.send_json(400, {"error": "Invalid JSON"})
                return
            status, response = run_validation_harness(payload)
            self.send_json(status, response)
            return

        if parsed.path != "/api/chat":
            self.send_json(404, {"error": "Not found"})
            return

        length = int(self.headers.get("Content-Length", "0"))
        if length > 20000:
            self.send_json(413, {"error": "Request is too large"})
            return

        try:
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
        except json.JSONDecodeError:
            self.send_json(400, {"error": "Invalid JSON"})
            return

        if not str(payload.get("message", "")).strip():
            self.send_json(400, {"error": "Message is required"})
            return

        identity = identity_from_authorization(self.headers.get("Authorization"))
        if not identity:
            self.send_json(401, {"error": "Sign in is required"})
            return
        payload["_identity"] = identity
        if payload.get("uploadBatchId"):
            try:
                payload["_uploadedBatchContext"] = uploaded_batch_chat_context(payload["uploadBatchId"], identity)
            except FileNotFoundError:
                self.send_json(404, {"error": "Uploaded batch was not found"})
                return
            except PermissionError as exc:
                self.send_json(403, {"error": str(exc)})
                return
            except (UploadValidationError, ValueError) as exc:
                self.send_json(422, {"error": str(exc)})
                return

        request_id = str(payload.get("sessionId") or uuid.uuid4())
        graph = build_conversation_graph(ServerConversationServices(payload, identity), checkpointer=MIA_GRAPH_CHECKPOINTER)
        synthesis_started = time.perf_counter()
        trace_state = {
            "graph": "conversation", "graph_version": "mia-graph-v1", "request_id": request_id,
            "retrieval_version": "hybrid-retrieval-v1", "prompt_version": "conversation-answer-v1", "config_version": "mia-config-v1",
        }
        with RedactedTracer(configured_langfuse_client()).span("mia.conversation.run", trace_state) as span:
            graph_state = graph.invoke(
                {"message": str(payload["message"]), "actor_id": str(identity["id"]), "request_id": request_id, "graph_version": "mia-graph-v1", "idempotency_key": hashlib.sha256(f"conversation:{identity['id']}:{request_id}".encode("utf-8")).hexdigest(), "trace_id": span.trace_id},
                config={"configurable": {"thread_id": f"conversation:{request_id}"}},
            )
        status = graph_state.get("status", 500)
        result = graph_state.get("result", {"error": "Conversation graph did not return an answer"})
        rag_context = graph_state.get("rag_context", {"retrieval": {"layers": []}})
        rag_context["retrieval"]["layers"].append(
            {"name": "response_synthesis", "durationMs": round((time.perf_counter() - synthesis_started) * 1000, 1), "records": 1 if status == 200 else 0}
        )
        result["graph"] = graph_audit_record({
            "graph": "conversation", "graph_version": "mia-graph-v1", "request_id": graph_state.get("request_id"), "trace_id": graph_state.get("trace_id"),
            "status": graph_state.get("publication_state", "failed"), "evidence_ids": graph_state.get("evidence_ids", []),
            "verdict": (graph_state.get("review") or {}).get("reasons", []),
        })
        record_graph_audit(result["graph"])
        result["review"] = graph_state.get("review", {})
        result["publicationState"] = graph_state.get("publication_state", "failed")
        if result["publicationState"] == "pending_human_review":
            try:
                review_item_id = queue_graph_review_item(
                    graph_name="conversation",
                    graph_run_id=str(graph_state.get("request_id") or request_id),
                    trace_id=str(result["graph"].get("trace_id") or ""),
                    title="Mia conversation requires human review",
                    reason="; ".join((graph_state.get("review") or {}).get("reasons") or ["Review policy requires approval"]),
                    recommendation=str((graph_state.get("answer") or {}).get("recommendation") or ""),
                    evidence_ids=[str(item) for item in graph_state.get("evidence_ids") or []],
                    route_key=str(payload.get("routeKey") or "all"),
                )
                result["reviewItemId"] = review_item_id
            except Exception:
                result["reviewPersistence"] = "unavailable"
        if status == 200:
            usage = measure_usage(payload, result.get("answer", ""), rag_context, result.get("providerUsage"))
            session_id = store_chat_turn(payload, result.get("answer", ""), rag_context, identity)
            record_rag_usage(identity, session_id, result.get("answer", ""), rag_context, usage)
            result["sessionId"] = session_id
            result["usage"] = usage
            try:
                validation_case = build_chat_validation_case(payload, result, rag_context, session_id=session_id)
                validation_status, validation_result = run_validation_harness(
                    {
                        "run_key": f"chat-{session_id}",
                        "cases": [validation_case],
                    }
                )
                if validation_status == 200:
                    result["validation"] = validation_result
                else:
                    result["validation"] = {
                        "error": validation_result.get("error", "Validation run failed"),
                        "status": validation_status,
                    }
            except Exception as exc:
                result["validation"] = {
                    "error": f"Validation harness failed: {exc}",
                }
        trace_state.update({
            "trace_id": graph_state.get("trace_id"),
            "status": status,
            "publication_state": result.get("publicationState"),
            "evidence_ids": graph_state.get("evidence_ids", []),
            "model_usage": result.get("usage") or result.get("providerUsage") or {},
            "validation_verdict": result.get("validation") or {},
            "final_outcome": result.get("publicationState") or ("returned" if status == 200 else "failed"),
        })
        record_graph_result_trace(configured_langfuse_client(), name="mia.conversation.result", state=trace_state, trace_id=graph_state.get("trace_id"))
        self.send_json(status, result)

    def do_DELETE(self):
        parsed = urlparse(self.path)
        if not (parsed.path.startswith("/api/upload-batches/") and parsed.path.count("/") == 3):
            self.send_json(404, {"error": "Not found"})
            return
        identity = identity_from_authorization(self.headers.get("Authorization"))
        if not identity:
            self.send_json(401, {"error": "Sign in is required"})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8")) if length else {}
        except (ValueError, json.JSONDecodeError):
            self.send_json(400, {"error": "Invalid JSON"})
            return
        status, response = delete_upload_batch(
            parsed.path.split("/")[3], identity, str(payload.get("reason") or "User deletion request")[:500]
        )
        self.send_json(status, response)


def main():
    global MIA_GRAPH_CHECKPOINTER
    load_local_env()
    database_ready = False
    try:
        database_ready = initialize_database()
    except Exception as exc:
        print(f"Database initialization skipped: {exc}")

    with checkpoint_session() as checkpointer:
        MIA_GRAPH_CHECKPOINTER = checkpointer
        if database_ready:
            backfill_upload_batch_lineage()
        resume_dataset_run_outbox()
        port = int(os.environ.get("PORT", "8766"))
        # Render and other container platforms probe the public container interface.
        host = os.environ.get("HOST", "0.0.0.0")
        server = ThreadingHTTPServer((host, port), Handler)
        print(f"Mia's Cruises dashboard listening on {host}:{port}")
        print(f"Database-backed Evidence Knowledge Base: {'ready' if database_ready else 'not connected'}")
        print("Set DEEPSEEK_API_KEY in .env or the shell for live LLM answers.")
        server.serve_forever()


if __name__ == "__main__":
    main()
